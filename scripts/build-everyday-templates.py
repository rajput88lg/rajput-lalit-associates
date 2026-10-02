"""
Builds the "everyday" Excel templates (personal finance + small business)
into /private-downloads.

    pip install openpyxl python-docx
    python scripts/build-everyday-templates.py

Reuses the styling helpers of scripts/build-digital-products.py so every kit
looks the same. Files in /private-downloads are NOT public — buyers get them
only through signed links from /api/download (see lib/paidServices.ts).
"""

import importlib.util
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.formatting.rule import CellIsRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.dont_write_bytecode = True  # keep scripts/__pycache__ out of the repo
_spec = importlib.util.spec_from_file_location(
    "kits", Path(__file__).resolve().parent / "build-digital-products.py"
)
kits = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(kits)

OUT, NAVY, GOLD = kits.OUT, kits.NAVY, kits.GOLD
BOX, H_FONT, H_FILL, IN_FILL, TOT_FILL = kits.BOX, kits.H_FONT, kits.H_FILL, kits.IN_FILL, kits.TOT_FILL
MONEY, DATE = kits.MONEY, kits.DATE
header, title, how_to, label_rows, input_block, money_cols = (
    kits.header, kits.title, kits.how_to, kits.label_rows, kits.input_block, kits.money_cols)

RED_FILL = PatternFill("solid", fgColor="FDE2E1")
GREEN_FILL = PatternFill("solid", fgColor="E7F5EC")
AMBER_FILL = PatternFill("solid", fgColor="FFF4D6")
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
HELP = ("Tax, GST or accounts question? Rajput Lalit & Associates — WhatsApp +91 93549 53603 · "
        "www.rajputlalitassociates.in")


def dv_list(ws, formula, ref, blank=True):
    dv = DataValidation(type="list", formula1=formula, allow_blank=blank)
    ws.add_data_validation(dv)
    dv.add(ref)
    return dv


def list_sheet(wb, name, heading, items, width=30):
    """Editable list (categories etc.) — returns an absolute range for drop-downs."""
    ws = wb.create_sheet(name)
    ws["A1"] = heading
    ws["A1"].font = Font(bold=True, color=NAVY)
    ws.column_dimensions["A"].width = width
    for i, it in enumerate(items, start=2):
        c = ws.cell(row=i, column=1, value=it)
        c.fill = IN_FILL
        c.border = BOX
    last = len(items) + 11  # room to add 10 more
    for r in range(len(items) + 2, last + 1):
        ws.cell(row=r, column=1).fill = IN_FILL
        ws.cell(row=r, column=1).border = BOX
    return ws, f"'{name}'!$A$2:$A${last}", last


def tile(ws, row, label, formula, fmt=MONEY, note=None):
    ws.cell(row=row, column=1, value=label).font = Font(bold=True)
    c = ws.cell(row=row, column=2, value=formula)
    c.number_format = fmt
    c.font = Font(bold=True, size=12, color=NAVY)
    c.fill = TOT_FILL
    c.border = BOX
    if note:
        ws.cell(row=row, column=3, value=note).font = Font(italic=True, color="6B7280")
    ws.row_dimensions[row].height = 22


def footer(ws, row, span=6):
    c = ws.cell(row=row, column=1, value=HELP)
    c.font = Font(size=9, bold=True, color=GOLD)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)


# ======================================================== 1. Budget planner --

def budget_planner():
    wb = Workbook()
    how_to(wb, "Monthly & Annual Budget Planner", [
        "In 'Setup' type the year you are planning for and your monthly budget for each category. Rename or add "
        "categories in the 'Categories' sheet — each one is marked Need, Want or Savings.",
        "Every time money comes in or goes out, add one line in 'Transactions' (date, category, amount). That's the "
        "only regular work.",
        "'Month View' — pick a month at the top to see budget vs actual for every category, and what is left.",
        "'Year Summary' shows all 12 months side by side, your savings rate and a chart.",
        "'50-30-20 Check' compares your spending with the popular 50% needs / 30% wants / 20% savings rule.",
    ], ["Works in Excel, Google Sheets and LibreOffice.", "Amounts are in rupees but work in any currency."])

    cats = [("Salary", "Income"), ("Business / freelance income", "Income"), ("Rent received", "Income"),
            ("Interest & dividends", "Income"), ("Other income", "Income"),
            ("Rent / home loan EMI", "Need"), ("Groceries", "Need"), ("Electricity, gas & water", "Need"),
            ("Mobile & internet", "Need"), ("Fuel & transport", "Need"), ("School / college fees", "Need"),
            ("Medical & insurance", "Need"), ("Other loan EMIs", "Need"), ("Household help", "Need"),
            ("Eating out & ordering in", "Want"), ("Shopping & clothes", "Want"), ("Entertainment & OTT", "Want"),
            ("Travel & holidays", "Want"), ("Gifts & functions", "Want"), ("Personal care", "Want"),
            ("SIP / mutual funds", "Savings"), ("PPF / EPF / NPS", "Savings"), ("FD / RD", "Savings"),
            ("Emergency fund", "Savings"), ("Miscellaneous", "Want")]
    ct = wb.create_sheet("Categories")
    title(ct, "Categories", "Edit names freely. Type: Income, Need, Want or Savings.", 3)
    header(ct, 4, ["Category", "Type", "Monthly budget (₹)"], [34, 12, 18])
    n_rows = 40
    for i in range(n_rows):
        r = 5 + i
        if i < len(cats):
            ct.cell(row=r, column=1, value=cats[i][0])
            ct.cell(row=r, column=2, value=cats[i][1])
        input_block(ct, r, r, (1, 2, 3))
        ct.cell(row=r, column=3).number_format = MONEY
    last = 4 + n_rows
    dv_list(ct, '"Income,Need,Want,Savings"', f"B5:B{last}")
    cat_rng = f"Categories!$A$5:$A${last}"

    st = wb.create_sheet("Setup", 1)
    title(st, "Setup", None, 2)
    label_rows(st, 3, [("Your name", ""), ("Year", 2026), ("Starting bank balance (optional)", 0)])
    st["B5"].number_format = MONEY
    st["A7"] = "Monthly budgets are entered in the 'Categories' sheet, column C."
    st["A7"].font = Font(italic=True, color="6B7280")

    tx = wb.create_sheet("Transactions", 2)
    title(tx, "Transactions", "One line per income or expense. Type and month fill in by themselves.", 7)
    header(tx, 4, ["Date", "Category", "Amount (₹)", "Paid by (Cash/UPI/Card/Bank)", "Note", "Type", "Month"],
           [12, 30, 14, 18, 30, 10, 10])
    N = 2000
    for r in range(5, 5 + N):
        input_block(tx, r, r, (1, 2, 3, 4, 5))
        tx.cell(row=r, column=1).number_format = DATE
        tx.cell(row=r, column=3).number_format = MONEY
        tx.cell(row=r, column=6, value=f'=IF(B{r}="","",IFERROR(INDEX(Categories!$B$5:$B${last},MATCH(B{r},{cat_rng},0)),"?"))')
        tx.cell(row=r, column=7, value=f'=IF(A{r}="","",MONTH(A{r}))')
    dv_list(tx, cat_rng, f"B5:B{4 + N}")
    dv_list(tx, '"Cash,UPI,Card,Bank"', f"D5:D{4 + N}")
    tx.freeze_panes = "A5"
    TA, TB, TC = f"Transactions!$A$5:$A${4 + N}", f"Transactions!$B$5:$B${4 + N}", f"Transactions!$C$5:$C${4 + N}"
    TF = f"Transactions!$F$5:$F${4 + N}"

    def in_month(m_cell):
        return (f'{TA},">="&DATE(Setup!$B$4,{m_cell},1),{TA},"<="&EOMONTH(DATE(Setup!$B$4,{m_cell},1),0)')

    mv = wb.create_sheet("Month View", 1)
    title(mv, "Month View — budget vs actual", "Pick the month in B3", 5)
    mv["A3"] = "Month (1–12)"
    mv["A3"].font = Font(bold=True)
    mv["B3"] = 4
    mv["B3"].fill = IN_FILL
    mv["B3"].border = BOX
    mv["C3"] = '=TEXT(DATE(Setup!B4,B3,1),"mmmm yyyy")'
    mv["C3"].font = Font(bold=True, color=NAVY, size=12)
    dv_list(mv, '"1,2,3,4,5,6,7,8,9,10,11,12"', "B3", blank=False)
    header(mv, 5, ["Category", "Type", "Budget", "Actual", "Left / (over)"], [34, 10, 16, 16, 16])
    for i in range(n_rows):
        r = 6 + i
        src = 5 + i
        mv.cell(row=r, column=1, value=f'=IF(Categories!A{src}="","",Categories!A{src})')
        mv.cell(row=r, column=2, value=f'=IF(Categories!A{src}="","",Categories!B{src})')
        mv.cell(row=r, column=3, value=f'=IF(A{r}="","",N(Categories!C{src}))')
        mv.cell(row=r, column=4, value=f'=IF(A{r}="","",SUMIFS({TC},{TB},A{r},{in_month("$B$3")}))')
        mv.cell(row=r, column=5, value=f'=IF(OR(A{r}="",B{r}="Income"),"",C{r}-D{r})')
        money_cols(mv, r, (3, 4, 5))
    end = 5 + n_rows
    mv.conditional_formatting.add(f"E6:E{end}", CellIsRule(operator="lessThan", formula=["0"], fill=RED_FILL))
    t = end + 2
    tile(mv, t, "Income this month", f'=SUMIFS(D6:D{end},B6:B{end},"Income")')
    tile(mv, t + 1, "Spent (needs + wants)", f'=SUMIFS(D6:D{end},B6:B{end},"Need")+SUMIFS(D6:D{end},B6:B{end},"Want")')
    tile(mv, t + 2, "Saved / invested", f'=SUMIFS(D6:D{end},B6:B{end},"Savings")')
    tile(mv, t + 3, "Money not yet allocated", f"=B{t}-B{t + 1}-B{t + 2}",
         note="Positive = extra cash; negative = you spent more than you earned")
    mv.freeze_panes = "A6"

    ys = wb.create_sheet("Year Summary", 2)
    title(ys, "Year Summary", "Actuals by month from 'Transactions'", 15)
    ys.column_dimensions["A"].width = 30
    ys.cell(row=4, column=1, value="Category")
    ys.cell(row=4, column=2, value="Type")
    for m in range(1, 13):
        ys.cell(row=4, column=2 + m, value=MONTHS[m - 1])
        ys.column_dimensions[get_column_letter(2 + m)].width = 11
    ys.cell(row=4, column=15, value="Total")
    ys.column_dimensions["O"].width = 13
    for c in range(1, 16):
        cell = ys.cell(row=4, column=c)
        cell.font = H_FONT
        cell.fill = H_FILL
    for i in range(n_rows):
        r = 5 + i
        src = 5 + i
        ys.cell(row=r, column=1, value=f'=IF(Categories!A{src}="","",Categories!A{src})')
        ys.cell(row=r, column=2, value=f'=IF(Categories!A{src}="","",Categories!B{src})')
        for m in range(1, 13):
            ys.cell(row=r, column=2 + m,
                    value=f'=IF($A{r}="","",SUMIFS({TC},{TB},$A{r},{in_month(m)}))').number_format = '#,##0'
        ys.cell(row=r, column=15, value=f'=IF($A{r}="","",SUM(C{r}:N{r}))').number_format = '#,##0'
    end = 4 + n_rows
    labels = [("TOTAL INCOME", "Income"), ("TOTAL NEEDS", "Need"), ("TOTAL WANTS", "Want"), ("TOTAL SAVINGS", "Savings")]
    tr = end + 2
    for j, (lab, typ) in enumerate(labels):
        r = tr + j
        ys.cell(row=r, column=1, value=lab).font = Font(bold=True)
        for c in range(3, 16):
            L = get_column_letter(c)
            cell = ys.cell(row=r, column=c, value=f'=SUMIFS({L}5:{L}{end},$B$5:$B${end},"{typ}")')
            cell.number_format = '#,##0'
            cell.fill = TOT_FILL
    r = tr + 4
    ys.cell(row=r, column=1, value="LEFT OVER (income − all outflows)").font = Font(bold=True, color=NAVY)
    for c in range(3, 16):
        L = get_column_letter(c)
        cell = ys.cell(row=r, column=c, value=f"={L}{tr}-{L}{tr + 1}-{L}{tr + 2}-{L}{tr + 3}")
        cell.number_format = '#,##0'
        cell.font = Font(bold=True)
        cell.fill = GREEN_FILL
    r2 = tr + 5
    ys.cell(row=r2, column=1, value="Savings rate").font = Font(bold=True)
    for c in range(3, 16):
        L = get_column_letter(c)
        ys.cell(row=r2, column=c, value=f'=IF({L}{tr}=0,"",{L}{tr + 3}/{L}{tr})').number_format = "0%"
    ys.freeze_panes = "C5"
    ch = BarChart()
    ch.title = "Income vs spending by month"
    ch.y_axis.title = "₹"
    ch.height, ch.width = 8, 22
    data = Reference(ys, min_col=3, max_col=14, min_row=tr, max_row=tr + 2)
    ch.add_data(data, from_rows=True, titles_from_data=False)
    ch.set_categories(Reference(ys, min_col=3, max_col=14, min_row=4))
    for ser, name in zip(ch.series, ("Income", "Needs", "Wants")):
        ser.tx = SeriesLabel(v=name)
    ys.add_chart(ch, f"A{r2 + 3}")

    fr = wb.create_sheet("50-30-20 Check", 3)
    title(fr, "50 / 30 / 20 Check — whole year", "A simple rule of thumb: 50% needs, 30% wants, 20% savings", 4)
    header(fr, 4, ["Bucket", "Your share of income", "Rule of thumb", "Comment"], [16, 20, 16, 50])
    inc = f"'Year Summary'!O{tr}"
    for i, (lab, row, rule) in enumerate((("Needs", tr + 1, 0.5), ("Wants", tr + 2, 0.3), ("Savings", tr + 3, 0.2))):
        r = 5 + i
        fr.cell(row=r, column=1, value=lab).font = Font(bold=True)
        fr.cell(row=r, column=2, value=f"=IF({inc}=0,0,'Year Summary'!O{row}/{inc})").number_format = "0%"
        fr.cell(row=r, column=3, value=rule).number_format = "0%"
        cmp = ">=" if lab == "Savings" else "<="
        fr.cell(row=r, column=4, value=f'=IF({inc}=0,"Add some transactions first",IF(B{r}{cmp}C{r},"On track ✔",'
                                       f'"{"Try to save a little more" if lab == "Savings" else "Higher than the rule — look for cuts here"}"))')
    footer(fr, 10, 4)
    wb.active = 1
    wb.save(OUT / "Budget-Planner.xlsx")


# ===================================================== 2. Debt payoff tracker --

def debt_payoff_tracker():
    wb = Workbook()
    how_to(wb, "Debt & EMI Payoff Tracker — Snowball / Avalanche", [
        "In 'My Loans' list every loan and card: amount outstanding today, yearly interest rate and the EMI / minimum "
        "payment.",
        "Choose a method: Snowball pays off the smallest balance first (quick wins); Avalanche pays the highest "
        "interest first (saves the most money).",
        "Enter the extra amount you can pay every month on top of all EMIs — even ₹1,000 makes a difference.",
        "'Dashboard' shows your debt-free date, total interest, and the month each loan closes. 'Schedule' shows "
        "every month and the chart shows your total debt going down.",
        "When a loan closes, its EMI automatically moves to the next loan in the order — that is what makes the "
        "method powerful.",
    ], ["Before prepaying a home or other loan, check prepayment charges with your bank.",
        "Interest is worked out monthly on the reducing balance — the usual method for EMIs and cards.",
        "If the minimum payment is less than the monthly interest, that loan never closes — the sheet warns you."])

    D = 10
    lo = wb.create_sheet("My Loans", 1)
    title(lo, "My Loans & Cards", "Up to 10. Yellow cells are yours to fill.", 7)
    header(lo, 5, ["#", "Loan / card", "Outstanding today (₹)", "Interest % per year", "EMI / minimum (₹)",
                   "Monthly interest now", "Sort key"], [5, 30, 20, 16, 18, 18, 10])
    lo["A3"] = "Method"
    lo["A3"].font = Font(bold=True)
    lo["B3"] = "Avalanche"
    lo["B3"].fill = IN_FILL
    lo["B3"].border = BOX
    dv_list(lo, '"Avalanche,Snowball"', "B3", blank=False)
    lo["C3"] = "Extra per month (₹)"
    lo["C3"].font = Font(bold=True)
    lo["D3"] = 2000
    lo["D3"].fill = IN_FILL
    lo["D3"].border = BOX
    lo["D3"].number_format = MONEY
    samples = [("Credit card", 60000, 36, 3000), ("Personal loan", 200000, 14, 7000), ("Car loan", 350000, 9.5, 9000)]
    for i in range(D):
        r = 6 + i
        lo.cell(row=r, column=1, value=i + 1)
        if i < len(samples):
            for j, v in enumerate(samples[i]):
                lo.cell(row=r, column=2 + j, value=v)
        input_block(lo, r, r, (2, 3, 4, 5))
        money_cols(lo, r, (3, 5, 6))
        lo.cell(row=r, column=6, value=f'=IF(C{r}="","",C{r}*D{r}/1200)')
        lo.cell(row=r, column=7, value=(f'=IF(N(C{r})<=0,1E+15,IF($B$3="Snowball",C{r},-D{r}))+ROW()/1E+6'))
    lo.column_dimensions["G"].hidden = True
    lo.conditional_formatting.add("E6:E15", FormulaRule(formula=["AND($C6>0,$E6<=$F6)"], fill=RED_FILL))
    lo["B17"] = "Red EMI = it doesn't even cover the monthly interest, so that loan will never close. Increase it."
    lo["B17"].font = Font(italic=True, color="B42318")
    lo["B18"] = "The sample rows are only an example — overwrite or clear them."
    lo["B18"].font = Font(italic=True, color="6B7280")

    # Schedule
    sc = wb.create_sheet("Schedule")
    T = 360
    title(sc, "Month-by-month schedule", "Calculated — no entries needed here", 8)
    blocks = ["Opening", "With interest", "EMI paid", "Extra paid", "Closing"]
    base = 4  # first block column
    col = {b: base + k * D for k, b in enumerate(blocks)}
    sc.cell(row=3, column=1, value="Order →")
    for i in range(D):
        lr = 6 + i
        for b in blocks:
            sc.cell(row=3, column=col[b] + i, value=f"=COUNTIF('My Loans'!$G$6:$G$15,\"<\"&'My Loans'!G{lr})+1")
            sc.cell(row=4, column=col[b] + i, value=f"='My Loans'!B{lr}&\" — {b}\"")
    sc.cell(row=4, column=1, value="Month")
    sc.cell(row=4, column=2, value="Date")
    sc.cell(row=4, column=3, value="Money for extra")
    total_col = base + 5 * D
    sc.cell(row=4, column=total_col, value="TOTAL DEBT")
    for c in range(1, total_col + 1):
        cell = sc.cell(row=4, column=c)
        cell.font = H_FONT
        cell.fill = H_FILL
        cell.alignment = Alignment(wrap_text=True)
        sc.column_dimensions[get_column_letter(c)].width = 12
    sc.row_dimensions[4].height = 45
    budget = "(SUM('My Loans'!$E$6:$E$15)+'My Loans'!$D$3)"
    L = lambda b, i: get_column_letter(col[b] + i)  # noqa: E731
    wr = f"${L('With interest', 0)}{{r}}:${L('With interest', D - 1)}{{r}}"
    mr = f"${L('EMI paid', 0)}{{r}}:${L('EMI paid', D - 1)}{{r}}"
    order = f"${L('Opening', 0)}$3:${L('Opening', D - 1)}$3"
    for t in range(1, T + 1):
        r = 4 + t
        sc.cell(row=r, column=1, value=t)
        sc.cell(row=r, column=2, value=f"=EOMONTH(TODAY(),{t - 1})").number_format = "mmm-yy"
        sc.cell(row=r, column=3, value=f"=MAX(0,{budget}-SUM({mr.format(r=r)}))").number_format = '#,##0'
        for i in range(D):
            lr = 6 + i
            o, w, m, x, c = (L(b, i) for b in blocks)
            prev = f"N('My Loans'!C{lr})" if t == 1 else f"{c}{r - 1}"
            sc[f"{o}{r}"] = f"=MAX(0,{prev})"
            sc[f"{w}{r}"] = f"=IF({o}{r}<0.5,0,{o}{r}*(1+N('My Loans'!D{lr})/1200))"
            sc[f"{m}{r}"] = f"=MIN({w}{r},N('My Loans'!E{lr}))"
            sc[f"{x}{r}"] = (f"=MIN({w}{r}-{m}{r},MAX(0,$C{r}-SUMPRODUCT(({order}<{o}$3)*"
                             f"({wr.format(r=r)}-{mr.format(r=r)}))))")
            sc[f"{c}{r}"] = f"=ROUND({w}{r}-{m}{r}-{x}{r},2)"
            for cc in (o, w, m, x, c):
                sc[f"{cc}{r}"].number_format = '#,##0'
        sc.cell(row=r, column=total_col,
                value=f"=SUM({L('Closing', 0)}{r}:{L('Closing', D - 1)}{r})").number_format = '#,##0'
    sc.freeze_panes = "D5"
    TC = get_column_letter(total_col)

    # Dashboard
    db = wb.create_sheet("Dashboard", 1)
    title(db, "Dashboard", "Updates as you change 'My Loans'", 4)
    db.column_dimensions["A"].width = 40
    db.column_dimensions["B"].width = 20
    db.column_dimensions["C"].width = 50
    tile(db, 4, "Total debt today", "=SUM('My Loans'!C6:C15)")
    tile(db, 5, "Paying every month (EMIs + extra)", f"={budget}")
    tile(db, 6, "Months to become debt-free",
         f'=IF(Schedule!{TC}{4 + T}>0.5,"More than 30 years — increase payments",COUNTIF(Schedule!{TC}5:{TC}{4 + T},">0.5")+1)',
         fmt="0")
    tile(db, 7, "Debt-free by", f'=IF(ISNUMBER(B6),EOMONTH(TODAY(),B6-1),"—")', fmt="mmmm yyyy")
    tile(db, 8, "Total interest you will pay",
         f"=SUM(Schedule!{L('With interest', 0)}5:{L('With interest', D - 1)}{4 + T})-SUM(Schedule!{L('Opening', 0)}5:{L('Opening', D - 1)}{4 + T})")
    header(db, 10, ["Loan", "Order", "Closes in month #", "Closes by"])
    for i in range(D):
        r = 11 + i
        lr = 6 + i
        cl = L("Closing", i)
        db.cell(row=r, column=1, value=f"=IF(N('My Loans'!C{lr})<=0,\"\",'My Loans'!B{lr})")
        db.cell(row=r, column=2, value=f"=IF(A{r}=\"\",\"\",Schedule!{L('Opening', i)}3)")
        db.cell(row=r, column=3, value=(f'=IF(A{r}="","",IF(Schedule!{cl}{4 + T}>0.5,"Never — EMI too low",'
                                        f'COUNTIF(Schedule!{cl}5:{cl}{4 + T},">0.5")+1))'))
        db.cell(row=r, column=4, value=f'=IF(ISNUMBER(C{r}),EOMONTH(TODAY(),C{r}-1),"")').number_format = "mmm yyyy"
    ch = LineChart()
    ch.title = "Total debt going down"
    ch.height, ch.width = 8, 18
    ch.add_data(Reference(sc, min_col=total_col, min_row=4, max_row=4 + 120), titles_from_data=True)
    ch.set_categories(Reference(sc, min_col=2, min_row=5, max_row=4 + 120))
    db.add_chart(ch, "F4")
    footer(db, 23, 4)
    wb.active = 1
    wb.save(OUT / "Debt-EMI-Payoff-Tracker.xlsx")


# ===================================================== 3. Savings goal tracker --

def savings_tracker():
    wb = Workbook()
    how_to(wb, "Savings Goal Tracker + 52-Week Challenge", [
        "In 'Goals' write each goal (emergency fund, bike, holiday, child's education…), the target amount, the date "
        "you want it by and what you have already saved.",
        "Every time you put money aside, add a line in 'Deposits' and pick the goal.",
        "'Goals' then shows progress, how much is still needed and how much to save each month to reach it on time.",
        "Want a fun habit? Try the '52-Week Challenge' — save a little more every week and tick it off.",
    ])
    g = wb.create_sheet("Goals", 1)
    title(g, "My Savings Goals", None, 9)
    header(g, 4, ["Goal", "Target (₹)", "Target date", "Saved before starting", "Saved via deposits", "Total saved",
                  "Progress", "Still needed", "Save per month to finish on time"], [28, 14, 13, 16, 16, 14, 12, 14, 20])
    examples = [("Emergency fund (6 months' expenses)", 300000), ("Holiday", 80000), ("New phone", 40000)]
    for i in range(15):
        r = 5 + i
        if i < len(examples):
            g.cell(row=r, column=1, value=examples[i][0])
            g.cell(row=r, column=2, value=examples[i][1])
        input_block(g, r, r, (1, 2, 3, 4))
        g.cell(row=r, column=3).number_format = DATE
        g.cell(row=r, column=5, value=f'=IF(A{r}="","",SUMIF(Deposits!$B$5:$B$1004,A{r},Deposits!$C$5:$C$1004))')
        g.cell(row=r, column=6, value=f'=IF(A{r}="","",N(D{r})+E{r})')
        g.cell(row=r, column=7, value=f'=IF(OR(A{r}="",N(B{r})=0),"",MIN(1,F{r}/B{r}))').number_format = "0%"
        g.cell(row=r, column=8, value=f'=IF(A{r}="","",MAX(0,B{r}-F{r}))')
        g.cell(row=r, column=9, value=(f'=IF(OR(A{r}="",C{r}=""),"",IF(H{r}=0,"Done ✔",IF(C{r}<=TODAY(),"Date passed",'
                                       f'ROUND(H{r}/MAX(1,DATEDIF(TODAY(),C{r},"m")+1),0))))'))
        money_cols(g, r, (2, 4, 5, 6, 8, 9))
    g.conditional_formatting.add("G5:G19", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1,
                                                       color="2E9E5B"))
    d = wb.create_sheet("Deposits", 2)
    title(d, "Deposits", "One line each time you save", 4)
    header(d, 4, ["Date", "Goal", "Amount (₹)", "Note"], [12, 34, 14, 30])
    for r in range(5, 1005):
        input_block(d, r, r, (1, 2, 3, 4))
        d.cell(row=r, column=1).number_format = DATE
        d.cell(row=r, column=3).number_format = MONEY
    dv_list(d, "Goals!$A$5:$A$19", "B5:B1004")
    d.freeze_panes = "A5"

    w = wb.create_sheet("52-Week Challenge", 3)
    title(w, "52-Week Savings Challenge", "Week 1 save ₹X, week 2 save 2×X … Change the starting amount in B3.", 5)
    w["A3"] = "Starting amount (₹)"
    w["A3"].font = Font(bold=True)
    w["B3"] = 100
    w["B3"].fill = IN_FILL
    w["B3"].border = BOX
    w["C3"] = "Year total if you complete it:"
    w["D3"] = "=B3*52*53/2"
    w["D3"].number_format = MONEY
    w["D3"].font = Font(bold=True, color=NAVY)
    header(w, 5, ["Week", "Save this week", "Done? (Y)", "Saved so far"], [8, 16, 12, 16])
    for k in range(1, 53):
        r = 5 + k
        w.cell(row=r, column=1, value=k)
        w.cell(row=r, column=2, value=f"=$B$3*A{r}").number_format = MONEY
        input_block(w, r, r, (3,))
        w.cell(row=r, column=4, value=f'=SUMIF($C$6:C{r},"Y",$B$6:B{r})').number_format = MONEY
    dv_list(w, '"Y"', "C6:C57")
    w.conditional_formatting.add("A6:D57", FormulaRule(formula=['$C6="Y"'], fill=GREEN_FILL))
    footer(w, 59, 4)
    wb.active = 1
    wb.save(OUT / "Savings-Goal-Tracker.xlsx")


# ======================================================== 4. Bills & EMI dates --

def bill_tracker():
    wb = Workbook()
    how_to(wb, "Bill & EMI Due Date Tracker", [
        "List every bill and EMI once in 'My Bills': amount, the day of the month it is due and how often "
        "(Monthly, Quarterly, Half-yearly, Yearly). For non-monthly bills also give the month it is due.",
        "The sheet works out the next due date for each one and sorts out what's due in the next 7 days.",
        "Tick each payment in 'Paid Register' (Y) — you'll never wonder 'did I pay the electricity bill?' again.",
        "'Summary' shows what you commit every month and every year.",
    ], ["Credit card due date = the 'payment due date' on your statement, not the statement date.",
        "Turn on auto-pay where you can, and mark it — the sheet still reminds you to keep money in the account."])
    b = wb.create_sheet("My Bills", 1)
    title(b, "My Bills & EMIs", "Status updates every time you open the file", 10)
    header(b, 4, ["Bill / EMI", "Category", "Amount (₹)", "Due day (1–31)", "How often", "Due month (non-monthly)",
                  "Auto-pay?", "Next due date", "Days left", "Status"], [26, 16, 13, 10, 13, 13, 10, 14, 10, 22])
    ex = [("Home loan EMI", "EMI", 25000, 5, "Monthly", None, "Y"), ("Electricity", "Utility", 2500, 15, "Monthly", None, "N"),
          ("Mobile postpaid", "Utility", 599, 22, "Monthly", None, "Y"), ("Health insurance", "Insurance", 24000, 10, "Yearly", 3, "N"),
          ("Property tax", "Tax", 6000, 31, "Yearly", 7, "N"), ("Advance tax", "Tax", 0, 15, "Quarterly", 6, "N")]
    cats = '"EMI,Utility,Credit card,Insurance,Rent,Tax,School fees,Subscription,SIP,Other"'
    for i in range(40):
        r = 5 + i
        if i < len(ex):
            for j, v in enumerate(ex[i]):
                b.cell(row=r, column=1 + j, value=v)
        input_block(b, r, r, range(1, 8))
        b.cell(row=r, column=3).number_format = MONEY
        f = f'IF(E{r}="Quarterly",3,IF(E{r}="Half-yearly",6,IF(E{r}="Yearly",12,1)))'
        a = f'IF(N(F{r})=0,MONTH(TODAY()),F{r})'
        n = "(YEAR(TODAY())*12+MONTH(TODAY())-1)"
        k0 = f"({n}+MOD(({a}-1)-{n},{f}))"
        dt = lambda k: f"DATE(INT({k}/12),MOD({k},12)+1,MIN(D{r},DAY(EOMONTH(DATE(INT({k}/12),MOD({k},12)+1,1),0))))"  # noqa: E731
        b.cell(row=r, column=8, value=f'=IF(OR(A{r}="",N(D{r})=0),"",IF({dt(k0)}<TODAY(),{dt(f"({k0}+{f})")},{dt(k0)}))').number_format = "dd-mmm-yy"
        b.cell(row=r, column=9, value=f'=IF(H{r}="","",H{r}-TODAY())')
        b.cell(row=r, column=10, value=(f'=IF(H{r}="","",IF(I{r}=0,"DUE TODAY",IF(I{r}<=3,"Due in "&I{r}&" day(s) — pay now",'
                                        f'IF(I{r}<=7,"Due this week","OK"))))'))
    dv_list(b, cats, "B5:B44")
    dv_list(b, '"Monthly,Quarterly,Half-yearly,Yearly"', "E5:E44")
    dv_list(b, '"1,2,3,4,5,6,7,8,9,10,11,12"', "F5:F44")
    dv_list(b, '"Y,N"', "G5:G44")
    b.conditional_formatting.add("A5:J44", FormulaRule(formula=['AND($I5<>"",$I5<=3)'], fill=RED_FILL))
    b.conditional_formatting.add("A5:J44", FormulaRule(formula=['AND($I5<>"",$I5>3,$I5<=7)'], fill=AMBER_FILL))
    b.freeze_panes = "B5"

    p = wb.create_sheet("Paid Register", 2)
    title(p, "Paid Register", "Type Y in the month you paid. Year starts in the month you choose in B3.", 14)
    p["A3"] = "First month"
    p["B3"] = "=DATE(YEAR(TODAY()),MONTH(TODAY()),1)"
    p["B3"].number_format = "mmm-yy"
    p["B3"].fill = IN_FILL
    p.column_dimensions["A"].width = 26
    p.cell(row=5, column=1, value="Bill / EMI")
    for m in range(12):
        c = p.cell(row=5, column=2 + m, value=f"=EDATE($B$3,{m})")
        c.number_format = "mmm-yy"
        p.column_dimensions[get_column_letter(2 + m)].width = 9
    for c in range(1, 14):
        cell = p.cell(row=5, column=c)
        cell.font = H_FONT
        cell.fill = H_FILL
    for i in range(40):
        r = 6 + i
        p.cell(row=r, column=1, value=f"=IF('My Bills'!A{5 + i}=\"\",\"\",'My Bills'!A{5 + i})")
        input_block(p, r, r, range(2, 14))
    dv_list(p, '"Y,—"', "B6:M45")
    p.conditional_formatting.add("B6:M45", CellIsRule(operator="equal", formula=['"Y"'], fill=GREEN_FILL))

    s = wb.create_sheet("Summary", 3)
    title(s, "Summary", None, 3)
    s.column_dimensions["A"].width = 40
    s.column_dimensions["B"].width = 18
    eq = ("SUMPRODUCT('My Bills'!$C$5:$C$44/IF('My Bills'!$E$5:$E$44=\"Quarterly\",3,IF('My Bills'!$E$5:$E$44=\"Half-yearly\",6,"
          "IF('My Bills'!$E$5:$E$44=\"Yearly\",12,1))))")
    tile(s, 4, "Average monthly commitment", f"={eq}")
    tile(s, 5, "Yearly commitment", "=B4*12")
    tile(s, 6, "Bills due in the next 7 days", "=COUNTIFS('My Bills'!I5:I44,\">=0\",'My Bills'!I5:I44,\"<=7\")", fmt="0")
    tile(s, 7, "Amount due in the next 7 days",
         "=SUMIFS('My Bills'!C5:C44,'My Bills'!I5:I44,\">=0\",'My Bills'!I5:I44,\"<=7\")")
    footer(s, 10, 3)
    wb.active = 1
    wb.save(OUT / "Bill-EMI-Due-Date-Tracker.xlsx")


# ============================================ 5. Small business income-expense --

def business_income_expense():
    wb = Workbook()
    how_to(wb, "Small Business Income & Expense Tracker", [
        "In 'Setup' enter your business name and the financial year start (1 April by default).",
        "Record every sale / receipt and every expense in 'Daily Entries' — date, Income or Expense, category, "
        "amount, cash/UPI/bank. Categories can be edited in 'Categories'.",
        "'Monthly Report' and 'Category Report' fill themselves — income, expenses and profit for each month.",
        "'Dashboard' shows the year so far: total sales, expenses, profit, cash vs digital and the best month.",
    ], ["Simple single-entry records — perfect for shops, home businesses and service providers.",
        "Need full books with GST registers and a balance sheet? See our Small Business Accounting Kit.",
        "Cash expenses above ₹10,000 to one person in a day are not allowed as a business deduction — pay digitally."])
    inc = ["Sales — counter / cash", "Sales — online / UPI", "Services", "Other income"]
    exp = ["Purchases / stock", "Rent", "Salaries & wages", "Electricity", "Mobile & internet", "Transport & fuel",
           "Packing material", "Marketing & ads", "Repairs", "Bank charges", "Loan EMI (interest part)",
           "Professional fees", "Miscellaneous"]
    c = wb.create_sheet("Categories")
    title(c, "Categories", "Edit names; keep the Type column correct", 2)
    header(c, 4, ["Category", "Type"], [32, 12])
    rows = [(x, "Income") for x in inc] + [(x, "Expense") for x in exp]
    for i in range(35):
        r = 5 + i
        if i < len(rows):
            c.cell(row=r, column=1, value=rows[i][0])
            c.cell(row=r, column=2, value=rows[i][1])
        input_block(c, r, r, (1, 2))
    dv_list(c, '"Income,Expense"', "B5:B39")
    st = wb.create_sheet("Setup", 1)
    title(st, "Setup", None, 2)
    label_rows(st, 3, [("Business name", "Your Business"), ("Financial year starts on", None), ("Opening cash + bank (₹)", 0)])
    st["B4"] = "=DATE(2026,4,1)"
    st["B4"].number_format = DATE
    st["B5"].number_format = MONEY

    de = wb.create_sheet("Daily Entries", 2)
    title(de, "Daily Entries", None, 7)
    header(de, 4, ["Date", "Income / Expense", "Category", "Details", "Amount (₹)", "Mode", "Month"],
           [12, 14, 28, 34, 14, 10, 10])
    N = 3000
    for r in range(5, 5 + N):
        input_block(de, r, r, range(1, 7))
        de.cell(row=r, column=1).number_format = DATE
        de.cell(row=r, column=5).number_format = MONEY
        de.cell(row=r, column=7, value=f'=IF(A{r}="","",TEXT(A{r},"mmm-yy"))')
    dv_list(de, '"Income,Expense"', f"B5:B{4 + N}")
    dv_list(de, "Categories!$A$5:$A$39", f"C5:C{4 + N}")
    dv_list(de, '"Cash,UPI,Bank,Card"', f"F5:F{4 + N}")
    de.freeze_panes = "A5"
    DA, DB, DC, DE, DF = (f"'Daily Entries'!${x}$5:${x}${4 + N}" for x in "ABCEF")

    mr = wb.create_sheet("Monthly Report", 1)
    title(mr, "Monthly Report", "From 'Daily Entries'", 6)
    header(mr, 4, ["Month", "Income", "Expenses", "Profit / (loss)", "Margin", "Cash receipts"], [12, 16, 16, 16, 10, 16])
    for m in range(12):
        r = 5 + m
        mr.cell(row=r, column=1, value=f"=EDATE(Setup!$B$4,{m})").number_format = "mmm-yy"
        crit = f'{DA},">="&A{r},{DA},"<="&EOMONTH(A{r},0)'
        mr.cell(row=r, column=2, value=f'=SUMIFS({DE},{crit},{DB},"Income")')
        mr.cell(row=r, column=3, value=f'=SUMIFS({DE},{crit},{DB},"Expense")')
        mr.cell(row=r, column=4, value=f"=B{r}-C{r}")
        mr.cell(row=r, column=5, value=f'=IF(B{r}=0,"",D{r}/B{r})').number_format = "0%"
        mr.cell(row=r, column=6, value=f'=SUMIFS({DE},{crit},{DB},"Income",{DF},"Cash")')
        money_cols(mr, r, (2, 3, 4, 6))
    mr.cell(row=17, column=1, value="Total").font = Font(bold=True)
    for col in (2, 3, 4, 6):
        Lc = get_column_letter(col)
        cell = mr.cell(row=17, column=col, value=f"=SUM({Lc}5:{Lc}16)")
        cell.number_format = MONEY
        cell.font = Font(bold=True)
        cell.fill = TOT_FILL
    mr.conditional_formatting.add("D5:D17", CellIsRule(operator="lessThan", formula=["0"], fill=RED_FILL))
    ch = BarChart()
    ch.title = "Income vs expenses"
    ch.height, ch.width = 8, 18
    ch.add_data(Reference(mr, min_col=2, max_col=3, min_row=4, max_row=16), titles_from_data=True)
    ch.set_categories(Reference(mr, min_col=1, min_row=5, max_row=16))
    mr.add_chart(ch, "H4")

    cr = wb.create_sheet("Category Report", 2)
    title(cr, "Category Report — whole year", None, 3)
    header(cr, 4, ["Category", "Type", "Amount", "Share of its type"], [32, 12, 16, 16])
    for i in range(35):
        r = 5 + i
        cr.cell(row=r, column=1, value=f'=IF(Categories!A{5 + i}="","",Categories!A{5 + i})')
        cr.cell(row=r, column=2, value=f'=IF(A{r}="","",Categories!B{5 + i})')
        cr.cell(row=r, column=3, value=f'=IF(A{r}="","",SUMIFS({DE},{DC},A{r},{DA},">="&Setup!$B$4,{DA},"<"&EDATE(Setup!$B$4,12)))').number_format = MONEY
        cr.cell(row=r, column=4, value=f'=IF(A{r}="","",IFERROR(C{r}/SUMIF($B$5:$B$39,B{r},$C$5:$C$39),0))').number_format = "0%"

    db = wb.create_sheet("Dashboard", 1)
    title(db, "Dashboard", "=Setup!B3", 3)
    db.column_dimensions["A"].width = 36
    db.column_dimensions["B"].width = 20
    db.column_dimensions["C"].width = 40
    tile(db, 4, "Total income (year so far)", "='Monthly Report'!B17")
    tile(db, 5, "Total expenses", "='Monthly Report'!C17")
    tile(db, 6, "Profit / (loss)", "='Monthly Report'!D17")
    tile(db, 7, "Profit margin", '=IF(B4=0,"",B6/B4)', fmt="0%")
    tile(db, 8, "Cash in hand + bank (estimated)", "=Setup!B5+B6", note="Opening balance + profit (ignores loans/drawings)")
    tile(db, 9, "Share of income received in cash", '=IF(B4=0,"",\'Monthly Report\'!F17/B4)', fmt="0%")
    tile(db, 10, "Best month", "=IFERROR(INDEX('Monthly Report'!A5:A16,MATCH(MAX('Monthly Report'!D5:D16),'Monthly Report'!D5:D16,0)),\"\")", fmt="mmmm yyyy")
    footer(db, 13, 3)
    wb.active = 1
    wb.save(OUT / "Small-Business-Income-Expense-Tracker.xlsx")


# ================================================ 6. Salary + attendance kit --

def salary_kit():
    wb = Workbook()
    how_to(wb, "Salary Sheet + Attendance + Salary Slip Kit", [
        "'Setup': company name, address and the month you are processing (any date in that month).",
        "'Employees': one row per employee — monthly Basic, HRA and other allowances, and whether PF / ESI apply. "
        "Fill once; change only when salary changes.",
        "'Attendance': mark each day with P (present), A (absent), H (half day), L (paid leave), WO (weekly off) or "
        "HO (holiday). Paid days are counted automatically.",
        "'Salary Sheet' calculates earned salary for the paid days, PF, ESI, professional tax, TDS, advances and net "
        "pay for everyone.",
        "'Salary Slip': choose an employee code in C5 — a print-ready slip appears. Print or save as PDF, then pick "
        "the next code.",
        "Next month: File → Save As with the new month's name, change the month in Setup and clear the attendance.",
    ], [
        "PF: 12% of Basic for the employee and 12% by the employer, on Basic up to ₹15,000 a month (the usual wage "
        "ceiling; you can pay on higher Basic if you choose).",
        "ESI: 0.75% employee + 3.25% employer, for employees whose gross pay is up to ₹21,000 a month.",
        "Under the new Labour Codes (in force from 21 November 2025), 'wages' (Basic + DA + retaining allowance) "
        "should be at least 50% of total pay — the Employees sheet flags anyone below 50%.",
        "Professional tax and its slabs differ by state — enter the amount that applies to you. Haryana has no "
        "professional tax.",
        "Rates as of FY 2026-27 — check with your consultant if a rate changes.",
    ])
    st = wb.create_sheet("Setup", 1)
    title(st, "Setup", None, 2)
    label_rows(st, 3, [("Company / shop name", "Your Company"), ("Address", "City, State"),
                       ("Salary month (any date in it)", None), ("PF wage ceiling (₹ per month)", 15000),
                       ("ESI limit (gross ₹ per month)", 21000)])
    st["B5"] = "=DATE(YEAR(TODAY()),MONTH(TODAY()),1)"
    st["B5"].number_format = "mmmm yyyy"
    st["A9"] = "Days in this month"
    st["B9"] = "=DAY(EOMONTH(B5,0))"
    st["A9"].font = Font(bold=True)

    E = 50
    em = wb.create_sheet("Employees", 2)
    title(em, "Employees", "Monthly (full-month) figures", 14)
    header(em, 4, ["Emp. code", "Name", "Designation", "Date of joining", "PAN", "UAN (PF)", "ESI no.", "Bank A/c & IFSC",
                   "Basic (₹)", "HRA (₹)", "Other allowances (₹)", "Gross (₹)", "PF? (Y/N)", "ESI? (Y/N)",
                   "Professional tax (₹)", "Basic share"],
           [10, 24, 16, 12, 13, 14, 14, 22, 12, 12, 14, 13, 8, 8, 12, 10])
    sample = [("E001", "Ramesh Kumar", "Salesman", None, "", "", "", "", 12000, 4800, 3200, None, "Y", "Y", 0),
              ("E002", "Sunita Devi", "Accountant", None, "", "", "", "", 16000, 6400, 2600, None, "Y", "N", 0)]
    for i in range(E):
        r = 5 + i
        if i < len(sample):
            for j, v in enumerate(sample[i]):
                if v is not None:
                    em.cell(row=r, column=1 + j, value=v)
        input_block(em, r, r, list(range(1, 12)) + [13, 14, 15])
        em.cell(row=r, column=4).number_format = DATE
        em.cell(row=r, column=12, value=f'=IF(A{r}="","",N(I{r})+N(J{r})+N(K{r}))')
        em.cell(row=r, column=16, value=f'=IF(OR(A{r}="",N(L{r})=0),"",I{r}/L{r})').number_format = "0%"
        money_cols(em, r, (9, 10, 11, 12, 15))
    dv_list(em, '"Y,N"', f"M5:N{4 + E}")
    em.conditional_formatting.add(f"P5:P{4 + E}", CellIsRule(operator="lessThan", formula=["0.5"], fill=RED_FILL))
    em.freeze_panes = "C5"

    at = wb.create_sheet("Attendance", 3)
    title(at, "Attendance", "P present · A absent · H half day · L paid leave · WO weekly off · HO holiday", 40)
    at.cell(row=4, column=1, value="Code")
    at.cell(row=4, column=2, value="Name")
    at.column_dimensions["A"].width = 9
    at.column_dimensions["B"].width = 20
    for d in range(1, 32):
        col = 2 + d
        at.cell(row=4, column=col, value=f'=IF({d}>Setup!$B$9,"",DATE(YEAR(Setup!$B$5),MONTH(Setup!$B$5),{d}))').number_format = "dd ddd"
        at.column_dimensions[get_column_letter(col)].width = 5.5
    names = ["P", "A", "H", "L", "WO+HO", "Paid days"]
    for k, nm in enumerate(names):
        at.cell(row=4, column=34 + k, value=nm)
        at.column_dimensions[get_column_letter(34 + k)].width = 8
    for c in range(1, 40):
        cell = at.cell(row=4, column=c)
        cell.font = H_FONT
        cell.fill = H_FILL
        cell.alignment = Alignment(horizontal="center", wrap_text=True, textRotation=90 if 3 <= c <= 33 else 0)
    at.row_dimensions[4].height = 50
    for i in range(E):
        r = 5 + i
        at.cell(row=r, column=1, value=f'=IF(Employees!A{r}="","",Employees!A{r})')
        at.cell(row=r, column=2, value=f'=IF(Employees!B{r}="","",Employees!B{r})')
        input_block(at, r, r, range(3, 34))
        rng = f"C{r}:AG{r}"
        at.cell(row=r, column=34, value=f'=COUNTIF({rng},"P")')
        at.cell(row=r, column=35, value=f'=COUNTIF({rng},"A")')
        at.cell(row=r, column=36, value=f'=COUNTIF({rng},"H")')
        at.cell(row=r, column=37, value=f'=COUNTIF({rng},"L")')
        at.cell(row=r, column=38, value=f'=COUNTIF({rng},"WO")+COUNTIF({rng},"HO")')
        at.cell(row=r, column=39, value=f'=IF(A{r}="","",MIN(Setup!$B$9,AH{r}+AK{r}+AL{r}+AJ{r}/2))')
    dv_list(at, '"P,A,H,L,WO,HO"', f"C5:AG{4 + E}")
    at.conditional_formatting.add(f"C5:AG{4 + E}", CellIsRule(operator="equal", formula=['"A"'], fill=RED_FILL))
    at.conditional_formatting.add(f"C5:AG{4 + E}", CellIsRule(operator="equal", formula=['"WO"'], fill=TOT_FILL))
    at.freeze_panes = "C5"

    ss = wb.create_sheet("Salary Sheet", 4)
    title(ss, "Salary Sheet", '="For "&TEXT(Setup!B5,"mmmm yyyy")', 20)
    cols = ["Code", "Name", "Paid days", "Basic earned", "HRA earned", "Other earned", "Gross earned",
            "PF (employee 12%)", "ESI (employee 0.75%)", "Professional tax", "TDS", "Advance / other deduction",
            "Total deductions", "NET PAY", "PF (employer 12%)", "ESI (employer 3.25%)", "Cost to company"]
    header(ss, 4, cols, [9, 22, 9] + [12] * 14)
    for i in range(E):
        r = 5 + i
        er = 5 + i
        ss.cell(row=r, column=1, value=f'=IF(Employees!A{er}="","",Employees!A{er})')
        ss.cell(row=r, column=2, value=f'=IF(A{r}="","",Employees!B{er})')
        ss.cell(row=r, column=3, value=f'=IF(A{r}="","",Attendance!AM{er})')
        f = f"/Setup!$B$9*C{r}"
        ss.cell(row=r, column=4, value=f'=IF(A{r}="","",ROUND(N(Employees!I{er}){f},0))')
        ss.cell(row=r, column=5, value=f'=IF(A{r}="","",ROUND(N(Employees!J{er}){f},0))')
        ss.cell(row=r, column=6, value=f'=IF(A{r}="","",ROUND(N(Employees!K{er}){f},0))')
        ss.cell(row=r, column=7, value=f'=IF(A{r}="","",D{r}+E{r}+F{r})')
        ss.cell(row=r, column=8, value=f'=IF(OR(A{r}="",Employees!M{er}<>"Y"),0,ROUND(MIN(D{r},Setup!$B$6)*12%,0))')
        ss.cell(row=r, column=9, value=(f'=IF(OR(A{r}="",Employees!N{er}<>"Y",N(Employees!L{er})>Setup!$B$7),0,'
                                        f'ROUNDUP(G{r}*0.75%,0))'))
        ss.cell(row=r, column=10, value=f'=IF(A{r}="","",N(Employees!O{er}))')
        input_block(ss, r, r, (11, 12))
        ss.cell(row=r, column=13, value=f'=IF(A{r}="","",H{r}+I{r}+J{r}+N(K{r})+N(L{r}))')
        ss.cell(row=r, column=14, value=f'=IF(A{r}="","",G{r}-M{r})')
        ss.cell(row=r, column=15, value=f'=IF(OR(A{r}="",Employees!M{er}<>"Y"),0,ROUND(MIN(D{r},Setup!$B$6)*12%,0))')
        ss.cell(row=r, column=16, value=(f'=IF(OR(A{r}="",Employees!N{er}<>"Y",N(Employees!L{er})>Setup!$B$7),0,'
                                         f'ROUNDUP(G{r}*3.25%,0))'))
        ss.cell(row=r, column=17, value=f'=IF(A{r}="","",G{r}+O{r}+P{r})')
        money_cols(ss, r, range(4, 18))
        ss.cell(row=r, column=14).font = Font(bold=True, color=NAVY)
    tr = 5 + E
    ss.cell(row=tr, column=2, value="TOTAL").font = Font(bold=True)
    for c in range(4, 18):
        Lc = get_column_letter(c)
        cell = ss.cell(row=tr, column=c, value=f"=SUM({Lc}5:{Lc}{tr - 1})")
        cell.number_format = MONEY
        cell.font = Font(bold=True)
        cell.fill = TOT_FILL
    ss.freeze_panes = "C5"

    sl = wb.create_sheet("Salary Slip", 1)
    for col, w in zip("ABCDEF", (3, 26, 18, 4, 26, 18)):
        sl.column_dimensions[col].width = w
    sl.merge_cells("B2:F2")
    sl["B2"] = "=Setup!B3"
    sl["B2"].font = Font(bold=True, size=16, color="FFFFFF")
    sl["B2"].fill = H_FILL
    sl["B2"].alignment = Alignment(horizontal="center")
    sl.merge_cells("B3:F3")
    sl["B3"] = '=Setup!B4'
    sl["B3"].alignment = Alignment(horizontal="center")
    sl.merge_cells("B4:F4")
    sl["B4"] = '="SALARY SLIP — "&UPPER(TEXT(Setup!B5,"mmmm yyyy"))'
    sl["B4"].font = Font(bold=True, color=NAVY, size=12)
    sl["B4"].alignment = Alignment(horizontal="center")
    sl["B5"] = "Employee code ▶"
    sl["B5"].font = Font(bold=True)
    sl["C5"] = "E001"
    sl["C5"].fill = IN_FILL
    sl["C5"].border = BOX
    dv_list(sl, f"Employees!$A$5:$A${4 + E}", "C5", blank=False)
    m = f"MATCH($C$5,Employees!$A$5:$A${4 + E},0)"
    ms = f"MATCH($C$5,'Salary Sheet'!$A$5:$A${4 + E},0)"
    emp = lambda col: f'=IFERROR(INDEX(Employees!${col}$5:${col}${4 + E},{m}),"")'  # noqa: E731
    sal = lambda col: f"=IFERROR(INDEX('Salary Sheet'!${col}$5:${col}${4 + E},{ms}),0)"  # noqa: E731
    info = [("Name", emp("B"), "Designation", emp("C")), ("Date of joining", emp("D"), "PAN", emp("E")),
            ("UAN", emp("F"), "ESI no.", emp("G")), ("Bank A/c", emp("H"), "Paid days",
                                                     f"=IFERROR(INDEX('Salary Sheet'!$C$5:$C${4 + E},{ms}),\"\")")]
    for i, (a, b_, c_, d_) in enumerate(info):
        r = 7 + i
        sl.cell(row=r, column=2, value=a).font = Font(bold=True)
        sl.cell(row=r, column=3, value=b_)
        sl.cell(row=r, column=5, value=c_).font = Font(bold=True)
        sl.cell(row=r, column=6, value=d_)
    sl["C8"].number_format = DATE
    for c_, t_ in ((2, "Earnings"), (3, "₹"), (5, "Deductions"), (6, "₹")):
        cell = sl.cell(row=12, column=c_, value=t_)
        cell.font = H_FONT
        cell.fill = H_FILL
    earn = [("Basic", "D"), ("HRA", "E"), ("Other allowances", "F")]
    ded = [("Provident Fund", "H"), ("ESI", "I"), ("Professional tax", "J"), ("TDS", "K"), ("Advance / other", "L")]
    for i in range(5):
        r = 13 + i
        if i < len(earn):
            sl.cell(row=r, column=2, value=earn[i][0])
            sl.cell(row=r, column=3, value=sal(earn[i][1])).number_format = MONEY
        sl.cell(row=r, column=5, value=ded[i][0])
        sl.cell(row=r, column=6, value=sal(ded[i][1])).number_format = MONEY
    sl["B18"] = "Gross earnings"
    sl["C18"] = sal("G")
    sl["E18"] = "Total deductions"
    sl["F18"] = sal("M")
    for ref in ("B18", "C18", "E18", "F18"):
        sl[ref].font = Font(bold=True)
        sl[ref].fill = TOT_FILL
    sl["C18"].number_format = MONEY
    sl["F18"].number_format = MONEY
    sl.merge_cells("B20:E20")
    sl["B20"] = "NET PAY"
    sl["B20"].font = Font(bold=True, size=13, color=NAVY)
    sl["F20"] = sal("N")
    sl["F20"].number_format = MONEY
    sl["F20"].font = Font(bold=True, size=13, color=NAVY)
    sl["F20"].fill = GREEN_FILL
    sl.merge_cells("B23:F23")
    sl["B23"] = "This is a computer-generated salary slip."
    sl["B23"].font = Font(italic=True, size=9, color="6B7280")
    sl["E26"] = "Authorised signatory"
    sl.page_setup.paperSize = sl.PAPERSIZE_A4
    sl.sheet_properties.pageSetUpPr.fitToPage = True
    sl.page_setup.fitToHeight = 1
    sl.print_area = "B2:F27"
    wb.active = 1
    wb.save(OUT / "Salary-Attendance-Salary-Slip-Kit.xlsx")


# ======================================================== 7. Inventory tracker --

def inventory_tracker():
    wb = Workbook()
    how_to(wb, "Inventory & Stock Tracker", [
        "List your products once in 'Items': code, name, unit, opening stock, purchase rate, selling rate and the "
        "reorder level (the quantity at which you want to re-order).",
        "Every purchase / goods received → one line in 'Stock In'. Every sale / goods issued → one line in "
        "'Stock Out'. Pick the item code from the drop-down.",
        "'Stock Summary' shows current stock, its value and which items to reorder (red) or are finished.",
        "'Dashboard' gives total stock value, items to reorder and the month's sales.",
    ], ["Stock value is at purchase rate — the way it is usually valued for accounts (cost or market, whichever "
        "is lower).", "Do a physical count every few months and adjust with an entry so the sheet matches the shelf."])
    I = 300
    it = wb.create_sheet("Items", 1)
    title(it, "Items", None, 9)
    header(it, 4, ["Item code", "Item name", "Category", "Unit", "Opening stock", "Purchase rate (₹)", "Selling rate (₹)",
                   "Reorder level", "HSN (optional)"], [12, 30, 16, 8, 12, 14, 14, 12, 12])
    ex = [("A001", "Sample item 1", "General", "Pcs", 50, 80, 120, 10), ("A002", "Sample item 2", "General", "Kg", 20, 45, 60, 5)]
    for i in range(I):
        r = 5 + i
        if i < len(ex):
            for j, v in enumerate(ex[i]):
                it.cell(row=r, column=1 + j, value=v)
        input_block(it, r, r, range(1, 10))
        money_cols(it, r, (6, 7))
    dv_list(it, '"Pcs,Kg,Gm,Ltr,Mtr,Box,Dozen,Pair,Set,Bag"', f"D5:D{4 + I}")
    it.freeze_panes = "B5"
    codes = f"Items!$A$5:$A${4 + I}"

    def mov(name, who):
        ws = wb.create_sheet(name)
        title(ws, name, None, 7)
        header(ws, 4, ["Date", "Item code", "Item name", "Quantity", "Rate (₹)", "Amount (₹)", who],
               [12, 12, 28, 10, 12, 14, 26])
        for r in range(5, 3005):
            input_block(ws, r, r, (1, 2, 4, 5, 7))
            ws.cell(row=r, column=1).number_format = DATE
            ws.cell(row=r, column=3, value=f'=IF(B{r}="","",IFERROR(INDEX(Items!$B$5:$B${4 + I},MATCH(B{r},{codes},0)),"Unknown code"))')
            ws.cell(row=r, column=6, value=f'=IF(OR(D{r}="",E{r}=""),"",D{r}*E{r})')
            money_cols(ws, r, (5, 6))
        dv_list(ws, codes, "B5:B3004")
        ws.freeze_panes = "A5"
    mov("Stock In", "Supplier")
    mov("Stock Out", "Customer")
    su = wb.create_sheet("Stock Summary", 2)
    title(su, "Stock Summary", "Updates by itself", 9)
    header(su, 4, ["Item code", "Item name", "Opening", "In", "Out", "In stock", "Stock value (₹)", "Status", "Sales value (₹)"],
           [12, 28, 10, 10, 10, 10, 16, 18, 16])
    for i in range(I):
        r = 5 + i
        src = 5 + i
        su.cell(row=r, column=1, value=f'=IF(Items!A{src}="","",Items!A{src})')
        su.cell(row=r, column=2, value=f'=IF(A{r}="","",Items!B{src})')
        su.cell(row=r, column=3, value=f'=IF(A{r}="","",N(Items!E{src}))')
        su.cell(row=r, column=4, value=f"=IF(A{r}=\"\",\"\",SUMIF('Stock In'!$B$5:$B$3004,A{r},'Stock In'!$D$5:$D$3004))")
        su.cell(row=r, column=5, value=f"=IF(A{r}=\"\",\"\",SUMIF('Stock Out'!$B$5:$B$3004,A{r},'Stock Out'!$D$5:$D$3004))")
        su.cell(row=r, column=6, value=f'=IF(A{r}="","",C{r}+D{r}-E{r})')
        su.cell(row=r, column=7, value=f'=IF(A{r}="","",F{r}*N(Items!F{src}))').number_format = MONEY
        su.cell(row=r, column=8, value=f'=IF(A{r}="","",IF(F{r}<=0,"OUT OF STOCK",IF(F{r}<=N(Items!H{src}),"REORDER","OK")))')
        su.cell(row=r, column=9, value=f"=IF(A{r}=\"\",\"\",SUMIF('Stock Out'!$B$5:$B$3004,A{r},'Stock Out'!$F$5:$F$3004))").number_format = MONEY
    su.conditional_formatting.add(f"A5:I{4 + I}", FormulaRule(formula=['$H5="OUT OF STOCK"'], fill=RED_FILL))
    su.conditional_formatting.add(f"A5:I{4 + I}", FormulaRule(formula=['$H5="REORDER"'], fill=AMBER_FILL))
    su.auto_filter.ref = f"A4:I{4 + I}"
    su.freeze_panes = "C5"
    db = wb.create_sheet("Dashboard", 1)
    title(db, "Dashboard", None, 3)
    db.column_dimensions["A"].width = 36
    db.column_dimensions["B"].width = 18
    db.column_dimensions["C"].width = 44
    tile(db, 4, "Total stock value (at cost)", f"=SUM('Stock Summary'!G5:G{4 + I})")
    tile(db, 5, "Items to reorder", f"=COUNTIF('Stock Summary'!H5:H{4 + I},\"REORDER\")", fmt="0",
         note="Filter 'Stock Summary' column H to see them")
    tile(db, 6, "Items out of stock", f"=COUNTIF('Stock Summary'!H5:H{4 + I},\"OUT OF STOCK\")", fmt="0")
    tile(db, 7, "Sales this month", "=SUMIFS('Stock Out'!F5:F3004,'Stock Out'!A5:A3004,\">=\"&DATE(YEAR(TODAY()),MONTH(TODAY()),1))")
    tile(db, 8, "Purchases this month", "=SUMIFS('Stock In'!F5:F3004,'Stock In'!A5:A3004,\">=\"&DATE(YEAR(TODAY()),MONTH(TODAY()),1))")
    footer(db, 11, 3)
    wb.active = 1
    wb.save(OUT / "Inventory-Stock-Tracker.xlsx")


# =========================================================== 8. Wedding budget --

def wedding_planner():
    wb = Workbook()
    how_to(wb, "Wedding Budget Planner", [
        "In 'Setup' enter the total budget, the wedding date and the expected number of guests.",
        "'Budget' already has the usual heads of an Indian wedding with a suggested share of the budget — change the "
        "estimates to suit you.",
        "Record every vendor and every payment in 'Vendors & Payments'. Spent and balance per head fill in by "
        "themselves.",
        "Keep the guest list (with RSVP and number of people) in 'Guests' — catering numbers come from here.",
        "'Checklist' tells you what to finish and by when, counted back from the wedding date.",
    ], ["Ask vendors for a written quote and receipt for every advance.",
        "Pay big amounts digitally — under income tax rules a vendor cannot accept ₹2 lakh or more in cash from you "
        "in a day or for one event."])
    st = wb.create_sheet("Setup", 1)
    title(st, "Setup", None, 2)
    label_rows(st, 3, [("Couple", "Bride & Groom"), ("Wedding date", None), ("Total budget (₹)", 1500000),
                       ("Expected guests", 300), ("Cost per plate (₹)", 900)])
    st["B4"] = "=DATE(2027,2,14)"
    st["B4"].number_format = DATE
    st["B5"].number_format = MONEY
    st["B7"].number_format = MONEY
    st["A9"] = "Days to go"
    st["A9"].font = Font(bold=True)
    st["B9"] = "=MAX(0,B4-TODAY())"
    st["B9"].font = Font(bold=True, size=14, color=GOLD)
    heads = [("Venue & tent", 0.15), ("Catering (plates × guests)", None), ("Decoration & flowers", 0.08),
             ("Bride's clothes & makeup", 0.08), ("Groom's clothes", 0.03), ("Jewellery", 0.12),
             ("Photography & video", 0.06), ("Music / DJ / band", 0.03), ("Mehendi & haldi", 0.03),
             ("Sangeet / cocktail", 0.04), ("Invitations & gifts for guests", 0.03), ("Travel & stay for guests", 0.04),
             ("Pandit / rituals", 0.01), ("Transport (car, ghodi, baraat)", 0.02), ("Shagun / gifts to family", 0.03),
             ("Miscellaneous / buffer", 0.05)]
    bu = wb.create_sheet("Budget", 2)
    title(bu, "Budget by head", None, 7)
    header(bu, 4, ["Head", "Suggested share", "Estimate (₹)", "Committed (quotes)", "Paid so far", "Balance to pay",
                   "Over / (under) estimate"], [32, 14, 16, 18, 16, 16, 18])
    for i in range(25):
        r = 5 + i
        if i < len(heads):
            h, share = heads[i]
            bu.cell(row=r, column=1, value=h)
            if share is not None:
                bu.cell(row=r, column=2, value=share).number_format = "0%"
                bu.cell(row=r, column=3, value=f"=ROUND(Setup!$B$5*B{r},-3)")
            else:
                bu.cell(row=r, column=2, value="=IF(Setup!$B$5=0,0,C6/Setup!$B$5)").number_format = "0%"
                bu.cell(row=r, column=3, value="=Setup!B6*Setup!B7")
        input_block(bu, r, r, (1, 3))
        V = "'Vendors & Payments'"
        bu.cell(row=r, column=4, value=f'=IF(A{r}="","",SUMIF({V}!$B$5:$B$204,A{r},{V}!$E$5:$E$204))')
        bu.cell(row=r, column=5, value=f'=IF(A{r}="","",SUMIF({V}!$B$5:$B$204,A{r},{V}!$I$5:$I$204))')
        bu.cell(row=r, column=6, value=f'=IF(A{r}="","",D{r}-E{r})')
        bu.cell(row=r, column=7, value=f'=IF(A{r}="","",D{r}-N(C{r}))')
        money_cols(bu, r, (3, 4, 5, 6, 7))
    bu.cell(row=30, column=1, value="TOTAL").font = Font(bold=True)
    for c in (3, 4, 5, 6, 7):
        Lc = get_column_letter(c)
        cell = bu.cell(row=30, column=c, value=f"=SUM({Lc}5:{Lc}29)")
        cell.number_format = MONEY
        cell.font = Font(bold=True)
        cell.fill = TOT_FILL
    bu["A32"] = '=IF(C30>Setup!B5,"Estimates are ₹"&TEXT(C30-Setup!B5,"#,##0")&" above your total budget","Estimates are within budget ✔")'
    bu["A32"].font = Font(bold=True, color=NAVY)
    bu.conditional_formatting.add("G5:G29", CellIsRule(operator="greaterThan", formula=["0"], fill=RED_FILL))
    vp = wb.create_sheet("Vendors & Payments", 3)
    title(vp, "Vendors & Payments", None, 10)
    header(vp, 4, ["Vendor", "Head", "Contact", "Booked on", "Total quote (₹)", "Advance paid", "2nd payment",
                   "Final payment", "Total paid", "Balance"], [24, 28, 16, 12, 14, 13, 13, 13, 14, 14])
    for r in range(5, 205):
        input_block(vp, r, r, range(1, 9))
        vp.cell(row=r, column=4).number_format = DATE
        vp.cell(row=r, column=9, value=f'=IF(A{r}="","",N(F{r})+N(G{r})+N(H{r}))')
        vp.cell(row=r, column=10, value=f'=IF(A{r}="","",N(E{r})-I{r})')
        money_cols(vp, r, range(5, 11))
    dv_list(vp, "Budget!$A$5:$A$29", "B5:B204")
    vp.freeze_panes = "B5"
    gu = wb.create_sheet("Guests", 4)
    title(gu, "Guest List", None, 7)
    header(gu, 4, ["Name / family", "Side", "Phone", "City", "No. of people", "RSVP", "Gift / shagun (₹)"],
           [28, 10, 14, 14, 12, 12, 14])
    for r in range(5, 805):
        input_block(gu, r, r, range(1, 8))
        gu.cell(row=r, column=7).number_format = MONEY
    dv_list(gu, '"Bride,Groom,Both"', "B5:B804")
    dv_list(gu, '"Invited,Coming,Not coming,Maybe"', "F5:F804")
    gu["I4"] = "People confirmed"
    gu["J4"] = '=SUMIF(F5:F804,"Coming",E5:E804)'
    gu["I5"] = "People invited"
    gu["J5"] = "=SUM(E5:E804)"
    gu["I6"] = "Shagun received"
    gu["J6"] = "=SUM(G5:G804)"
    gu["J6"].number_format = MONEY
    gu.column_dimensions["I"].width = 18
    gu.freeze_panes = "A5"
    ck = wb.create_sheet("Checklist", 5)
    title(ck, "Checklist", "Target dates are counted back from the wedding date", 4)
    header(ck, 4, ["Task", "Months before", "Do by", "Done? (Y)"], [46, 14, 14, 12])
    tasks = [("Fix budget and guest count", 9), ("Book venue", 8), ("Book caterer & taste menu", 7),
             ("Book photographer & decorator", 6), ("Shop jewellery", 5), ("Order outfits", 4),
             ("Book makeup artist & mehendi", 4), ("Print / send invitations", 2), ("Book guest rooms & transport", 2),
             ("Final guest count to caterer", 0.5), ("Collect outfits, final payments plan", 0.5),
             ("Pack emergency kit & cash envelopes", 0.1)]
    for i, (t, mb) in enumerate(tasks):
        r = 5 + i
        ck.cell(row=r, column=1, value=t)
        ck.cell(row=r, column=2, value=mb)
        ck.cell(row=r, column=3, value=f"=Setup!$B$4-ROUND(B{r}*30,0)").number_format = DATE
        input_block(ck, r, r, (4,))
    for r in range(5 + len(tasks), 35):
        input_block(ck, r, r, (1, 2, 4))
        ck.cell(row=r, column=3, value=f'=IF(B{r}="","",Setup!$B$4-ROUND(B{r}*30,0))').number_format = DATE
    ck.conditional_formatting.add("A5:D34", FormulaRule(formula=['$D5="Y"'], fill=GREEN_FILL))
    ck.conditional_formatting.add("A5:D34", FormulaRule(formula=['AND($D5<>"Y",$C5<>"",$C5<TODAY())'], fill=RED_FILL))
    footer(ck, 36, 4)
    wb.active = 1
    wb.save(OUT / "Wedding-Budget-Planner.xlsx")


# ========================================================= 9. Rental property --

def rental_tracker():
    wb = Workbook()
    how_to(wb, "Rental Property Tracker", [
        "In 'Properties' add each house / flat / shop: tenant, monthly rent, deposit, rent due day and agreement dates.",
        "Each time rent comes in, add it in 'Rent Received' (date, property, month it is for, amount, TDS if the "
        "tenant deducted any).",
        "Record property expenses in 'Expenses' — municipal tax, repairs, society charges, insurance, loan interest.",
        "'Summary' shows rent due vs received (arrears), expenses and net income per property for the year, and a "
        "simple house-property income working for your income tax return.",
        "Agreements expiring in the next 60 days are highlighted so you can renew or revise the rent in time.",
    ], ["Income-tax working uses the long-standing rules: municipal taxes paid are deducted, then a flat 30% "
        "standard deduction, then interest on a loan for that property. Confirm for your year before filing.",
        "Tenants paying rent above ₹50,000 a month must deduct TDS — it shows in your Form 26AS / AIS.",
        "Rent receipts and a registered agreement protect you — keep copies."])
    P = 20
    pr = wb.create_sheet("Properties", 1)
    title(pr, "Properties", None, 10)
    header(pr, 4, ["Property", "Address", "Tenant", "Tenant phone", "Monthly rent (₹)", "Security deposit (₹)",
                   "Rent due day", "Agreement start", "Agreement end", "Agreement status"],
           [20, 28, 20, 14, 14, 16, 10, 13, 13, 20])
    for i in range(P):
        r = 5 + i
        if i == 0:
            for j, v in enumerate(("Flat 101", "Sector 9, Ambala", "Tenant name", "", 15000, 30000, 5)):
                pr.cell(row=r, column=1 + j, value=v)
        input_block(pr, r, r, range(1, 10))
        money_cols(pr, r, (5, 6))
        pr.cell(row=r, column=8).number_format = DATE
        pr.cell(row=r, column=9).number_format = DATE
        pr.cell(row=r, column=10, value=(f'=IF(I{r}="","",IF(I{r}<TODAY(),"EXPIRED — renew",'
                                         f'IF(I{r}-TODAY()<=60,"Ends in "&(I{r}-TODAY())&" days","Active")))'))
    pr.conditional_formatting.add(f"J5:J{4 + P}", FormulaRule(formula=['LEFT(J5,4)="Ends"'], fill=AMBER_FILL))
    pr.conditional_formatting.add(f"J5:J{4 + P}", FormulaRule(formula=['LEFT(J5,7)="EXPIRED"'], fill=RED_FILL))
    props = f"Properties!$A$5:$A${4 + P}"
    rr = wb.create_sheet("Rent Received", 2)
    title(rr, "Rent Received", None, 7)
    header(rr, 4, ["Date received", "Property", "Rent for month", "Amount received (₹)", "TDS deducted by tenant (₹)",
                   "Mode", "Note"], [13, 20, 14, 18, 18, 10, 26])
    for r in range(5, 1205):
        input_block(rr, r, r, range(1, 8))
        rr.cell(row=r, column=1).number_format = DATE
        rr.cell(row=r, column=3).number_format = "mmm-yy"
        money_cols(rr, r, (4, 5))
    dv_list(rr, props, "B5:B1204")
    dv_list(rr, '"Bank,UPI,Cash,Cheque"', "F5:F1204")
    rr.freeze_panes = "A5"
    ex = wb.create_sheet("Expenses", 3)
    title(ex, "Property Expenses", None, 5)
    header(ex, 4, ["Date", "Property", "Expense type", "Amount (₹)", "Note"], [12, 20, 26, 14, 30])
    for r in range(5, 1005):
        input_block(ex, r, r, range(1, 6))
        ex.cell(row=r, column=1).number_format = DATE
        ex.cell(row=r, column=4).number_format = MONEY
    dv_list(ex, props, "B5:B1004")
    dv_list(ex, '"Municipal / property tax,Home loan interest,Repairs & maintenance,Society charges,Insurance,'
                'Brokerage,Other"', "C5:C1004")
    ex.freeze_panes = "A5"
    su = wb.create_sheet("Summary", 1)
    title(su, "Summary — financial year", "Set the year start in B3", 11)
    su["A3"] = "Year starts"
    su["B3"] = "=DATE(2026,4,1)"
    su["B3"].number_format = DATE
    su["B3"].fill = IN_FILL
    header(su, 5, ["Property", "Rent due (months × rent)", "Rent received", "Arrears", "TDS by tenant",
                   "Municipal tax paid", "Other expenses", "Loan interest", "Net cash income",
                   "Taxable house-property income (est.)", ""], [20, 18, 16, 14, 14, 16, 16, 14, 16, 22, 2])
    yr = lambda rng, crit_col: f'{rng},{crit_col}'  # noqa: E731
    R = "'Rent Received'"
    for i in range(P):
        r = 6 + i
        src = 5 + i
        su.cell(row=r, column=1, value=f'=IF(Properties!A{src}="","",Properties!A{src})')
        months = (f"IFERROR(MAX(0,MIN(12,DATEDIF(MAX($B$3,IF(Properties!H{src}=\"\",$B$3,Properties!H{src})),"
                  f"MIN(TODAY(),EDATE($B$3,12)-1,IF(Properties!I{src}=\"\",TODAY(),Properties!I{src})),\"m\")+1)),0)")
        su.cell(row=r, column=2, value=f'=IF(A{r}="","",N(Properties!E{src})*{months})')
        inyr_r = f'{R}!$A$5:$A$1204,">="&$B$3,{R}!$A$5:$A$1204,"<"&EDATE($B$3,12)'
        su.cell(row=r, column=3, value=f'=IF(A{r}="","",SUMIFS({R}!$D$5:$D$1204,{R}!$B$5:$B$1204,A{r},{inyr_r}))')
        su.cell(row=r, column=4, value=f'=IF(A{r}="","",MAX(0,B{r}-C{r}-E{r}))')
        su.cell(row=r, column=5, value=f'=IF(A{r}="","",SUMIFS({R}!$E$5:$E$1204,{R}!$B$5:$B$1204,A{r},{inyr_r}))')
        inyr_e = 'Expenses!$A$5:$A$1004,">="&$B$3,Expenses!$A$5:$A$1004,"<"&EDATE($B$3,12)'
        exp = lambda t: f'SUMIFS(Expenses!$D$5:$D$1004,Expenses!$B$5:$B$1004,A{r},Expenses!$C$5:$C$1004,"{t}",{inyr_e})'  # noqa: E731
        su.cell(row=r, column=6, value=f'=IF(A{r}="","",{exp("Municipal / property tax")})')
        su.cell(row=r, column=8, value=f'=IF(A{r}="","",{exp("Home loan interest")})')
        su.cell(row=r, column=7, value=(f'=IF(A{r}="","",SUMIFS(Expenses!$D$5:$D$1004,Expenses!$B$5:$B$1004,A{r},{inyr_e})'
                                        f'-F{r}-H{r})'))
        su.cell(row=r, column=9, value=f'=IF(A{r}="","",C{r}+E{r}-F{r}-G{r}-H{r})')
        su.cell(row=r, column=10, value=f'=IF(A{r}="","",ROUND((C{r}+E{r}-F{r})*70%,0)-H{r})')
        money_cols(su, r, range(2, 11))
    tr = 6 + P
    su.cell(row=tr, column=1, value="TOTAL").font = Font(bold=True)
    for c in range(2, 11):
        Lc = get_column_letter(c)
        cell = su.cell(row=tr, column=c, value=f"=SUM({Lc}6:{Lc}{tr - 1})")
        cell.number_format = MONEY
        cell.font = Font(bold=True)
        cell.fill = TOT_FILL
    su.cell(row=tr + 2, column=1, value=(
        "Taxable house-property income = (rent + TDS − municipal tax) × 70% − loan interest. A negative figure is a "
        "loss; how much of it can be set off depends on the regime and the year's limits."))
    su.cell(row=tr + 2, column=1).alignment = Alignment(wrap_text=True)
    su.merge_cells(start_row=tr + 2, start_column=1, end_row=tr + 3, end_column=10)
    su.row_dimensions[tr + 2].height = 30
    su.conditional_formatting.add(f"D6:D{tr - 1}", CellIsRule(operator="greaterThan", formula=["0"], fill=RED_FILL))
    footer(su, tr + 5, 10)
    wb.active = 1
    wb.save(OUT / "Rental-Property-Tracker.xlsx")


# ===================================================== 10. Net worth tracker --

def networth_tracker():
    wb = Workbook()
    how_to(wb, "Investment & Net Worth Tracker", [
        "In 'Assets' list everything you own: bank balances, FDs, PPF, EPF, NPS, mutual funds, shares, gold, "
        "property, crypto — with the amount invested and today's value.",
        "In 'Loans' list what you owe: home loan, car loan, personal loan, card balance.",
        "'Net Worth' shows your total, the asset mix chart and the gain on your investments.",
        "Once a month (or quarter) copy the net worth figure into 'History' — the chart shows how you are growing.",
        "FDs and policies maturing in the next 90 days are highlighted in 'Assets'.",
    ], ["Update current values from your bank / broker / CAMS statement — the sheet does not fetch prices.",
        "Write the nominee for every asset — it saves your family a lot of trouble."])
    types = '"Bank,FD / RD,PPF,EPF,NPS,Mutual funds,Shares,Gold,Property,Insurance (savings),Crypto,Other"'
    a = wb.create_sheet("Assets", 1)
    title(a, "Assets", None, 9)
    header(a, 4, ["Asset", "Type", "Where held", "Amount invested (₹)", "Current value (₹)", "Gain (₹)", "Gain %",
                  "Maturity date", "Nominee"], [26, 16, 18, 16, 16, 14, 10, 13, 18])
    ex = [("Savings account", "Bank", "SBI", 50000, 50000), ("Equity SIP", "Mutual funds", "CAMS", 200000, 236000),
          ("PPF", "PPF", "Post office", 300000, 352000)]
    for i in range(60):
        r = 5 + i
        if i < len(ex):
            for j, v in enumerate(ex[i]):
                a.cell(row=r, column=1 + j, value=v)
        input_block(a, r, r, (1, 2, 3, 4, 5, 8, 9))
        a.cell(row=r, column=6, value=f'=IF(A{r}="","",N(E{r})-N(D{r}))')
        a.cell(row=r, column=7, value=f'=IF(OR(A{r}="",N(D{r})=0),"",F{r}/D{r})').number_format = "0.0%"
        a.cell(row=r, column=8).number_format = DATE
        money_cols(a, r, (4, 5, 6))
    dv_list(a, types, "B5:B64")
    a.conditional_formatting.add("A5:I64", FormulaRule(formula=['AND($H5<>"",$H5-TODAY()<=90,$H5>=TODAY())'], fill=AMBER_FILL))
    a.freeze_panes = "B5"
    lo = wb.create_sheet("Loans", 2)
    title(lo, "Loans & Dues", None, 5)
    header(lo, 4, ["Loan", "Lender", "Outstanding today (₹)", "EMI (₹)", "Ends on"], [26, 18, 18, 14, 13])
    for r in range(5, 25):
        input_block(lo, r, r, range(1, 6))
        money_cols(lo, r, (3, 4))
        lo.cell(row=r, column=5).number_format = DATE
    nw = wb.create_sheet("Net Worth", 1)
    title(nw, "Net Worth", '="As on "&TEXT(TODAY(),"dd mmm yyyy")', 4)
    nw.column_dimensions["A"].width = 34
    nw.column_dimensions["B"].width = 18
    nw.column_dimensions["C"].width = 14
    tile(nw, 4, "Total assets (current value)", "=SUM(Assets!E5:E64)")
    tile(nw, 5, "Total loans", "=SUM(Loans!C5:C24)")
    tile(nw, 6, "NET WORTH", "=B4-B5")
    nw["B6"].fill = GREEN_FILL
    tile(nw, 7, "Gain on investments", "=SUM(Assets!F5:F64)")
    header(nw, 9, ["Asset type", "Current value", "Share"])
    tl = ["Bank", "FD / RD", "PPF", "EPF", "NPS", "Mutual funds", "Shares", "Gold", "Property", "Insurance (savings)",
          "Crypto", "Other"]
    for i, t in enumerate(tl):
        r = 10 + i
        nw.cell(row=r, column=1, value=t)
        nw.cell(row=r, column=2, value=f'=SUMIF(Assets!$B$5:$B$64,A{r},Assets!$E$5:$E$64)').number_format = MONEY
        nw.cell(row=r, column=3, value=f"=IF($B$4=0,0,B{r}/$B$4)").number_format = "0%"
    pie = PieChart()
    pie.title = "Where your money is"
    pie.height, pie.width = 9, 12
    pie.add_data(Reference(nw, min_col=2, min_row=10, max_row=21))
    pie.set_categories(Reference(nw, min_col=1, min_row=10, max_row=21))
    nw.add_chart(pie, "E4")
    hi = wb.create_sheet("History", 3)
    title(hi, "Net Worth History", "Copy the figures from 'Net Worth' once a month", 4)
    header(hi, 4, ["Date", "Total assets", "Total loans", "Net worth"], [12, 16, 16, 16])
    for r in range(5, 125):
        input_block(hi, r, r, (1, 2, 3))
        hi.cell(row=r, column=1).number_format = "mmm-yy"
        hi.cell(row=r, column=4, value=f'=IF(A{r}="","",N(B{r})-N(C{r}))')
        money_cols(hi, r, (2, 3, 4))
    lc = LineChart()
    lc.title = "Net worth over time"
    lc.height, lc.width = 8, 18
    lc.add_data(Reference(hi, min_col=4, min_row=4, max_row=64), titles_from_data=True)
    lc.set_categories(Reference(hi, min_col=1, min_row=5, max_row=64))
    hi.add_chart(lc, "F4")
    footer(nw, 24, 3)
    wb.active = 1
    wb.save(OUT / "Investment-Net-Worth-Tracker.xlsx")


if __name__ == "__main__":
    for fn in (budget_planner, debt_payoff_tracker, savings_tracker, bill_tracker, business_income_expense,
               salary_kit, inventory_tracker, wedding_planner, rental_tracker, networth_tracker):
        fn()
    for f in sorted(OUT.iterdir()):
        print(f"{f.name}: {f.stat().st_size:,} bytes")
