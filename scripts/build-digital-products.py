"""
Builds the paid download kits into /private-downloads.

    pip install openpyxl python-docx
    python scripts/build-digital-products.py

Files in /private-downloads are NOT public — buyers get them only through
signed links from /api/download after paying (see lib/paidServices.ts).
Edit the content here and re-run to update a kit.
"""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.pagebreak import Break

OUT = Path(__file__).resolve().parent.parent / "private-downloads"
OUT.mkdir(exist_ok=True)

NAVY = "002B5C"
GOLD = "D99A2B"
LIGHT = "F3F6FA"
INPUT = "FFF7E6"
FIRM = "Rajput Lalit & Associates · www.rajputlalitassociates.in · WhatsApp +91 93549 53603"

thin = Side(style="thin", color="C9D2DE")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
H_FONT = Font(bold=True, color="FFFFFF")
H_FILL = PatternFill("solid", fgColor=NAVY)
IN_FILL = PatternFill("solid", fgColor=INPUT)
TOT_FILL = PatternFill("solid", fgColor=LIGHT)
MONEY = '#,##0.00'
DATE = "dd-mm-yyyy"

STATES = [
    ("01", "Jammu and Kashmir"), ("02", "Himachal Pradesh"), ("03", "Punjab"), ("04", "Chandigarh"),
    ("05", "Uttarakhand"), ("06", "Haryana"), ("07", "Delhi"), ("08", "Rajasthan"), ("09", "Uttar Pradesh"),
    ("10", "Bihar"), ("11", "Sikkim"), ("12", "Arunachal Pradesh"), ("13", "Nagaland"), ("14", "Manipur"),
    ("15", "Mizoram"), ("16", "Tripura"), ("17", "Meghalaya"), ("18", "Assam"), ("19", "West Bengal"),
    ("20", "Jharkhand"), ("21", "Odisha"), ("22", "Chhattisgarh"), ("23", "Madhya Pradesh"), ("24", "Gujarat"),
    ("26", "Dadra and Nagar Haveli and Daman and Diu"), ("27", "Maharashtra"), ("29", "Karnataka"),
    ("30", "Goa"), ("31", "Lakshadweep"), ("32", "Kerala"), ("33", "Tamil Nadu"), ("34", "Puducherry"),
    ("35", "Andaman and Nicobar Islands"), ("36", "Telangana"), ("37", "Andhra Pradesh"), ("38", "Ladakh"),
    ("97", "Other Territory"),
]
GST_RATES = "0,0.25,3,5,18,40"
FY_MONTHS = [(2026, m) for m in range(4, 13)] + [(2027, m) for m in range(1, 4)]


# ---------------------------------------------------------------- helpers --

def header(ws, row, titles, widths=None):
    for i, t in enumerate(titles, start=1):
        c = ws.cell(row=row, column=i, value=t)
        c.font = H_FONT
        c.fill = H_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX
    ws.row_dimensions[row].height = 32
    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w


def title(ws, text, sub=None, span=8):
    ws["A1"] = text
    ws["A1"].font = Font(bold=True, size=16, color=NAVY)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=span)
    if sub:
        ws["A2"] = sub
        ws["A2"].font = Font(italic=True, size=10, color="6B7280")
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=span)


def how_to(wb, heading, steps, notes=()):
    ws = wb.active
    ws.title = "How to use"
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 110
    ws["B1"] = heading
    ws["B1"].font = Font(bold=True, size=18, color=NAVY)
    ws["B2"] = FIRM
    ws["B2"].font = Font(size=10, color=GOLD, bold=True)
    r = 4
    ws.cell(row=r, column=2, value="Steps").font = Font(bold=True, size=13, color=NAVY)
    for i, s in enumerate(steps, start=1):
        r += 1
        ws.cell(row=r, column=1, value=i).font = Font(bold=True, color=GOLD)
        c = ws.cell(row=r, column=2, value=s)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    if notes:
        r += 2
        ws.cell(row=r, column=2, value="Good to know").font = Font(bold=True, size=13, color=NAVY)
        for n in notes:
            r += 1
            c = ws.cell(row=r, column=2, value="• " + n)
            c.alignment = Alignment(wrap_text=True, vertical="top")
    r += 2
    c = ws.cell(row=r, column=2, value="Cells shaded light yellow are for your entries. White cells contain formulas — "
                                         "please don't type over them.")
    c.fill = IN_FILL
    c.alignment = Alignment(wrap_text=True)
    r += 2
    ws.cell(row=r, column=2, value="Licensed for use in your own business. Please do not resell or share publicly. "
                                     "This template is a working aid, not professional advice.").font = Font(size=9, color="6B7280")
    return ws


def label_rows(ws, start_row, rows, col_w=(34, 50)):
    ws.column_dimensions["A"].width = col_w[0]
    ws.column_dimensions["B"].width = col_w[1]
    for i, (label, value) in enumerate(rows):
        r = start_row + i
        ws.cell(row=r, column=1, value=label).font = Font(bold=True)
        c = ws.cell(row=r, column=2, value=value)
        c.fill = IN_FILL
        c.border = BOX


def input_block(ws, first_row, last_row, cols):
    for r in range(first_row, last_row + 1):
        for col in cols:
            c = ws.cell(row=r, column=col)
            c.fill = IN_FILL
            c.border = BOX


def states_sheet(wb):
    ws = wb.create_sheet("States")
    header(ws, 1, ["Code", "State / UT", "Code + Name"], [8, 42, 50])
    for i, (code, name) in enumerate(STATES, start=2):
        ws.cell(row=i, column=1, value=code)
        ws.cell(row=i, column=2, value=name)
        ws.cell(row=i, column=3, value=f"{code} - {name}")
    ws.sheet_state = "hidden"
    return f"States!$C$2:$C${len(STATES) + 1}"


# --------------------------------------------------------- GST invoice kit --

def gst_invoice_kit():
    wb = Workbook()
    how_to(
        wb,
        "GST Compliance Kit — Invoice, Registers, 3B Working, ITC Reconciliation & Due Dates",
        [
            "Open the 'Settings' sheet and enter your business name, address, GSTIN, state and bank details once.",
            "Open 'Invoice'. Enter invoice number, date, buyer details and pick the buyer's state (place of supply).",
            "Add items (up to 15 lines): description, HSN/SAC, quantity, rate, discount % and GST rate. Taxable value, "
            "CGST+SGST (same state) or IGST (other state) and totals are calculated automatically.",
            "Type the amount in words, then print (File → Print, A4, 'Fit sheet on one page') or save as PDF.",
            "Copy the invoice summary into 'Sales Register' — one row per invoice. 'Monthly Summary' fills itself "
            "and gives you the numbers for GSTR-1 / GSTR-3B.",
            "For the HSN-wise table in GSTR-1, log item lines in 'HSN Register'; 'HSN Summary' adds them up per HSN code.",
            "For the next invoice: File → Save As with a new name, or clear the yellow cells and change the invoice number.",
            "Record every purchase bill in 'Purchase Register' and mark whether its ITC is eligible.",
            "Each month download GSTR-2B from the GST portal (Returns → GSTR-2B → Download Excel) and paste the B2B "
            "invoice lines into 'GSTR-2B Data'. Both sheets then show, invoice by invoice, what matches and what doesn't.",
            "'ITC Reconciliation' summarises the result — chase suppliers for invoices that are not in your 2B before you "
            "claim that credit.",
            "'GSTR-3B Working' puts output tax and matched ITC side by side for each month and shows the cash you need "
            "to arrange.",
            "'Due Date Calendar' lists every GST due date for FY 2026-27. Enter the date you filed — the status column "
            "turns to 'OVERDUE' on its own when a date is missed.",
            "Filed late? 'Late Fee & Interest' works out the late fee (with the legal cap) and 18% interest.",
        ],
        [
            "Rates in the drop-down follow the GST rate structure in force from 22 September 2025 "
            "(0%, 0.25%, 3%, 5%, 18%, 40%). You can type any other rate if needed.",
            "If the buyer has a GSTIN the sale is B2B, otherwise B2C — the register marks this for you.",
            "Invoice numbers must be consecutive, unique for the financial year and at most 16 characters.",
            "If your turnover is above the e-invoicing limit (currently ₹5 crore), invoices must also be reported on the IRP.",
        ],
    )
    states = states_sheet(wb)

    # Settings
    s = wb.create_sheet("Settings", 1)
    title(s, "Your business details", "Fill once — used on every invoice", 2)
    label_rows(s, 4, [
        ("Business / trade name", "Your Business Name"),
        ("Address line 1", "Shop No., Street"),
        ("Address line 2", "City, State - PIN"),
        ("GSTIN", "06ABCDE1234F1Z5"),
        ("Your state (code - name)", "06 - Haryana"),
        ("Phone", "+91 "),
        ("Email", ""),
        ("Bank name", ""),
        ("Account number", ""),
        ("IFSC", ""),
        ("UPI ID (optional)", ""),
        ("Terms line", "Goods once sold will not be taken back. Subject to local jurisdiction."),
    ])
    dv = DataValidation(type="list", formula1=states, allow_blank=False)
    s.add_data_validation(dv)
    dv.add("B8")

    # Invoice
    ws = wb.create_sheet("Invoice", 2)
    widths = [5, 34, 11, 8, 7, 12, 8, 14, 7, 12, 12, 12, 14]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.merge_cells("A1:M1")
    ws["A1"] = "TAX INVOICE"
    ws["A1"].font = Font(bold=True, size=16, color="FFFFFF")
    ws["A1"].fill = H_FILL
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.merge_cells("A2:G2")
    ws["A2"] = "=Settings!B4"
    ws["A2"].font = Font(bold=True, size=14, color=NAVY)
    ws.merge_cells("A3:G3")
    ws["A3"] = '=Settings!B5&", "&Settings!B6'
    ws.merge_cells("A4:G4")
    ws["A4"] = '="GSTIN: "&Settings!B7&"   |   State: "&Settings!B8'
    ws.merge_cells("A5:G5")
    ws["A5"] = '="Ph: "&Settings!B9&"   "&Settings!B10'

    meta = [("Invoice No.", "INV/26-27/001"), ("Invoice date", None), ("Reverse charge", "No"), ("Supply type", None)]
    for i, (k, v) in enumerate(meta):
        r = 2 + i
        ws.merge_cells(start_row=r, start_column=9, end_row=r, end_column=10)
        ws.cell(row=r, column=9, value=k).font = Font(bold=True)
        ws.merge_cells(start_row=r, start_column=11, end_row=r, end_column=13)
        c = ws.cell(row=r, column=11, value=v)
        c.border = BOX
        if k != "Supply type":
            c.fill = IN_FILL
    ws["K3"].number_format = DATE
    ws["K3"] = "=TODAY()"

    ws["A7"] = "Bill to"
    ws["A7"].font = Font(bold=True, color=NAVY)
    buyer = [("Name", ""), ("Address", ""), ("GSTIN (blank if unregistered)", ""), ("Place of supply", "06 - Haryana")]
    for i, (k, v) in enumerate(buyer):
        r = 8 + i
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        ws.cell(row=r, column=2, value=k).font = Font(bold=True)
        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=10)
        c = ws.cell(row=r, column=4, value=v)
        c.fill = IN_FILL
        c.border = BOX
    dvp = DataValidation(type="list", formula1=states, allow_blank=False)
    ws.add_data_validation(dvp)
    dvp.add("D11")
    # Supply type follows the place of supply (D11) vs your own state
    ws["K5"] = '=IF(D11="","",IF(LEFT(D11,2)=LEFT(Settings!B8,2),"Intra-state (CGST+SGST)","Inter-state (IGST)"))'

    hdr = ["#", "Description of goods / services", "HSN/SAC", "Qty", "Unit", "Rate", "Disc %", "Taxable value",
           "GST %", "CGST", "SGST/UTGST", "IGST", "Total"]
    header(ws, 13, hdr)
    first, last = 14, 28
    dvr = DataValidation(type="list", formula1=f'"{GST_RATES}"', allow_blank=True, showErrorMessage=False)
    ws.add_data_validation(dvr)
    for r in range(first, last + 1):
        ws.cell(row=r, column=1, value=r - first + 1)
        for col in (2, 3, 4, 5, 6, 7, 9):
            c = ws.cell(row=r, column=col)
            c.fill = IN_FILL
            c.border = BOX
        dvr.add(f"I{r}")
        ws.cell(row=r, column=8, value=f'=IF(D{r}="","",ROUND(D{r}*F{r}*(1-G{r}/100),2))')
        intra = 'LEFT($D$11,2)=LEFT(Settings!$B$8,2)'
        ws.cell(row=r, column=10, value=f'=IF(H{r}="","",IF({intra},ROUND(H{r}*I{r}/200,2),0))')
        ws.cell(row=r, column=11, value=f'=IF(H{r}="","",IF({intra},ROUND(H{r}*I{r}/200,2),0))')
        ws.cell(row=r, column=12, value=f'=IF(H{r}="","",IF({intra},0,ROUND(H{r}*I{r}/100,2)))')
        ws.cell(row=r, column=13, value=f'=IF(H{r}="","",H{r}+J{r}+K{r}+L{r})')
        for col in (6, 8, 10, 11, 12, 13):
            ws.cell(row=r, column=col).number_format = MONEY
        for col in (8, 10, 11, 12, 13):
            ws.cell(row=r, column=col).border = BOX
    t = last + 1
    ws.cell(row=t, column=2, value="Total").font = Font(bold=True)
    for col in (4, 8, 10, 11, 12, 13):
        L = get_column_letter(col)
        c = ws.cell(row=t, column=col, value=f"=SUM({L}{first}:{L}{last})")
        c.font = Font(bold=True)
        c.fill = TOT_FILL
        c.number_format = MONEY if col != 4 else "0.##"
        c.border = BOX
    ws.cell(row=t + 1, column=11, value="Round off").font = Font(bold=True)
    ws.merge_cells(start_row=t + 1, start_column=11, end_row=t + 1, end_column=12)
    ws.cell(row=t + 1, column=13, value=f"=ROUND(M{t},0)-M{t}").number_format = MONEY
    ws.cell(row=t + 2, column=11, value="GRAND TOTAL").font = Font(bold=True, color=NAVY)
    ws.merge_cells(start_row=t + 2, start_column=11, end_row=t + 2, end_column=12)
    g = ws.cell(row=t + 2, column=13, value=f"=ROUND(M{t},0)")
    g.font = Font(bold=True, size=13, color=NAVY)
    g.number_format = '"₹" #,##0.00'
    g.fill = TOT_FILL

    ws.cell(row=t + 1, column=1, value="Amount in words:").font = Font(bold=True)
    ws.merge_cells(start_row=t + 1, start_column=3, end_row=t + 1, end_column=9)
    ws.cell(row=t + 1, column=3, value="Rupees ______________ only").fill = IN_FILL
    ws.cell(row=t + 3, column=1, value="Bank details:").font = Font(bold=True)
    ws.merge_cells(start_row=t + 3, start_column=3, end_row=t + 3, end_column=9)
    ws.cell(row=t + 3, column=3, value='=Settings!B11&"  A/c "&Settings!B12&"  IFSC "&Settings!B13&IF(Settings!B14<>"","  UPI "&Settings!B14,"")')
    ws.merge_cells(start_row=t + 4, start_column=1, end_row=t + 4, end_column=9)
    ws.cell(row=t + 4, column=1, value="=Settings!B15").font = Font(size=9, italic=True)
    ws.merge_cells(start_row=t + 4, start_column=10, end_row=t + 4, end_column=13)
    ws.cell(row=t + 4, column=10, value='="For "&Settings!B4').alignment = Alignment(horizontal="right")
    ws.merge_cells(start_row=t + 7, start_column=10, end_row=t + 7, end_column=13)
    ws.cell(row=t + 7, column=10, value="Authorised Signatory").alignment = Alignment(horizontal="right")
    ws.print_area = f"A1:M{t + 7}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    # Sales register
    sr = wb.create_sheet("Sales Register")
    title(sr, "Sales Register", "One row per invoice. Taxable value and tax come from the invoice totals.", 12)
    header(sr, 4, ["Date", "Invoice No.", "Customer", "Customer GSTIN", "Place of supply", "Taxable value",
                   "CGST", "SGST", "IGST", "Invoice total", "B2B / B2C", "Month"],
           [12, 16, 28, 18, 24, 14, 12, 12, 12, 14, 10, 10])
    for r in range(5, 505):
        input_block(sr, r, r, range(1, 10))
        sr.cell(row=r, column=1).number_format = DATE
        sr.cell(row=r, column=10, value=f'=IF(F{r}="","",F{r}+G{r}+H{r}+I{r})')
        sr.cell(row=r, column=11, value=f'=IF(F{r}="","",IF(D{r}<>"","B2B","B2C"))')
        sr.cell(row=r, column=12, value=f'=IF(A{r}="","",TEXT(A{r},"mmm-yy"))')
        for col in range(6, 11):
            sr.cell(row=r, column=col).number_format = MONEY
    sr.freeze_panes = "A5"
    dvs = DataValidation(type="list", formula1=states, allow_blank=True)
    sr.add_data_validation(dvs)
    dvs.add("E5:E504")

    # Monthly summary
    ms = wb.create_sheet("Monthly Summary")
    title(ms, "Monthly GST Summary — FY 2026-27", "Filled automatically from the Sales Register", 9)
    header(ms, 4, ["Month", "Invoices", "Taxable value", "CGST", "SGST", "IGST", "Total tax", "B2B taxable", "B2C taxable"],
           [12, 10, 16, 14, 14, 14, 14, 16, 16])
    rng = "'Sales Register'!$A$5:$A$504"
    for i, (y, m) in enumerate(FY_MONTHS):
        r = 5 + i
        ms.cell(row=r, column=1, value=f"=DATE({y},{m},1)").number_format = "mmm-yy"
        crit = f'{rng},">="&A{r},{rng},"<="&EOMONTH(A{r},0)'
        ms.cell(row=r, column=2, value=f"=COUNTIFS({crit})")
        for col, src in ((3, "F"), (4, "G"), (5, "H"), (6, "I")):
            ms.cell(row=r, column=col, value=f"=SUMIFS('Sales Register'!${src}$5:${src}$504,{crit})")
        ms.cell(row=r, column=7, value=f"=D{r}+E{r}+F{r}")
        ms.cell(row=r, column=8, value=f"=SUMIFS('Sales Register'!$F$5:$F$504,{crit},'Sales Register'!$K$5:$K$504,\"B2B\")")
        ms.cell(row=r, column=9, value=f"=SUMIFS('Sales Register'!$F$5:$F$504,{crit},'Sales Register'!$K$5:$K$504,\"B2C\")")
        for col in range(3, 10):
            ms.cell(row=r, column=col).number_format = MONEY
    tr = 5 + len(FY_MONTHS)
    ms.cell(row=tr, column=1, value="Total").font = Font(bold=True)
    for col in range(2, 10):
        L = get_column_letter(col)
        c = ms.cell(row=tr, column=col, value=f"=SUM({L}5:{L}{tr - 1})")
        c.font = Font(bold=True)
        c.fill = TOT_FILL
        c.number_format = MONEY if col > 2 else "0"

    # HSN register + summary
    hr = wb.create_sheet("HSN Register")
    title(hr, "HSN Register (item lines)", "Log each invoice line here for the GSTR-1 HSN table", 9)
    header(hr, 4, ["Date", "Invoice No.", "HSN/SAC", "Description", "UQC (unit)", "Qty", "Taxable value", "GST %", "Tax amount"],
           [12, 16, 12, 30, 10, 10, 14, 8, 14])
    for r in range(5, 1005):
        input_block(hr, r, r, range(1, 9))
        hr.cell(row=r, column=1).number_format = DATE
        hr.cell(row=r, column=9, value=f'=IF(G{r}="","",ROUND(G{r}*H{r}/100,2))').number_format = MONEY
        hr.cell(row=r, column=7).number_format = MONEY
    hr.freeze_panes = "A5"

    hs = wb.create_sheet("HSN Summary")
    title(hs, "HSN-wise Summary", "Type each HSN code once in column A — totals fill in automatically", 5)
    header(hs, 4, ["HSN/SAC", "Description", "Total qty", "Taxable value", "Tax amount"], [14, 34, 12, 16, 16])
    for r in range(5, 45):
        input_block(hs, r, r, (1, 2))
        hs.cell(row=r, column=3, value=f"=IF(A{r}=\"\",\"\",SUMIF('HSN Register'!$C$5:$C$1004,A{r},'HSN Register'!$F$5:$F$1004))")
        hs.cell(row=r, column=4, value=f"=IF(A{r}=\"\",\"\",SUMIF('HSN Register'!$C$5:$C$1004,A{r},'HSN Register'!$G$5:$G$1004))").number_format = MONEY
        hs.cell(row=r, column=5, value=f"=IF(A{r}=\"\",\"\",SUMIF('HSN Register'!$C$5:$C$1004,A{r},'HSN Register'!$I$5:$I$1004))").number_format = MONEY

    gst_compliance_sheets(wb)

    wb.active = 1
    wb.save(OUT / "GST-Compliance-Kit.xlsx")


# ------------------------------------------- GST compliance (extra sheets) --

# QRMP GSTR-3B is due on the 22nd for these state codes, on the 24th for the rest.
QRMP_22_STATES = '{"22","23","24","25","26","27","29","30","31","32","33","34","35","36","37"}'


def money_cols(ws, row, cols):
    for col in cols:
        ws.cell(row=row, column=col).number_format = MONEY


def gst_compliance_sheets(wb):
    """Purchase register, GSTR-2B matching, 3B working, due dates and late fee — added to the invoice kit."""
    n = 500  # rows 5..504

    # Purchase register with a match key against GSTR-2B
    pr = wb.create_sheet("Purchase Register")
    title(pr, "Purchase Register (inward supplies)",
          "One row per purchase bill. The last column checks the bill against your GSTR-2B data.", 12)
    header(pr, 4, ["Date", "Supplier", "Supplier GSTIN", "Bill / Invoice No.", "Taxable value", "CGST", "SGST", "IGST",
                   "Total", "ITC eligible? (Y/N)", "Match key", "Status vs GSTR-2B"],
           [12, 28, 18, 16, 14, 12, 12, 12, 14, 12, 26, 26])
    yn = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
    pr.add_data_validation(yn)
    yn.add(f"J5:J{n + 4}")
    for r in range(5, n + 5):
        input_block(pr, r, r, list(range(1, 9)) + [10])
        pr.cell(row=r, column=1).number_format = DATE
        pr.cell(row=r, column=9, value=f'=IF(E{r}="","",E{r}+F{r}+G{r}+H{r})')
        money_cols(pr, r, range(5, 10))
        pr.cell(row=r, column=11, value=f'=IF(D{r}="","",UPPER(TRIM(C{r}))&"|"&UPPER(TRIM(D{r})))')
        pr.cell(row=r, column=12, value=(
            f'=IF(K{r}="","",IF(COUNTIF(\'GSTR-2B Data\'!$J$5:$J${n + 4},K{r})=0,"Not in 2B — follow up",'
            f'IF(ABS(SUMIF(\'GSTR-2B Data\'!$J$5:$J${n + 4},K{r},\'GSTR-2B Data\'!$I$5:$I${n + 4})-(F{r}+G{r}+H{r}))>1,'
            f'"Tax amount differs","Matched")))'))
    pr.freeze_panes = "A5"

    # GSTR-2B data pasted from the portal
    g2 = wb.create_sheet("GSTR-2B Data")
    title(g2, "GSTR-2B Data (paste from the portal)",
          "GST portal → Returns → GSTR-2B → Download (Excel) → 'B2B' tab. Paste supplier GSTIN, name, invoice no., "
          "date, taxable value and tax into the yellow columns.", 11)
    header(g2, 4, ["Supplier GSTIN", "Supplier name", "Invoice No.", "Invoice date", "Taxable value", "CGST", "SGST",
                   "IGST", "Total tax", "Match key", "In your books?"],
           [18, 28, 16, 12, 14, 12, 12, 12, 14, 26, 26])
    for r in range(5, n + 5):
        input_block(g2, r, r, range(1, 9))
        g2.cell(row=r, column=4).number_format = DATE
        g2.cell(row=r, column=9, value=f'=IF(C{r}="","",F{r}+G{r}+H{r})')
        money_cols(g2, r, range(5, 10))
        g2.cell(row=r, column=10, value=f'=IF(C{r}="","",UPPER(TRIM(A{r}))&"|"&UPPER(TRIM(C{r})))')
        g2.cell(row=r, column=11, value=(
            f'=IF(J{r}="","",IF(COUNTIF(\'Purchase Register\'!$K$5:$K${n + 4},J{r})=0,'
            f'"In 2B, not in your books","In books"))'))
    g2.freeze_panes = "A5"

    # Reconciliation summary
    rc = wb.create_sheet("ITC Reconciliation")
    title(rc, "ITC Reconciliation — Books vs GSTR-2B", "Fills automatically from the two sheets", 4)
    header(rc, 4, ["Result", "No. of invoices", "Tax amount", "What to do"], [32, 16, 18, 70])
    prs = f"'Purchase Register'!$L$5:$L${n + 4}"
    prt = f"'Purchase Register'!$F$5:$F${n + 4}+'Purchase Register'!$G$5:$G${n + 4}+'Purchase Register'!$H$5:$H${n + 4}"
    rows = [
        ("Matched", f'=COUNTIF({prs},"Matched")', f'=SUMPRODUCT(({prs}="Matched")*({prt}))',
         "Safe to claim (if otherwise eligible)."),
        ("Not in 2B — follow up", f'=COUNTIF({prs},"Not in 2B — follow up")',
         f'=SUMPRODUCT(({prs}="Not in 2B — follow up")*({prt}))',
         "Ask the supplier to upload the invoice in their GSTR-1/IFF. Claim the credit only when it appears in your 2B."),
        ("Tax amount differs", f'=COUNTIF({prs},"Tax amount differs")',
         f'=SUMPRODUCT(({prs}="Tax amount differs")*({prt}))',
         "Check the invoice — wrong rate, typing error, or the supplier needs to amend."),
        ("In 2B, not in your books", f"=COUNTIF('GSTR-2B Data'!$K$5:$K${n + 4},\"In 2B, not in your books\")",
         f"=SUMIF('GSTR-2B Data'!$K$5:$K${n + 4},\"In 2B, not in your books\",'GSTR-2B Data'!$I$5:$I${n + 4})",
         "Either a bill you forgot to record, or a supplier reported a sale to you by mistake — check before claiming."),
    ]
    for i, (lab, cnt, amt, todo) in enumerate(rows):
        r = 5 + i
        rc.cell(row=r, column=1, value=lab).font = Font(bold=True)
        rc.cell(row=r, column=2, value=cnt)
        rc.cell(row=r, column=3, value=amt).number_format = MONEY
        c = rc.cell(row=r, column=4, value=todo)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        rc.row_dimensions[r].height = 32
    rc["A11"] = ("Matching is on supplier GSTIN + invoice number (spaces and capital letters ignored). If a supplier "
                 "writes the invoice number differently, correct it in one of the sheets so the two match.")
    rc["A11"].alignment = Alignment(wrap_text=True)
    rc.merge_cells("A11:D12")
    rc.row_dimensions[11].height = 30

    # GSTR-3B working
    w3 = wb.create_sheet("GSTR-3B Working")
    title(w3, "GSTR-3B Working — FY 2026-27",
          "Output tax from 'Monthly Summary'; ITC = eligible purchase bills that are matched with GSTR-2B "
          "(by bill date). Use the adjustment column for anything else.", 11)
    header(w3, 4, ["Month", "Output CGST", "Output SGST", "Output IGST", "Total output tax", "Matched ITC (books)",
                   "ITC adjustment (+/−)", "ITC available incl. brought forward", "Cash to pay (approx.)",
                   "ITC carried forward", "Note"],
           [12, 14, 14, 14, 16, 16, 16, 20, 16, 16, 36])
    pdt = f"'Purchase Register'!$A$5:$A${n + 4}"
    for i, (y, m) in enumerate(FY_MONTHS):
        r = 5 + i
        ms_row = 5 + i
        w3.cell(row=r, column=1, value=f"=DATE({y},{m},1)").number_format = "mmm-yy"
        w3.cell(row=r, column=2, value=f"='Monthly Summary'!D{ms_row}")
        w3.cell(row=r, column=3, value=f"='Monthly Summary'!E{ms_row}")
        w3.cell(row=r, column=4, value=f"='Monthly Summary'!F{ms_row}")
        w3.cell(row=r, column=5, value=f"=B{r}+C{r}+D{r}")
        w3.cell(row=r, column=6, value=(
            f"=SUMPRODUCT(({pdt}>=A{r})*({pdt}<=EOMONTH(A{r},0))*('Purchase Register'!$J$5:$J${n + 4}=\"Y\")"
            f"*({prs}=\"Matched\")*({prt}))"))
        input_block(w3, r, r, (7,))
        prev = f"+J{r - 1}" if i else ""
        w3.cell(row=r, column=8, value=f"=F{r}+G{r}{prev}")
        w3.cell(row=r, column=9, value=f"=MAX(0,E{r}-H{r})")
        w3.cell(row=r, column=10, value=f"=MAX(0,H{r}-E{r})")
        money_cols(w3, r, range(2, 11))
    tr = 5 + len(FY_MONTHS)
    w3.cell(row=tr, column=1, value="Total").font = Font(bold=True)
    for col in (2, 3, 4, 5, 6, 7, 9):
        L = get_column_letter(col)
        c = w3.cell(row=tr, column=col, value=f"=SUM({L}5:{L}{tr - 1})")
        c.font = Font(bold=True)
        c.fill = TOT_FILL
        c.number_format = MONEY
    w3.cell(row=tr + 2, column=1, value=(
        "This is a planning figure. The portal sets off credit head-wise in a fixed order (IGST first, then CGST/SGST) — "
        "the exact cash per head is shown in GSTR-3B Table 6.1. Blocked credits (Section 17(5)) must be marked 'N' "
        "in the Purchase Register.")).alignment = Alignment(wrap_text=True)
    w3.merge_cells(start_row=tr + 2, start_column=1, end_row=tr + 3, end_column=11)
    w3.row_dimensions[tr + 2].height = 30
    w3.freeze_panes = "B5"

    # Due date calendar
    dd = wb.create_sheet("Due Date Calendar")
    title(dd, "GST Due Date Calendar — FY 2026-27",
          "Statutory dates. The government sometimes extends a date — if so, just change it in column C. "
          "Enter the date you filed in column E.", 6)
    header(dd, 4, ["Return / payment", "Period", "Due date", "Who files it", "Filed on", "Status"],
           [30, 16, 13, 34, 13, 28])
    items = []
    month_names = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]

    def nxt(y, m):
        return (y + 1, 1) if m == 12 else (y, m + 1)

    for y, m in FY_MONTHS:
        ny, nm = nxt(y, m)
        per = f"{month_names[m - 1]}-{str(y)[2:]}"
        items.append(("GSTR-1", per, f"=DATE({ny},{nm},11)", "Monthly filers"))
        items.append(("GSTR-3B + tax payment", per, f"=DATE({ny},{nm},20)", "Monthly filers"))
        if m % 3 != 0:  # first two months of a quarter
            items.append(("IFF (optional B2B invoices)", per, f"=DATE({ny},{nm},13)", "QRMP (quarterly) filers"))
            items.append(("PMT-06 tax payment", per, f"=DATE({ny},{nm},25)", "QRMP (quarterly) filers"))
        else:
            q = f"{month_names[m - 3]}–{month_names[m - 1]}-{str(y)[2:]}"
            items.append(("GSTR-1 (quarterly)", q, f"=DATE({ny},{nm},13)", "QRMP (quarterly) filers"))
            items.append(("GSTR-3B (quarterly) + tax", q,
                          f"=DATE({ny},{nm},IF(ISNUMBER(MATCH(LEFT(Settings!$B$8,2),{QRMP_22_STATES},0)),22,24))",
                          "QRMP filers (22nd or 24th by your state — from Settings)"))
            items.append(("CMP-08 (composition tax)", q, f"=DATE({ny},{nm},18)", "Composition dealers"))
    items.append(("ITC last date for FY 2025-26 credits", "FY 2025-26", "=DATE(2026,11,30)",
                  "All regular taxpayers (claim in a return filed by this date)"))
    items.append(("GSTR-9 / 9C annual return", "FY 2025-26", "=DATE(2026,12,31)",
                  "Turnover above ₹2 crore (9C above ₹5 crore)"))
    items.append(("GSTR-4 annual return", "FY 2026-27", "=DATE(2027,4,30)", "Composition dealers"))
    for i, (ret, per, due, who) in enumerate(items):
        r = 5 + i
        dd.cell(row=r, column=1, value=ret)
        dd.cell(row=r, column=2, value=per)
        dd.cell(row=r, column=3, value=due).number_format = DATE
        dd.cell(row=r, column=4, value=who)
        input_block(dd, r, r, (5,))
        dd.cell(row=r, column=5).number_format = DATE
        dd.cell(row=r, column=6, value=(
            f'=IF(E{r}<>"",IF(E{r}<=C{r},"Filed on time","Filed late — see Late Fee sheet"),'
            f'IF(TODAY()>C{r},"OVERDUE",IF(C{r}-TODAY()<=5,"Due in "&(C{r}-TODAY())&" day(s)","Upcoming")))'))
    dd.freeze_panes = "A5"
    dd.auto_filter.ref = f"A4:F{4 + len(items)}"

    # Late fee & interest
    lf = wb.create_sheet("Late Fee & Interest")
    title(lf, "Late Fee & Interest Calculator — GSTR-1 / GSTR-3B", "Fill the yellow cells", 3)
    lf.column_dimensions["A"].width = 44
    lf.column_dimensions["B"].width = 22
    lf.column_dimensions["C"].width = 60
    ins = [
        ("Return", "GSTR-3B"),
        ("Nil return? (Y/N)", "N"),
        ("Aggregate turnover in the previous FY", "Up to ₹1.5 crore"),
        ("Due date", None),
        ("Date actually filed (or expected)", None),
        ("Tax paid late in cash (net of ITC), ₹", 0),
    ]
    for i, (k, v) in enumerate(ins):
        r = 4 + i
        lf.cell(row=r, column=1, value=k).font = Font(bold=True)
        c = lf.cell(row=r, column=2, value=v)
        c.fill = IN_FILL
        c.border = BOX
    lf["B7"].number_format = DATE
    lf["B8"].number_format = DATE
    lf["B9"].number_format = MONEY
    for ref, f in (("B4", '"GSTR-3B,GSTR-1"'), ("B5", '"Y,N"'),
                   ("B6", '"Up to ₹1.5 crore,₹1.5 crore to ₹5 crore,Above ₹5 crore"')):
        dv = DataValidation(type="list", formula1=f, allow_blank=False)
        lf.add_data_validation(dv)
        dv.add(ref)
    outs = [
        ("Days late", '=IF(OR(B7="",B8=""),0,MAX(0,B8-B7))', "0"),
        ("Late fee per day (CGST + SGST)", '=IF(B5="Y",20,50)', MONEY),
        ("Maximum late fee for this return", '=IF(B5="Y",500,IF(B6="Up to ₹1.5 crore",2000,IF(B6="Above ₹5 crore",10000,5000)))', MONEY),
        ("Late fee payable", "=MIN(B11*B12,B13)", MONEY),
        ("   of which CGST", "=B14/2", MONEY),
        ("   of which SGST", "=B14/2", MONEY),
        ("Interest @ 18% p.a. on tax paid late", '=IF(B4="GSTR-3B",ROUND(B9*18%*B11/365,0),0)', MONEY),
        ("TOTAL (late fee + interest)", "=B14+B17", MONEY),
    ]
    for i, (k, f, fmt) in enumerate(outs):
        r = 11 + i
        lf.cell(row=r, column=1, value=k).font = Font(bold=(k.startswith("TOTAL") or k == "Late fee payable"))
        c = lf.cell(row=r, column=2, value=f)
        c.number_format = fmt
        c.fill = TOT_FILL
    lf["B18"].font = Font(bold=True, color=NAVY, size=12)
    notes = [
        "Late fee: ₹50 per day (₹25 CGST + ₹25 SGST); ₹20 per day (₹10 + ₹10) for nil returns.",
        "Maximum late fee per return (Notifications 19/2021 & 20/2021 — CT): nil return ₹500; turnover up to "
        "₹1.5 crore ₹2,000; ₹1.5–5 crore ₹5,000; above ₹5 crore ₹10,000.",
        "Interest (Section 50): 18% per year on the tax paid in cash after the due date, counted from the day after "
        "the due date. GSTR-1 itself carries no interest.",
        "Late fee is paid in cash only — it cannot be paid from ITC.",
        "Rates as of FY 2026-27. Waivers/amnesty notified by the government can change the amount — check before paying.",
    ]
    for i, t in enumerate(notes):
        c = lf.cell(row=11 + i, column=3, value="• " + t)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        lf.row_dimensions[11 + i].height = 30


# ---------------------------------------------------------- Bookkeeping kit --

EXPENSE_HEADS = [
    "Rent", "Salary & Wages", "Electricity", "Telephone & Internet", "Travel & Conveyance", "Office Expenses",
    "Printing & Stationery", "Repairs & Maintenance", "Professional Fees", "Bank Charges", "Advertising",
    "Freight & Cartage", "Insurance", "Interest", "Miscellaneous",
]


def bookkeeping_kit():
    wb = Workbook()
    how_to(
        wb,
        "Small Business Accounting Kit — Books, P&L, Receivables & Balance Sheet",
        [
            "In 'Settings' enter your business name and the opening cash and bank balance on 1 April 2026.",
            "Record every cash receipt/payment in 'Cash Book' and every bank entry in 'Bank Book' — the balance runs automatically.",
            "Record each sale bill in 'Sales Register' and each purchase bill in 'Purchase Register'.",
            "Record expenses (rent, salary, electricity...) in 'Expenses' and choose the head from the drop-down.",
            "In the Purchase Register mark whether the bill appears in your GSTR-2B — 'ITC Tracker' then shows the month-wise gap.",
            "Open 'P&L Summary' any time to see month-wise sales, purchases, expenses and profit.",
            "Enter payments received against each sale (and paid against each purchase). 'Receivables' and 'Payables' "
            "then show who owes what, with an ageing of overdue bills — type each customer/supplier name once there.",
            "'Bank Reconciliation' explains the difference between your Bank Book and the bank statement.",
            "At year end fill the few yellow cells in 'Balance Sheet' (capital, loans, fixed assets, stock) — cash, bank, "
            "debtors, creditors and profit come in automatically, and a check line tells you if it balances.",
            "'Dashboard' shows the year so far on one screen.",
            "Share this file with your CA at year end — everything is already classified.",
        ],
        [
            "If a sale/purchase is paid in cash or bank, also enter it in the Cash/Bank Book so balances stay correct.",
            "Profit shown is an estimate (no stock valuation or depreciation) — your CA will finalise it.",
            "Cash expenses above ₹10,000 paid to one person in a day are not allowed as a deduction (Income Tax rule) — pay by bank.",
        ],
    )
    s = wb.create_sheet("Settings", 1)
    title(s, "Settings", None, 2)
    label_rows(s, 3, [
        ("Business name", "Your Business Name"),
        ("Financial year", "FY 2026-27"),
        ("Opening cash balance (1 Apr 2026)", 0),
        ("Opening bank balance (1 Apr 2026)", 0),
    ])
    heads = wb.create_sheet("Heads")
    heads["A1"] = "Expense heads (edit if you like)"
    heads["A1"].font = Font(bold=True)
    heads.column_dimensions["A"].width = 30
    for i, h in enumerate(EXPENSE_HEADS, start=2):
        heads.cell(row=i, column=1, value=h).fill = IN_FILL
    head_rng = f"Heads!$A$2:$A${len(EXPENSE_HEADS) + 1}"

    def book(name, opening_ref, extra_col=None):
        ws = wb.create_sheet(name)
        title(ws, name, "Balance is calculated automatically", 7)
        cols = ["Date", "Particulars", "Voucher / Ref"] + ([extra_col] if extra_col else []) + ["Receipt (In)", "Payment (Out)", "Balance"]
        header(ws, 4, cols, [12, 40, 14] + ([18] if extra_col else []) + [14, 14, 16])
        n_in = len(cols) - 3
        rc, pc, bc = n_in + 1, n_in + 2, n_in + 3
        R, P, B = get_column_letter(rc), get_column_letter(pc), get_column_letter(bc)
        ws.cell(row=5, column=2, value="Opening balance").font = Font(bold=True)
        ws.cell(row=5, column=bc, value=f"={opening_ref}").number_format = MONEY
        for r in range(6, 1006):
            input_block(ws, r, r, range(1, pc + 1))
            ws.cell(row=r, column=1).number_format = DATE
            ws.cell(row=r, column=rc).number_format = MONEY
            ws.cell(row=r, column=pc).number_format = MONEY
            ws.cell(row=r, column=bc,
                    value=f'=IF(AND({R}{r}="",{P}{r}=""),"",{opening_ref}+SUM({R}$6:{R}{r})-SUM({P}$6:{P}{r}))').number_format = MONEY
        ws.freeze_panes = "A6"

    book("Cash Book", "Settings!B5")
    book("Bank Book", "Settings!B6", "Cheque / UTR")

    sr = wb.create_sheet("Sales Register")
    title(sr, "Sales Register", "Enter what the customer has paid so far — 'Receivables' tracks who owes you and since when", 13)
    header(sr, 4, ["Date", "Invoice No.", "Customer", "GSTIN", "Taxable value", "CGST", "SGST", "IGST", "Total",
                   "Amount received", "Credit days", "Balance due", "Days overdue"],
           [12, 16, 30, 18, 14, 12, 12, 12, 14, 14, 10, 14, 12])
    for r in range(5, 1005):
        input_block(sr, r, r, list(range(1, 9)) + [10, 11])
        sr.cell(row=r, column=1).number_format = DATE
        sr.cell(row=r, column=9, value=f'=IF(E{r}="","",E{r}+F{r}+G{r}+H{r})')
        sr.cell(row=r, column=12, value=f'=IF(I{r}="","",I{r}-N(J{r}))')
        sr.cell(row=r, column=13, value=f'=IF(OR(L{r}="",N(L{r})<=0,A{r}=""),"",MAX(0,TODAY()-(A{r}+N(K{r}))))')
        for col in range(5, 11):
            sr.cell(row=r, column=col).number_format = MONEY
        sr.cell(row=r, column=12).number_format = MONEY
    sr.freeze_panes = "A5"

    pr = wb.create_sheet("Purchase Register")
    title(pr, "Purchase Register", "Mark ITC eligibility and whether the bill shows in your GSTR-2B", 11)
    header(pr, 4, ["Date", "Bill No.", "Supplier", "Supplier GSTIN", "Taxable value", "CGST", "SGST", "IGST", "Total",
                   "ITC eligible? (Y/N)", "In GSTR-2B? (Y/N)", "Amount paid", "Balance payable"],
           [12, 16, 30, 18, 14, 12, 12, 12, 14, 12, 12, 14, 14])
    yn = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
    pr.add_data_validation(yn)
    yn.add("J5:K1004")
    for r in range(5, 1005):
        input_block(pr, r, r, list(range(1, 9)) + [10, 11, 12])
        pr.cell(row=r, column=13, value=f'=IF(I{r}="","",I{r}-N(L{r}))')
        pr.cell(row=r, column=12).number_format = MONEY
        pr.cell(row=r, column=13).number_format = MONEY
        pr.cell(row=r, column=1).number_format = DATE
        pr.cell(row=r, column=9, value=f'=IF(E{r}="","",E{r}+F{r}+G{r}+H{r})')
        for col in range(5, 10):
            pr.cell(row=r, column=col).number_format = MONEY
    pr.freeze_panes = "A5"

    ex = wb.create_sheet("Expenses")
    title(ex, "Expenses", "Pick the head from the drop-down", 6)
    header(ex, 4, ["Date", "Expense head", "Paid to", "Amount", "Mode (Cash/Bank/UPI)", "Remarks"], [12, 24, 28, 14, 16, 30])
    dvh = DataValidation(type="list", formula1=head_rng, allow_blank=True)
    ex.add_data_validation(dvh)
    dvh.add("B5:B1004")
    mode = DataValidation(type="list", formula1='"Cash,Bank,UPI,Card"', allow_blank=True)
    ex.add_data_validation(mode)
    mode.add("E5:E1004")
    for r in range(5, 1005):
        input_block(ex, r, r, range(1, 7))
        ex.cell(row=r, column=1).number_format = DATE
        ex.cell(row=r, column=4).number_format = MONEY
    ex.freeze_panes = "A5"

    it = wb.create_sheet("ITC Tracker")
    title(it, "ITC Tracker — books vs GSTR-2B", "Difference = ITC you booked that is not yet in GSTR-2B (follow up with the supplier)", 5)
    header(it, 4, ["Month", "Eligible ITC in books", "ITC reflected in GSTR-2B", "Difference", "Status"], [12, 20, 22, 16, 30])
    pd = "'Purchase Register'!$A$5:$A$1004"
    tax = "('Purchase Register'!$F$5:$F$1004+'Purchase Register'!$G$5:$G$1004+'Purchase Register'!$H$5:$H$1004)"
    for i, (y, m) in enumerate(FY_MONTHS):
        r = 5 + i
        it.cell(row=r, column=1, value=f"=DATE({y},{m},1)").number_format = "mmm-yy"
        inmonth = f"({pd}>=A{r})*({pd}<=EOMONTH(A{r},0))"
        it.cell(row=r, column=2, value=f"=SUMPRODUCT({inmonth}*('Purchase Register'!$J$5:$J$1004=\"Y\")*{tax})").number_format = MONEY
        it.cell(row=r, column=3, value=f"=SUMPRODUCT({inmonth}*('Purchase Register'!$J$5:$J$1004=\"Y\")*('Purchase Register'!$K$5:$K$1004=\"Y\")*{tax})").number_format = MONEY
        it.cell(row=r, column=4, value=f"=B{r}-C{r}").number_format = MONEY
        it.cell(row=r, column=5, value=f'=IF(B{r}=0,"",IF(D{r}=0,"All matched","Follow up with suppliers"))')

    pl = wb.create_sheet("P&L Summary")
    title(pl, "Profit & Loss Summary — FY 2026-27", "Estimated, from your registers (taxable values, excluding GST)", 14)
    pl.column_dimensions["A"].width = 26
    pl.cell(row=4, column=1, value="Particulars")
    for i, (y, m) in enumerate(FY_MONTHS):
        c = pl.cell(row=4, column=2 + i, value=f"=DATE({y},{m},1)")
        c.number_format = "mmm-yy"
        pl.column_dimensions[get_column_letter(2 + i)].width = 12
    pl.cell(row=4, column=14, value="Total")
    pl.column_dimensions["N"].width = 14
    for col in range(1, 15):
        c = pl.cell(row=4, column=col)
        c.font = H_FONT
        c.fill = H_FILL
        c.alignment = Alignment(horizontal="center")

    def month_sum(sheet, amount_col, r_col, extra=""):
        rng = f"'{sheet}'!$A$5:$A$1004"
        return (f"=SUMIFS('{sheet}'!${amount_col}$5:${amount_col}$1004,{rng},\">=\"&{r_col}$4,"
                f"{rng},\"<=\"&EOMONTH({r_col}$4,0){extra})")

    rows = [("Sales", "Sales Register", "E", "")]
    rows.append(("Purchases", "Purchase Register", "E", ""))
    r = 5
    for label, sheet, col, extra in rows:
        pl.cell(row=r, column=1, value=label).font = Font(bold=True)
        for i in range(12):
            L = get_column_letter(2 + i)
            pl.cell(row=r, column=2 + i, value=month_sum(sheet, col, L, extra)).number_format = MONEY
        r += 1
    pl.cell(row=r, column=1, value="Gross profit").font = Font(bold=True, color=NAVY)
    for i in range(12):
        L = get_column_letter(2 + i)
        c = pl.cell(row=r, column=2 + i, value=f"={L}5-{L}6")
        c.number_format = MONEY
        c.fill = TOT_FILL
    gp_row = r
    r += 2
    first_exp = r
    for h_i, h in enumerate(EXPENSE_HEADS):
        pl.cell(row=r, column=1, value=f"=Heads!A{h_i + 2}")
        for i in range(12):
            L = get_column_letter(2 + i)
            pl.cell(row=r, column=2 + i, value=month_sum("Expenses", "D", L, f",'Expenses'!$B$5:$B$1004,$A{r}")).number_format = MONEY
        r += 1
    pl.cell(row=r, column=1, value="Total expenses").font = Font(bold=True)
    for i in range(12):
        L = get_column_letter(2 + i)
        c = pl.cell(row=r, column=2 + i, value=f"=SUM({L}{first_exp}:{L}{r - 1})")
        c.number_format = MONEY
        c.fill = TOT_FILL
    te_row = r
    r += 2
    pl.cell(row=r, column=1, value="NET PROFIT").font = Font(bold=True, size=12, color=NAVY)
    for i in range(12):
        L = get_column_letter(2 + i)
        c = pl.cell(row=r, column=2 + i, value=f"={L}{gp_row}-{L}{te_row}")
        c.number_format = MONEY
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E7F5EC")
    for rr in range(5, r + 1):
        if pl.cell(row=rr, column=2).value is not None:
            c = pl.cell(row=rr, column=14, value=f"=SUM(B{rr}:M{rr})")
            c.number_format = MONEY
            c.font = Font(bold=True)
    pl.freeze_panes = "B5"

    accounting_extras(wb, r)

    wb.active = 1
    wb.save(OUT / "Small-Business-Accounting-Kit.xlsx")


def accounting_extras(wb, np_row):
    """Receivables, payables, bank reconciliation, balance sheet and dashboard for the accounting kit."""
    S = "'Sales Register'"
    P = "'Purchase Register'"

    def party_sheet(name, sub, src, party_col, total_col, paid_col, bal_col, ageing):
        ws = wb.create_sheet(name)
        title(ws, name, sub, 8 if ageing else 5)
        cols = ["Name (type once)", "Total billed", "Paid / received", "Outstanding"]
        if ageing:
            cols += ["Not due / 0–30 days", "31–60 days", "61–90 days", "Over 90 days"]
        header(ws, 4, cols, [30, 16, 16, 16, 18, 14, 14, 14])
        rng = lambda c: f"{src}!${c}$5:${c}$1004"  # noqa: E731
        for r in range(5, 105):
            input_block(ws, r, r, (1,))
            ws.cell(row=r, column=2, value=f'=IF(A{r}="","",SUMIF({rng(party_col)},A{r},{rng(total_col)}))')
            ws.cell(row=r, column=3, value=f'=IF(A{r}="","",SUMIF({rng(party_col)},A{r},{rng(paid_col)}))')
            ws.cell(row=r, column=4, value=f'=IF(A{r}="","",B{r}-C{r})')
            if ageing:
                base = f"{rng(bal_col)},{rng(party_col)},A{r},{rng('M')}"
                ws.cell(row=r, column=5, value=f'=IF(A{r}="","",SUMIFS({base},"<=30"))')
                ws.cell(row=r, column=6, value=f'=IF(A{r}="","",SUMIFS({base},">30",{rng("M")},"<=60"))')
                ws.cell(row=r, column=7, value=f'=IF(A{r}="","",SUMIFS({base},">60",{rng("M")},"<=90"))')
                ws.cell(row=r, column=8, value=f'=IF(A{r}="","",SUMIFS({base},">90"))')
            money_cols(ws, r, range(2, 9 if ageing else 5))
        last = 8 if ageing else 4
        ws.cell(row=105, column=1, value="Total (all parties in the register)").font = Font(bold=True)
        ws.cell(row=105, column=2, value=f"=SUM({rng(total_col)})")
        ws.cell(row=105, column=3, value=f"=SUM({rng(paid_col)})")
        ws.cell(row=105, column=4, value=f"=SUM({rng(bal_col)})")
        for col in range(2, 5):
            c = ws.cell(row=105, column=col)
            c.font = Font(bold=True)
            c.fill = TOT_FILL
            c.number_format = MONEY
        ws.freeze_panes = "A5"
        return last

    party_sheet("Receivables", "Who owes you — from the Sales Register. Ageing counts days after the credit period.",
                S, "C", "I", "J", "L", True)
    party_sheet("Payables", "What you owe suppliers — from the Purchase Register.", P, "C", "I", "L", "M", False)

    # Bank reconciliation
    br = wb.create_sheet("Bank Reconciliation")
    title(br, "Bank Reconciliation Statement", "Compare your Bank Book with the bank statement on any date", 4)
    br.column_dimensions["A"].width = 52
    br.column_dimensions["B"].width = 18
    br.column_dimensions["C"].width = 18
    br.column_dimensions["D"].width = 40
    br["A4"] = "Balance as per your Bank Book (latest)"
    br["B4"] = "=Settings!B6+SUM('Bank Book'!E6:E1005)-SUM('Bank Book'!F6:F1005)"
    br["A5"] = "Add: cheques issued but not yet presented (list below)"
    br["B5"] = "=SUMIF(A12:A41,\"Cheque issued, not presented\",C12:C41)"
    br["A6"] = "Less: cheques / cash deposited but not yet credited"
    br["B6"] = "=SUMIF(A12:A41,\"Deposit not yet credited\",C12:C41)"
    br["A7"] = "Less: bank charges / debits not yet entered in Bank Book"
    br["B7"] = "=SUMIF(A12:A41,\"Bank debit not in books\",C12:C41)"
    br["A8"] = "Add: interest / credits not yet entered in Bank Book"
    br["B8"] = "=SUMIF(A12:A41,\"Bank credit not in books\",C12:C41)"
    br["A9"] = "Balance as per bank statement (expected)"
    br["B9"] = "=B4+B5-B6-B7+B8"
    br["C9"] = "Actual statement balance →"
    br["D9"] = 0
    br["D9"].fill = IN_FILL
    br["A10"] = "Unexplained difference"
    br["B10"] = '=D9-B9'
    br["C10"] = '=IF(ABS(B10)<1,"Reconciled ✔","Find the difference")'
    for ref in ("A4", "A9", "A10"):
        br[ref].font = Font(bold=True)
    for ref in ("B4", "B5", "B6", "B7", "B8", "B9", "B10", "D9"):
        br[ref].number_format = MONEY
    header(br, 11, ["Item type", "Date", "Amount", "Details (cheque no., party)"])
    kinds = DataValidation(type="list", formula1='"Cheque issued, not presented,Deposit not yet credited,'
                                                  'Bank debit not in books,Bank credit not in books"', allow_blank=True)
    br.add_data_validation(kinds)
    kinds.add("A12:A41")
    for r in range(12, 42):
        input_block(br, r, r, range(1, 5))
        br.cell(row=r, column=2).number_format = DATE
        br.cell(row=r, column=3).number_format = MONEY

    # Balance sheet
    bs = wb.create_sheet("Balance Sheet")
    title(bs, "Simple Balance Sheet (Statement of Affairs) — as on 31 March 2027",
          "Yellow = your entry. Everything else comes from the other sheets.", 4)
    bs.column_dimensions["A"].width = 46
    bs.column_dimensions["B"].width = 18
    bs.column_dimensions["C"].width = 46
    bs.column_dimensions["D"].width = 18
    header(bs, 4, ["Capital & Liabilities", "₹", "Assets", "₹"])
    left = [
        ("Opening capital (1 April 2026)", 0, True),
        ("Add: net profit from P&L Summary", f"='P&L Summary'!N{np_row}", False),
        ("Add: closing stock − opening stock (stock adjustment)", "=D12-B21", False),
        ("Less: drawings (personal withdrawals)", 0, True),
        ("Closing capital", "=B5+B6+B7-B8", False),
        ("Loans (bank / others)", 0, True),
        ("Sundry creditors (from Purchase Register)", f"=SUM({P}!M5:M1004)", False),
        ("GST / TDS payable", 0, True),
        ("Other liabilities", 0, True),
    ]
    right = [
        ("Fixed assets (machinery, furniture, vehicle…)", 0, True),
        ("Sundry debtors (from Sales Register)", f"=SUM({S}!L5:L1004)", False),
        ("Cash in hand (Cash Book)", "=Settings!B5+SUM('Cash Book'!D6:D1005)-SUM('Cash Book'!E6:E1005)", False),
        ("Cash at bank (Bank Book)", "=Settings!B6+SUM('Bank Book'!E6:E1005)-SUM('Bank Book'!F6:F1005)", False),
        ("Investments / deposits", 0, True),
        ("GST input credit balance", 0, True),
        ("Other assets / advances", 0, True),
        ("Closing stock (value at cost)", 0, True),
    ]
    for i, (lab, val, inp) in enumerate(left):
        r = 5 + i
        bs.cell(row=r, column=1, value=lab).font = Font(bold=(lab == "Closing capital"))
        c = bs.cell(row=r, column=2, value=val)
        c.number_format = MONEY
        if inp:
            c.fill = IN_FILL
            c.border = BOX
    for i, (lab, val, inp) in enumerate(right):
        r = 5 + i
        bs.cell(row=r, column=3, value=lab)
        c = bs.cell(row=r, column=4, value=val)
        c.number_format = MONEY
        if inp:
            c.fill = IN_FILL
            c.border = BOX
    bs["A15"] = "TOTAL"
    bs["B15"] = "=B9+B10+B11+B12+B13"
    bs["C15"] = "TOTAL"
    bs["D15"] = "=SUM(D5:D12)"
    for ref in ("A15", "B15", "C15", "D15"):
        bs[ref].font = Font(bold=True, color=NAVY)
        bs[ref].fill = TOT_FILL
    bs["B15"].number_format = MONEY
    bs["D15"].number_format = MONEY
    bs["A17"] = "Check"
    bs["B17"] = '=IF(ABS(B15-D15)<1,"Balanced ✔","Difference of ₹"&TEXT(ABS(B15-D15),"#,##0")&" — see note")'
    bs["A17"].font = Font(bold=True)
    bs["A18"] = ("A difference usually means an opening balance, a loan, an asset purchase or a capital introduction "
                 "has not been entered. The P&L here has no depreciation — your CA will add it when finalising.")
    bs["A18"].alignment = Alignment(wrap_text=True)
    bs.merge_cells("A18:D19")
    bs.row_dimensions[18].height = 30
    bs["A21"] = "Opening stock (1 April 2026)"
    bs["B21"] = 0
    bs["B21"].fill = IN_FILL
    bs["B21"].border = BOX
    bs["B21"].number_format = MONEY

    # Dashboard
    db = wb.create_sheet("Dashboard", 1)
    title(db, "Dashboard — FY 2026-27 so far", "Updates by itself as you fill the registers", 3)
    db.column_dimensions["A"].width = 40
    db.column_dimensions["B"].width = 20
    db.column_dimensions["C"].width = 50
    tiles = [
        ("Sales (taxable value)", "='P&L Summary'!N5", ""),
        ("Purchases (taxable value)", "='P&L Summary'!N6", ""),
        ("Net profit (before stock & depreciation)", f"='P&L Summary'!N{np_row}", ""),
        ("Cash in hand", "='Balance Sheet'!D7", ""),
        ("Bank balance (as per books)", "='Balance Sheet'!D8", ""),
        ("Customers owe you", "='Balance Sheet'!D6", "See 'Receivables' for names and ageing"),
        ("   of which overdue more than 90 days", "=SUM(Receivables!H5:H104)", "Follow up — these are getting hard to recover"),
        ("You owe suppliers", "='Balance Sheet'!B11", "See 'Payables'"),
        ("ITC booked but not yet in GSTR-2B", "=SUM('ITC Tracker'!D5:D16)", "Ask suppliers to file their GSTR-1"),
    ]
    for i, (lab, f, note) in enumerate(tiles):
        r = 4 + i
        db.cell(row=r, column=1, value=lab).font = Font(bold=not lab.startswith("   "))
        c = db.cell(row=r, column=2, value=f)
        c.number_format = MONEY
        c.font = Font(bold=True, size=12, color=NAVY)
        c.fill = TOT_FILL
        c.border = BOX
        db.cell(row=r, column=3, value=note).font = Font(italic=True, color="6B7280")
        db.row_dimensions[r].height = 22


# -------------------------------------------------------- Rent receipt kit --

def rent_receipt_kit():
    wb = Workbook()
    how_to(
        wb,
        "Rent Receipt & HRA Kit",
        [
            "Fill the yellow cells in 'Details' — your name, landlord's name and PAN, address, monthly rent and payment mode.",
            "Open 'Receipts' — 12 monthly receipts (April 2026 to March 2027) are ready. Print on A4.",
            "Get each receipt signed by the landlord. If rent is paid in cash and is more than ₹5,000, a ₹1 revenue stamp is needed.",
            "Submit the receipts to your employer with the investment declaration (Form 12BB), or keep them for your ITR.",
            "Open 'HRA Check' to see how much HRA is exempt and whether the landlord's PAN is required.",
        ],
        [
            "Landlord's PAN is required if the rent you pay is more than ₹1,00,000 in the year.",
            "HRA exemption is available only in the OLD tax regime.",
            "Metro cities for the 50% limit: Delhi, Mumbai, Kolkata and Chennai.",
            "Paying rent to parents is allowed if they own the house and show the rent in their own ITR. Rent to a spouse is generally not accepted.",
        ],
    )
    d = wb.create_sheet("Details", 1)
    title(d, "Details for your rent receipts", None, 2)
    label_rows(d, 3, [
        ("Tenant name (you)", "Your Name"),
        ("Landlord name", "Landlord Name"),
        ("Landlord PAN", "ABCDE1234F"),
        ("Rented house address", "House No., Street, City, PIN"),
        ("Monthly rent (₹)", 15000),
        ("Payment mode", "Bank transfer"),
        ("First month of rent", None),
        ("Metro city? (Y/N)", "N"),
        ("Annual Basic + DA (₹)", 600000),
        ("Annual HRA received (₹)", 240000),
    ])
    d["B9"] = "=DATE(2026,4,1)"
    d["B9"].number_format = "mmmm yyyy"
    d["B9"].fill = IN_FILL
    mode = DataValidation(type="list", formula1='"Bank transfer,UPI,Cheque,Cash"', allow_blank=False)
    d.add_data_validation(mode)
    mode.add("B8")
    yn = DataValidation(type="list", formula1='"Y,N"', allow_blank=False)
    d.add_data_validation(yn)
    yn.add("B10")

    r = wb.create_sheet("Receipts", 2)
    r.column_dimensions["A"].width = 3
    r.column_dimensions["B"].width = 95
    row = 1
    for i in range(12):
        top = row
        r.cell(row=row, column=2, value="RENT RECEIPT").font = Font(bold=True, size=13, color=NAVY)
        r.cell(row=row, column=2).alignment = Alignment(horizontal="center")
        row += 1
        r.cell(row=row, column=2, value=f'="Receipt No. {i + 1:02d}     Date: "&TEXT(EOMONTH(Details!$B$9,{i}),"dd-mm-yyyy")')
        row += 1
        c = r.cell(
            row=row,
            column=2,
            value=(
                f'="Received with thanks from "&Details!$B$3&" a sum of Rs. "&TEXT(Details!$B$7,"#,##0")&"/- "'
                f'&"by "&LOWER(Details!$B$8)&" towards rent for the month of "&TEXT(EDATE(Details!$B$9,{i}),"mmmm yyyy")'
                f'&" for the residential premises at "&Details!$B$6&"."'
            ),
        )
        c.alignment = Alignment(wrap_text=True, vertical="top")
        r.row_dimensions[row].height = 48
        row += 1
        r.cell(row=row, column=2, value='="Landlord: "&Details!$B$4&"     PAN: "&IF(Details!$B$7*12>100000,Details!$B$5,"(not required)")')
        row += 1
        r.cell(row=row, column=2, value='=IF(AND(Details!$B$8="Cash",Details!$B$7>5000),"[ Affix ₹1 revenue stamp ]                                        Signature of landlord","                                                                          Signature of landlord")')
        row += 1
        for rr in range(top, row):
            r.cell(row=rr, column=2).border = Border(
                left=thin, right=thin,
                top=thin if rr == top else None,
                bottom=thin if rr == row - 1 else None,
            )
        row += 2
        if i % 3 == 2:
            r.row_breaks.append(Break(id=row - 1))
    r.page_setup.paperSize = r.PAPERSIZE_A4
    r.page_setup.fitToWidth = 1
    r.sheet_properties.pageSetUpPr.fitToPage = True
    r.page_setup.fitToHeight = 0

    h = wb.create_sheet("HRA Check", 3)
    title(h, "HRA exemption check (old regime)", "Exemption = least of the three amounts below", 2)
    h.column_dimensions["A"].width = 52
    h.column_dimensions["B"].width = 20
    rows = [
        ("Annual rent paid", "=Details!B7*12"),
        ("1. Actual HRA received", "=Details!B12"),
        ("2. Rent paid minus 10% of Basic+DA", "=MAX(0,B4-10%*Details!B11)"),
        ("3. 50% (metro) / 40% (non-metro) of Basic+DA", '=IF(Details!B10="Y",50%,40%)*Details!B11'),
        ("HRA EXEMPT (least of 1, 2, 3)", "=MIN(B5,B6,B7)"),
        ("HRA taxable", "=B5-B8"),
        ("Landlord PAN needed?", '=IF(B4>100000,"Yes — rent is above ₹1,00,000 a year","No")'),
    ]
    for i, (k, f) in enumerate(rows):
        rr = 4 + i
        h.cell(row=rr, column=1, value=k).font = Font(bold=(i in (4, 6)))
        c = h.cell(row=rr, column=2, value=f)
        c.border = BOX
        if i < 6:
            c.number_format = MONEY
        if i == 4:
            c.fill = PatternFill("solid", fgColor="E7F5EC")
            c.font = Font(bold=True)

    wb.active = 1
    wb.save(OUT / "Rent-Receipt-HRA-Kit.xlsx")


# -------------------------------------------------------- Notice reply kit --

def notice_reply_kit():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)

    def h(text, level=1):
        p = doc.add_heading(text, level=level)
        for run in p.runs:
            run.font.color.rgb = RGBColor(0x00, 0x2B, 0x5C)
        return p

    def para(text, bold=False, italic=False, size=None):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        if size:
            run.font.size = Pt(size)
        return p

    def bullets(items):
        for it in items:
            doc.add_paragraph(it, style="List Bullet")

    def letter_head(subject, ref):
        para("[YOUR BUSINESS NAME]\n[Address]\nGSTIN: [__________]     Phone: [__________]     Email: [__________]", bold=True)
        para("Date: [DD-MM-YYYY]")
        para("To,\nThe Proper Officer / [Designation],\n[Ward / Range / Circle],\n[State GST / CGST Office Address]")
        para(f"Subject: {subject}", bold=True)
        para(f"Reference: {ref}")
        para("Respected Sir / Madam,")

    def sign_off(prayer):
        h("Prayer", 3)
        para(prayer)
        para("We also request an opportunity of personal hearing before any adverse view is taken, as required by "
             "the principles of natural justice.")
        para("Verification: I/We hereby solemnly affirm that the information given above is true and correct to the "
             "best of my/our knowledge and belief, and nothing has been concealed therefrom.")
        para("Thanking you,\nYours faithfully,\n\nFor [YOUR BUSINESS NAME]\n\n\n(Authorised Signatory)\nName: [__________]\n"
             "Designation: [Proprietor / Partner / Director]")
        para("Enclosures: [list each document annexed — e.g. Annexure A: reconciliation statement; Annexure B: "
             "copies of invoices; Annexure C: GSTR-2B extract]", italic=True)
        doc.add_page_break()

    # Cover
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = t.add_run("GST Notice Reply Formats")
    run.bold = True
    run.font.size = Pt(26)
    run.font.color.rgb = RGBColor(0x00, 0x2B, 0x5C)
    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = s.add_run("14 ready-to-edit drafts for the most common GST notices, orders and appeals\n" + FIRM)
    r2.font.color.rgb = RGBColor(0xD9, 0x9A, 0x2B)

    h("How to use these formats")
    bullets([
        "Read the notice carefully: note its form (ASMT-10, REG-17, DRC-01A...), reference number, date, tax period "
        "and the last date to reply.",
        "Pick the matching format below. Replace everything in [square brackets] and delete the explanations that do "
        "not apply to you.",
        "Give a point-wise reply to every discrepancy raised — never leave a point unanswered.",
        "Attach supporting documents (reconciliation, invoices, GSTR-2B extract, bank statements) and number them "
        "as annexures.",
        "File the reply on the GST portal: Services → User Services → View Notices and Orders (or View Additional "
        "Notices/Orders), in the form mentioned in the notice, and save the acknowledgement.",
        "If you agree with part of the demand, pay it with interest through Form DRC-03 and mention the ARN in your reply.",
    ])
    h("Time limits to watch", 2)
    bullets([
        "ASMT-10 (scrutiny of returns): reply in Form ASMT-11 within the time given in the notice — normally 30 days.",
        "REG-17 (show cause for cancellation): reply in Form REG-18 within 7 working days.",
        "DRC-01A (intimation before show cause notice): reply in Part B by the date mentioned in the intimation.",
        "DRC-01B (GSTR-1 vs GSTR-3B liability gap) and DRC-01C (ITC in 3B more than 2B): pay or explain in Part B "
        "within 7 days.",
        "DRC-01 (show cause notice / demand): reply in Form DRC-06 within the time stated in the notice — usually 30 days.",
        "REG-03 (query on a new registration application): reply in Form REG-04 within 7 working days.",
        "GSTR-3A (notice to a non-filer): file the pending return within 15 days.",
        "Cancellation already ordered (REG-19)? Apply for revocation in REG-21 within 90 days of the order "
        "(the law allows a further extension in some cases).",
        "Order passed (DRC-07)? Rectification request within 3 months (Section 161); appeal in APL-01 within "
        "3 months of the order (Section 107), with the pre-deposit.",
        "Missed the date? Reply anyway with an adjournment request (Format 6) — silence usually leads to an ex-parte order.",
    ])
    para("Important: These formats are drafting aids based on common situations. Every notice turns on its own "
         "facts and documents. Where the amount involved is significant, or a show cause notice / demand order has "
         "been issued, get the reply reviewed by a professional.", italic=True)
    doc.add_page_break()

    # 1. ASMT-10
    h("Format 1 — Reply to Notice in Form GST ASMT-10 (Scrutiny of Returns)")
    para("Filed in Form GST ASMT-11 under Section 61 of the CGST Act, 2017 read with Rule 99.", italic=True)
    letter_head(
        "Reply to notice in Form GST ASMT-10 for the period [MM/YYYY – MM/YYYY]",
        "Notice No. [__________] dated [DD-MM-YYYY]; GSTIN [__________]",
    )
    para("1. We refer to the above notice in which certain discrepancies have been pointed out in the returns filed "
         "by us for the period [__________]. We thank your good office for the opportunity to explain, and submit "
         "our point-wise reply as under.")
    table = doc.add_table(rows=1, cols=4)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    for i, t in enumerate(["S.No.", "Discrepancy as per notice", "Our explanation", "Supporting document"]):
        hdr[i].text = t
    for n in range(1, 4):
        row = table.add_row().cells
        row[0].text = str(n)
        row[1].text = "[Copy the discrepancy exactly as stated, with amount]"
        row[2].text = "[Explain — see sample explanations below]"
        row[3].text = f"Annexure [{chr(64 + n)}]"
    para("")
    para("Sample explanations (use only those that are true for you):", bold=True)
    bullets([
        "The difference is due to invoices of [month] reported in GSTR-1 of [later month] through amendment; tax on "
        "them was paid in GSTR-3B of [month]. Reconciliation at Annexure [ ].",
        "Credit notes of ₹[ ] issued under Section 34 were reported in GSTR-1 but adjusted in GSTR-3B of [month].",
        "The amount was a typographical error in Table [ ] of GSTR-3B; the correct liability was paid in [month] "
        "(DRC-03 ARN [ ] / GSTR-3B of [month]).",
        "Input tax credit claimed relates to invoices uploaded late by the supplier; all conditions of Section 16 "
        "are met and the credit was claimed within the time limit of Section 16(4).",
        "The turnover difference with e-way bills is due to e-way bills generated for [job work / stock transfer / "
        "returns / cancelled invoices], which are not supplies.",
    ])
    para("2. In view of the above, the discrepancies stand fully explained. Where any short payment was found, we "
         "have paid tax of ₹[ ] with interest of ₹[ ] vide DRC-03 ARN [ ] dated [ ].")
    sign_off("We therefore request your good office to accept this explanation and drop the proceedings, as "
             "provided in Section 61(2) of the CGST Act, 2017 and Rule 99(2), by issuing Form GST ASMT-12.")

    # 2. GSTR-1 vs 3B
    h("Format 2 — Reply on Mismatch between GSTR-1 and GSTR-3B (Output Tax)")
    letter_head(
        "Explanation for difference between outward supplies reported in GSTR-1 and GSTR-3B for [period]",
        "Notice / Intimation No. [__________] dated [DD-MM-YYYY]; GSTIN [__________]",
    )
    para("1. The notice points out that tax liability in GSTR-1 (₹[ ]) is higher than tax paid in GSTR-3B (₹[ ]), "
         "a difference of ₹[ ].")
    para("2. The month-wise reconciliation is attached as Annexure A. The difference arises for these reasons:")
    table = doc.add_table(rows=1, cols=3)
    table.style = "Light Grid Accent 1"
    for i, t in enumerate(["Reason", "Amount (₹)", "Where the tax is paid / adjusted"]):
        table.rows[0].cells[i].text = t
    for reason in ["Invoices amended in later GSTR-1", "Credit notes / debit notes timing", "Advances received and adjusted",
                   "Duplicate / wrongly reported invoices in GSTR-1", "Balance short payment (if any)"]:
        row = table.add_row().cells
        row[0].text = reason
        row[1].text = "[ ]"
        row[2].text = "[GSTR-3B of month / DRC-03 ARN]"
    para("")
    para("3. The balance difference of ₹[ ], if any, has been paid with applicable interest under Section 50 through "
         "DRC-03 ARN [ ].")
    sign_off("We request that the reconciliation be accepted and no further liability be proposed for the period.")

    # 3. ITC mismatch
    h("Format 3 — Reply on Excess ITC: GSTR-3B vs GSTR-2A/2B")
    para("Use for scrutiny notices, and for the Part B reply to an intimation in Form DRC-01C (Rule 88D).", italic=True)
    letter_head(
        "Explanation for input tax credit claimed in GSTR-3B in excess of GSTR-2B for [period]",
        "Notice / DRC-01C reference No. [__________] dated [DD-MM-YYYY]; GSTIN [__________]",
    )
    para("1. The notice states that ITC of ₹[ ] claimed in GSTR-3B exceeds the ITC available in GSTR-2B by ₹[ ]. "
         "Supplier-wise reconciliation is enclosed as Annexure A.")
    para("2. The difference is explained as follows:", bold=True)
    bullets([
        "ITC of ₹[ ] relates to invoices uploaded by suppliers in a later period (list with GSTIN, invoice no., "
        "date and the period in which they appear in GSTR-2B at Annexure B).",
        "ITC of ₹[ ] is on imports of goods (Bill of Entry no. [ ]), which is claimed on the basis of the Bill of "
        "Entry and appears in GSTR-2B only through ICEGATE data.",
        "ITC of ₹[ ] was claimed under reverse charge after paying the tax in cash; it does not appear in GSTR-2B "
        "by design.",
        "ITC of ₹[ ] was reversed in GSTR-3B of [month] (Table 4B) under Rule 37 / 42 / 43; net ITC availed is "
        "therefore within GSTR-2B.",
    ])
    para("3. All conditions of Section 16(2) — possession of tax invoice, receipt of goods/services, payment of tax "
         "by supplier and filing of return — are satisfied, and the credit was availed within the time limit of "
         "Section 16(4).")
    para("4. For the balance of ₹[ ] for which no supporting document is available, we have reversed/paid the ITC "
         "with interest through DRC-03 ARN [ ].")
    sign_off("We request your good office to accept the above explanation and drop the proceedings.")

    # 4. REG-17
    h("Format 4 — Reply to Show Cause Notice for Cancellation of Registration (Form GST REG-17)")
    para("Reply in Form GST REG-18 within 7 working days (Rule 22).", italic=True)
    letter_head(
        "Reply to show cause notice in Form GST REG-17 for cancellation of registration",
        "SCN reference No. [__________] dated [DD-MM-YYYY]; GSTIN [__________]",
    )
    para("1. The notice proposes cancellation of our registration on the ground that [copy the ground, e.g. "
         "'returns have not been furnished for a continuous period of six months' / 'the business is not being "
         "conducted from the declared place of business'].")
    para("2. Our reply (keep the one that applies):", bold=True)
    bullets([
        "Non-filing of returns: All pending returns up to [month] have since been filed along with late fee and "
        "interest (ARNs listed at Annexure A). There is no return pending as on date.",
        "Place of business: We continue to carry on business from [address]. Photographs of the premises with name "
        "board, the latest electricity bill, rent agreement and recent purchase/sale invoices are enclosed as "
        "Annexure B. We shall be happy to facilitate a physical verification.",
        "Other grounds: [explain with documents].",
    ])
    para("3. The business is active and we intend to comply with all provisions of the Act going forward.")
    sign_off("We request that the show cause notice be dropped by issuing an order in Form GST REG-20 and that "
             "our registration be continued. If the registration has been suspended, we request that the "
             "suspension be revoked.")

    # 5. DRC-01A
    h("Format 5 — Reply to Intimation in Form GST DRC-01A (before Show Cause Notice)")
    para("Filed in Part B of DRC-01A. Either pay the ascertained amount through DRC-03, or give your submissions "
         "against it — or both, for different parts.", italic=True)
    letter_head(
        "Submissions against intimation of tax ascertained in Form GST DRC-01A for [period]",
        "Intimation No. [__________] dated [DD-MM-YYYY]; GSTIN [__________]",
    )
    para("1. By the above intimation, tax of ₹[ ], interest of ₹[ ] and penalty of ₹[ ] has been ascertained on "
         "the ground that [summarise].")
    para("2. Part of the amount accepted (if any): We accept liability of ₹[ ] relating to [ ], which has been "
         "paid with interest vide DRC-03 ARN [ ] dated [ ].")
    para("3. Part of the amount not accepted: The balance of ₹[ ] is not payable for the following reasons:")
    bullets([
        "[Point 1 — facts, with reference to the annexure proving it]",
        "[Point 2 — legal provision relied upon, e.g. Section 16, Section 17(5), notification no.]",
        "[Point 3 — any calculation error in the intimation]",
    ])
    sign_off("We request that no show cause notice be issued for the amount not accepted, and that the "
             "proceedings for the amount paid be concluded under Section 73(6) / 74(6) (or the corresponding "
             "provision for the relevant period).")

    # 6. Adjournment
    h("Format 6 — Request for Adjournment / Extension of Time")
    letter_head(
        "Request for extension of time to file reply to notice dated [DD-MM-YYYY]",
        "Notice No. [__________]; GSTIN [__________]",
    )
    para("1. We have received the above notice requiring us to [file a reply / appear for personal hearing] by "
         "[date].")
    para("2. The records needed for a complete reply — [books of account / supplier confirmations / bank "
         "statements] — are being compiled and reconciled [/ our authorised signatory is unwell / our consultant is "
         "reviewing the matter]. A complete and correct reply needs some more time.")
    para("3. We therefore request an extension of [15] days, up to [date], to file our reply. We assure full "
         "cooperation. As per Section 75(5), adjournment may be granted for sufficient cause.")
    para("Thanking you,\nYours faithfully,\n\nFor [YOUR BUSINESS NAME]\n\n(Authorised Signatory)")
    doc.add_page_break()

    # 7. Cover letter
    h("Format 7 — Cover Letter for Submitting Documents")
    letter_head(
        "Submission of documents in connection with notice dated [DD-MM-YYYY]",
        "Notice No. [__________]; GSTIN [__________]",
    )
    para("With reference to the above notice [and the personal hearing held on (date)], we submit the following "
         "documents for your kind consideration:")
    table = doc.add_table(rows=1, cols=3)
    table.style = "Light Grid Accent 1"
    for i, t in enumerate(["Annexure", "Document", "Pages"]):
        table.rows[0].cells[i].text = t
    for a in "ABCD":
        row = table.add_row().cells
        row[0].text = a
        row[1].text = "[ ]"
        row[2].text = "[ ]"
    para("")
    para("We request that these documents be taken on record. We shall be glad to furnish any further information "
         "required.")
    para("Thanking you,\nYours faithfully,\n\nFor [YOUR BUSINESS NAME]\n\n(Authorised Signatory)")

    doc.add_page_break()
    notice_formats_more(doc, h, para, bullets, letter_head, sign_off)

    para("")
    para("Need help with a notice? Rajput Lalit & Associates drafts and files GST notice replies — "
         "WhatsApp +91 93549 53603 · www.rajputlalitassociates.in/gst-notice-reply", bold=True)

    doc.save(OUT / "GST-Notice-Reply-Formats.docx")
    notice_recon_workbook()


def notice_formats_more(doc, h, para, bullets, letter_head, sign_off):
    """Formats 8–14 (added in the ₹2,999 edition)."""

    # 8. DRC-01 show cause notice
    h("Format 8 — Reply to Show Cause Notice in Form GST DRC-01 (Demand)")
    para("Filed in Form GST DRC-06 under Section 73 / 74 (periods up to FY 2023-24) or Section 74A (FY 2024-25 "
         "onwards). This is the most important reply — whatever is not raised here is hard to raise later in appeal.",
         italic=True)
    letter_head(
        "Reply to Show Cause Notice in Form GST DRC-01 for the period [__________]",
        "SCN No. [__________] dated [DD-MM-YYYY]; DIN [__________]; GSTIN [__________]",
    )
    h("A. Preliminary submissions (use only those that apply)", 3)
    bullets([
        "Limitation: the notice for [period] has been issued on [date], beyond the time limit prescribed for that "
        "period, and is therefore barred by limitation.",
        "The notice does not bear a valid Document Identification Number (DIN) / is not digitally signed [where "
        "applicable].",
        "The notice does not give the basis of calculation, nor the documents relied upon. We request copies of all "
        "relied-upon documents and reserve the right to file an additional reply after receiving them.",
        "Section 74 / extended period is invoked without any specific allegation or evidence of fraud, wilful "
        "misstatement or suppression of facts. All transactions are recorded in our books and returns.",
        "No DRC-01A intimation / pre-notice consultation was given before the notice [if relevant for the period].",
    ])
    h("B. Submissions on merits — point-wise", 3)
    para("Para [ ] of the SCN — [summarise the allegation and amount]")
    para("Our reply: [facts → documents (Annexure) → legal provision / circular / judgment relied upon → conclusion "
         "that the amount is not payable]")
    para("(Repeat for every paragraph of the notice. Do not leave any allegation unanswered.)", italic=True)
    h("C. Interest and penalty", 3)
    para("As no tax is payable for the reasons above, interest under Section 50 and penalty do not arise. Without "
         "prejudice, [if part is admitted]: tax of ₹[ ] with interest of ₹[ ] has been paid vide DRC-03 ARN [ ] dated "
         "[ ], within the period allowed for paying without penalty — the proceedings for that part may be concluded.")
    para("Tip: in non-fraud cases, paying the admitted tax with interest within 30 days of the SCN (Section 73) or "
         "within 60 days (Section 74A) generally means no penalty for that part.", italic=True)
    sign_off("We pray that the show cause notice be dropped in full, and that no demand of tax, interest or penalty "
             "be confirmed. We request a personal hearing under Section 75(4) before any order is passed.")

    # 9. GSTR-3A
    h("Format 9 — Reply to Notice in Form GSTR-3A (Return Not Filed)")
    para("Issued under Section 46. File the pending return(s) first — the letter only puts it on record.", italic=True)
    letter_head(
        "Reply to notice in Form GSTR-3A for the tax period [MM/YYYY]",
        "Notice No. [__________] dated [DD-MM-YYYY]; GSTIN [__________]",
    )
    para("1. The return in Form [GSTR-3B / GSTR-1] for [period] could not be filed by the due date because [brief, "
         "genuine reason — e.g. delay in finalising accounts / illness of the proprietor / technical issue].")
    para("2. The return has now been filed on [date] vide ARN [__________], with tax of ₹[ ], interest of ₹[ ] and "
         "late fee of ₹[ ].")
    para("3. Since the return has been filed with all dues within the time allowed, we request that the "
         "proceedings be dropped. If an assessment order in Form ASMT-13 has been passed, it stands withdrawn under "
         "Section 62(2) as the return has been furnished within the period allowed by that section.")
    para("Thanking you,\nYours faithfully,\n\nFor [YOUR BUSINESS NAME]\n\n(Authorised Signatory)")
    doc.add_page_break()

    # 10. REG-03
    h("Format 10 — Reply to Query on New Registration (Form GST REG-03 → reply in REG-04)")
    para("Reply within 7 working days or the application is likely to be rejected. Answer exactly what is asked "
         "and upload a clear document for each point.", italic=True)
    letter_head(
        "Clarification in Form GST REG-04 against notice in Form GST REG-03",
        "Application Reference No. (ARN) [__________] dated [DD-MM-YYYY]",
    )
    table = doc.add_table(rows=1, cols=3)
    table.style = "Light Grid Accent 1"
    for i, t in enumerate(["Query raised", "Our clarification", "Document uploaded"]):
        table.rows[0].cells[i].text = t
    for q, a, d in [
        ("Proof of principal place of business not clear",
         "The premises at [address] are [owned / rented]. [Rent agreement dated __ / ownership proof] uploaded.",
         "Rent agreement + latest electricity bill of the owner"),
        ("Consent / NOC of the owner",
         "The owner [name] has given consent for use of the premises as our place of business.",
         "Signed NOC + owner's ID proof"),
        ("Photographs of the business premises",
         "Photographs showing the name board and inside of the premises are uploaded.",
         "Geo-tagged photographs"),
        ("[Any other query]", "[Reply]", "[Document]"),
    ]:
        row = table.add_row().cells
        row[0].text, row[1].text, row[2].text = q, a, d
    para("")
    para("We request that the application be approved. We shall provide any further information required.")
    para("Thanking you,\nYours faithfully,\n\n[Name of applicant]\n(Proprietor / Authorised Signatory)")
    doc.add_page_break()

    # 11. REG-21 revocation
    h("Format 11 — Application for Revocation of Cancelled Registration (Form GST REG-21)")
    para("Under Section 30 read with Rule 23. File all pending returns and pay dues first — the portal and the "
         "officer will check this. Time limit: 90 days from the date of the cancellation order.", italic=True)
    letter_head(
        "Application for revocation of cancellation of registration",
        "Cancellation Order No. [__________] dated [DD-MM-YYYY] (Form REG-19); GSTIN [__________]",
    )
    para("1. Our registration was cancelled on the ground that [returns were not filed for (period) / business not "
         "found at the declared place / other].")
    para("2. The reason for the default was [genuine reason]. The business is running and we wish to continue "
         "complying with GST.")
    para("3. All returns up to the date of cancellation have now been filed, with tax, interest and late fee "
         "paid, as below:")
    t = doc.add_table(rows=1, cols=4)
    t.style = "Light Grid Accent 1"
    for i, x in enumerate(["Return", "Period", "Date filed", "ARN"]):
        t.rows[0].cells[i].text = x
    for _ in range(3):
        row = t.add_row().cells
        for c in row:
            c.text = "[ ]"
    para("")
    para("4. [If the ground was 'not found at the address':] The business continues at [address]; photographs, rent "
         "agreement and electricity bill are attached.")
    para("We request that the cancellation be revoked under Section 30 and the registration restored.")
    para("Thanking you,\nYours faithfully,\n\nFor [YOUR BUSINESS NAME]\n\n(Authorised Signatory)")
    doc.add_page_break()

    # 12. Written submissions after hearing
    h("Format 12 — Written Submissions after Personal Hearing")
    letter_head(
        "Written submissions following the personal hearing held on [DD-MM-YYYY]",
        "SCN / Notice No. [__________] dated [DD-MM-YYYY]; GSTIN [__________]",
    )
    para("1. We thank you for the personal hearing granted on [date], attended by [name, designation / authorised "
         "representative].")
    para("2. During the hearing, the following points were discussed and the following documents were sought:")
    bullets(["[Point / document 1]", "[Point / document 2]"])
    para("3. Our submissions on these points: [answer each point; refer to annexures]")
    para("4. We reiterate our reply dated [date] and request that the proceedings be dropped.")
    para("Thanking you,\nYours faithfully,\n\nFor [YOUR BUSINESS NAME]\n\n(Authorised Signatory)")
    doc.add_page_break()

    # 13. Rectification
    h("Format 13 — Request for Rectification of Order (Section 161)")
    para("For mistakes apparent on the record — wrong figures, a payment already made but not credited, a reply "
         "filed but not considered. Apply within 3 months of the order.", italic=True)
    letter_head(
        "Application for rectification of order under Section 161",
        "Order No. [__________] dated [DD-MM-YYYY] (Form DRC-07 / [ ]); GSTIN [__________]",
    )
    para("1. The above order contains the following error(s) apparent on the face of the record:")
    bullets([
        "Tax of ₹[ ] already paid vide [GSTR-3B / DRC-03 ARN ( )] dated [ ] has not been given credit.",
        "Our reply dated [date] (ARN [ ]) has not been considered in the order.",
        "Arithmetical error: [explain the correct calculation].",
    ])
    para("2. We request that the order be rectified as above and the demand be reduced to ₹[ ] / dropped. A "
         "corrected summary in Form DRC-07 may kindly be uploaded on the portal.")
    para("Thanking you,\nYours faithfully,\n\nFor [YOUR BUSINESS NAME]\n\n(Authorised Signatory)")
    doc.add_page_break()

    # 14. Appeal skeleton
    h("Format 14 — Statement of Facts & Grounds of Appeal (Form GST APL-01)")
    para("Appeal to the Appellate Authority under Section 107 within 3 months of the order (a further 1 month can be "
         "condoned for sufficient cause). Pre-deposit: the full admitted amount plus 10% of the disputed tax. "
         "Appeals are technical — we strongly suggest professional help.", italic=True)
    h("Statement of facts", 3)
    bullets([
        "The appellant is engaged in [business] and is registered under GSTIN [ ].",
        "Notice [type, no., date] was issued alleging [ ]. The appellant replied on [date] (Annexure [ ]).",
        "By the impugned order dated [date], the adjudicating authority confirmed a demand of tax ₹[ ], interest ₹[ ] "
        "and penalty ₹[ ].",
        "Being aggrieved, the appellant files this appeal on the following grounds, which are without prejudice to "
        "one another.",
    ])
    h("Grounds of appeal", 3)
    bullets([
        "A. The order is non-speaking: it does not deal with the submissions made in the reply dated [date].",
        "B. The order violates the principles of natural justice: no / inadequate personal hearing was given "
        "(Section 75(4)).",
        "C. The order travels beyond the show cause notice (Section 75(7)).",
        "D. On merits: [main legal and factual ground — with documents, provisions, circulars and judgments].",
        "E. The demand is barred by limitation for [period].",
        "F. Interest and penalty are not leviable as the tax itself is not payable; in any case, no penalty is "
        "justified as there was no intent to evade tax.",
        "The appellant craves leave to add, alter or amend any of the grounds at the time of hearing.",
    ])
    h("Prayer", 3)
    para("Set aside the impugned order, drop the demand of tax, interest and penalty, grant a personal hearing, and "
         "pass any other order deemed fit.")
    doc.add_page_break()


def notice_recon_workbook():
    """Excel working papers that go with the notice reply formats."""
    wb = Workbook()
    how_to(
        wb,
        "GST Notice Reconciliation Workbook",
        [
            "Most GST notices are about one of three gaps: books vs GSTR-1, GSTR-1 vs GSTR-3B, or ITC in 3B vs GSTR-2B. "
            "Fill the matching sheet month by month from your books and the portal.",
            "Write a short reason against every difference in the 'Reason' column — this becomes your reconciliation "
            "annexure. Print it or save it as PDF and attach it to the reply.",
            "'Demand & Interest' works out interest at 18% to any payment date and the pre-deposit needed for an appeal.",
            "'Notice Tracker' keeps every notice, its reply date and status in one place.",
        ],
        ["Use the Word formats for the reply letter; this workbook is for the numbers behind it."],
    )

    def recon(name, sub, cols, a_label, b_label, c_label=None):
        ws = wb.create_sheet(name)
        title(ws, name, sub, len(cols))
        header(ws, 4, cols, [12] + [16] * (len(cols) - 2) + [44])
        for i, (y, m) in enumerate(FY_MONTHS):
            r = 5 + i
            ws.cell(row=r, column=1, value=f"=DATE({y},{m},1)").number_format = "mmm-yy"
            input_block(ws, r, r, (2, 3) + ((4,) if c_label else ()))
            if c_label:
                ws.cell(row=r, column=5, value=f"=B{r}-C{r}")
                ws.cell(row=r, column=6, value=f"=C{r}-D{r}")
                money_cols(ws, r, range(2, 7))
                input_block(ws, r, r, (7,))
            else:
                ws.cell(row=r, column=4, value=f"=B{r}-C{r}")
                money_cols(ws, r, range(2, 5))
                input_block(ws, r, r, (5,))
        tr = 5 + len(FY_MONTHS)
        ws.cell(row=tr, column=1, value="Total").font = Font(bold=True)
        for col in range(2, len(cols)):
            L = get_column_letter(col)
            c = ws.cell(row=tr, column=col, value=f"=SUM({L}5:{L}{tr - 1})")
            c.font = Font(bold=True)
            c.fill = TOT_FILL
            c.number_format = MONEY
        ws.freeze_panes = "B5"

    recon("Output Tax Recon", "Total tax (CGST+SGST+IGST+cess) per month. Notices DRC-01B / ASMT-10 usually come from these gaps.",
          ["Month", "Tax as per books", "Tax as per GSTR-1", "Tax as per GSTR-3B", "Books − GSTR-1", "GSTR-1 − GSTR-3B",
           "Reason for difference"], "Books", "GSTR-1", "GSTR-3B")
    recon("ITC Recon", "ITC per month. A positive 'claimed − 2B' figure is what DRC-01C / ASMT-10 questions.",
          ["Month", "ITC claimed in GSTR-3B", "ITC as per GSTR-2B", "Claimed − 2B", "Reason for difference"],
          "3B", "2B")

    dm = wb.create_sheet("Demand & Interest")
    title(dm, "Demand, Interest & Appeal Pre-deposit", "Fill the yellow cells", 3)
    dm.column_dimensions["A"].width = 48
    dm.column_dimensions["B"].width = 18
    dm.column_dimensions["C"].width = 60
    rows = [
        ("Tax demanded in the notice / order", 0, True, MONEY),
        ("Of which you accept (will pay)", 0, True, MONEY),
        ("Original due date of that tax", None, True, DATE),
        ("Date you pay / expect to pay", None, True, DATE),
        ("Disputed tax", "=B4-B5", False, MONEY),
        ("Days of delay", '=IF(OR(B6="",B7=""),0,MAX(0,B7-B6))', False, "0"),
        ("Interest @ 18% p.a. on accepted tax", "=ROUND(B5*18%*B9/365,0)", False, MONEY),
        ("Pay now through DRC-03 (accepted tax + interest)", "=B5+B10", False, MONEY),
        ("Appeal pre-deposit (10% of disputed tax)", "=ROUND(B8*10%,0)", False, MONEY),
    ]
    for i, (k, v, inp, fmt) in enumerate(rows):
        r = 4 + i
        dm.cell(row=r, column=1, value=k).font = Font(bold=True)
        c = dm.cell(row=r, column=2, value=v)
        c.number_format = fmt
        c.fill = IN_FILL if inp else TOT_FILL
        c.border = BOX
    dm["C10"] = "Section 50: 18% per year from the day after the due date until payment."
    dm["C12"] = ("Section 107(6): admitted amount in full + 10% of the disputed tax (subject to the statutory upper "
                 "limit). Interest/penalty-only appeals have their own rule — check before filing.")
    for ref in ("C10", "C12"):
        dm[ref].alignment = Alignment(wrap_text=True)
    dm.row_dimensions[12].height = 45

    nt = wb.create_sheet("Notice Tracker")
    title(nt, "Notice Tracker", "One row per notice. Status turns red on its own as the reply date comes close.", 9)
    header(nt, 4, ["Form (ASMT-10, DRC-01…)", "Reference no.", "Notice date", "Period", "Amount ₹", "Reply due by",
                   "Replied on (ARN)", "Days left", "Status"], [20, 20, 12, 14, 14, 13, 20, 10, 24])
    for r in range(5, 55):
        input_block(nt, r, r, range(1, 8))
        nt.cell(row=r, column=3).number_format = DATE
        nt.cell(row=r, column=6).number_format = DATE
        nt.cell(row=r, column=5).number_format = MONEY
        nt.cell(row=r, column=8, value=f'=IF(OR(F{r}="",G{r}<>""),"",F{r}-TODAY())')
        nt.cell(row=r, column=9, value=(f'=IF(A{r}="","",IF(G{r}<>"","Replied",IF(F{r}="","Enter reply date",'
                                        f'IF(H{r}<0,"OVERDUE — reply / seek time now",IF(H{r}<=7,"URGENT","Open")))))'))
    nt.freeze_panes = "A5"

    wb.active = 1
    wb.save(OUT / "GST-Notice-Reconciliation-Workbook.xlsx")


# ------------------------------------------------------------ ITR organizer --

def itr_organizer_kit():
    wb = Workbook()
    how_to(
        wb,
        "ITR Filing Checklist & Organizer",
        [
            "Start with 'Which ITR Form' to see which form fits your income.",
            "Go through 'Document Checklist' — mark each item Have / Need / N.A. The progress bar at the top tells you "
            "when you are ready to file.",
            "Fill 'Income Organizer' from your Form 16, bank statements, rent agreement and capital-gains statements.",
            "Fill 'Deductions' only if you will choose the old tax regime (most deductions are not available in the "
            "new regime). Compare both regimes on our free calculator: www.rajputlalitassociates.in/income-tax-calculator",
            "Download your AIS and Form 26AS from the e-filing portal and enter their figures in 'AIS & 26AS Match'. "
            "Every difference must be explained or corrected before you file — that is how most notices are avoided.",
            "Keep 'Filing Details' updated with the bank account for refund, the acknowledgement number and the date "
            "you e-verified.",
        ],
        [
            "Due dates for tax year 2026-27 (FY April 2026 – March 2027): 31 July 2027 for salaried and other "
            "individuals without business income; 31 August 2027 for business/professional income without audit; "
            "31 October 2027 for audit cases.",
            "The return is complete only after e-verification (Aadhaar OTP / net banking / DSC) — do it within 30 days "
            "of filing.",
            "From 1 April 2026 the Income-tax Act, 2025 replaced the 1961 Act. Section numbers changed, but most "
            "concepts here (heads of income, AIS, regimes) stay the same. ITR forms are notified every year — confirm "
            "eligibility on the e-filing portal before filing.",
            "Works for any year — just change the year in 'Filing Details'.",
        ],
    )

    wf = wb.create_sheet("Which ITR Form", 1)
    title(wf, "Which ITR form should you file?", "General guide for individuals and HUFs — confirm on the portal for your year", 3)
    header(wf, 4, ["Form", "Who can use it", "Who cannot use it"], [12, 70, 60])
    forms = [
        ("ITR-1 (Sahaj)",
         "Resident individual with total income up to ₹50 lakh from salary/pension, house property, interest/"
         "dividends/family pension, agricultural income up to ₹5,000, and long-term gains on listed shares/equity "
         "funds up to ₹1.25 lakh.",
         "NRIs and RNOR; directors of a company; holders of unlisted shares; anyone with foreign assets/income, "
         "business income, or other capital gains."),
        ("ITR-2",
         "Individuals and HUFs with no business/professional income — e.g. capital gains of any kind, more than "
         "₹50 lakh income, NRIs, foreign assets, directors, unlisted shares.",
         "Anyone with business or professional income."),
        ("ITR-3",
         "Individuals and HUFs with business or professional income computed from books of account, partners in "
         "a firm, or those who opt out of the presumptive scheme.",
         "—"),
        ("ITR-4 (Sugam)",
         "Resident individuals, HUFs and firms (not LLPs) with total income up to ₹50 lakh and business/"
         "professional income under the presumptive scheme (old Sections 44AD / 44ADA / 44AE).",
         "NRIs; directors; holders of unlisted shares; those with foreign assets; income above ₹50 lakh."),
    ]
    for i, row in enumerate(forms):
        r = 5 + i
        for j, v in enumerate(row):
            c = wf.cell(row=r, column=1 + j, value=v)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = BOX
        wf.cell(row=r, column=1).font = Font(bold=True, color=NAVY)
        wf.row_dimensions[r].height = 75
    wf["A10"] = "Not sure? WhatsApp +91 93549 53603 — we'll tell you the right form free of charge."
    wf["A10"].font = Font(bold=True, color=GOLD)

    # Checklist
    cl = wb.create_sheet("Document Checklist", 2)
    title(cl, "Document Checklist", "Mark each line: Have / Need / N.A.", 4)
    header(cl, 5, ["Section", "Document", "Status", "Notes"], [22, 70, 12, 40])
    groups = {
        "Basic": ["PAN and Aadhaar (linked)", "Login for the income tax e-filing portal",
                  "Bank account details (pre-validated on the portal for refund)",
                  "Last year's ITR acknowledgement (if any)"],
        "AIS / 26AS": ["Annual Information Statement (AIS) and Taxpayer Information Summary (TIS)",
                       "Form 26AS (TDS/TCS and tax paid)"],
        "Salary / pension": ["Form 16 from each employer in the year", "Salary slips (for allowances / HRA)",
                             "Rent receipts and landlord's PAN (if rent above ₹1 lakh a year — old regime HRA)",
                             "Pension certificate (pensioners)"],
        "House property": ["Rent received / rent agreement for let-out property", "Municipal tax paid receipts",
                           "Home loan interest certificate (bank)", "Possession / completion date of the house"],
        "Interest & dividends": ["Savings and FD interest certificates (all banks / post office)",
                                 "Dividend statements", "Interest on income tax refund (shown in AIS)"],
        "Capital gains": ["Capital gains statement from broker / mutual fund (CAMS / KFintech)",
                          "Property sale deed and purchase deed, cost of improvement bills",
                          "Proof of investment for exemption (new house / specified bonds)"],
        "Business / profession": ["Bank statements of business accounts", "Sales and purchase summary / GST returns",
                                  "Profit & loss and balance sheet (or gross receipts for presumptive)",
                                  "Advance tax challans"],
        "Deductions (old regime)": ["Investments: PPF, ELSS, LIC, EPF, tuition fees, home loan principal",
                                    "Health insurance premium receipts (self, family, parents)",
                                    "NPS contribution statement", "Donation receipts with donee PAN",
                                    "Education loan interest certificate"],
        "Other": ["Foreign bank accounts / assets (residents must disclose)", "Crypto / virtual digital asset statements",
                  "Details of any income tax notice received"],
    }
    r = 6
    first = r
    st = DataValidation(type="list", formula1='"Have,Need,N.A."', allow_blank=True)
    cl.add_data_validation(st)
    for g, docs in groups.items():
        for d in docs:
            cl.cell(row=r, column=1, value=g).font = Font(bold=True, color=NAVY)
            cl.cell(row=r, column=2, value=d).alignment = Alignment(wrap_text=True)
            input_block(cl, r, r, (3, 4))
            r += 1
    last = r - 1
    st.add(f"C{first}:C{last}")
    cl["A3"] = "Ready:"
    cl["A3"].font = Font(bold=True)
    cl["B3"] = (f'=IF(COUNTA(C{first}:C{last})=0,"Start marking the list below",'
                f'TEXT(COUNTIF(C{first}:C{last},"Have")/MAX(1,COUNTA(C{first}:C{last})-COUNTIF(C{first}:C{last},"N.A.")),"0%")'
                f'&" collected — "&COUNTIF(C{first}:C{last},"Need")&" item(s) still needed")')
    cl["B3"].font = Font(bold=True, color=GOLD, size=12)
    cl.freeze_panes = "A6"

    # Income organizer
    io = wb.create_sheet("Income Organizer", 3)
    title(io, "Income Organizer", "Enter yearly figures in rupees", 3)
    io.column_dimensions["A"].width = 56
    io.column_dimensions["B"].width = 18
    io.column_dimensions["C"].width = 50
    sections = [
        ("Salary / pension", ["Gross salary (Form 16, Part B)", "Exempt allowances claimed (old regime: HRA, LTA…)",
                              "Professional tax paid", "Pension"]),
        ("House property", ["Annual rent received", "Municipal tax paid", "Home loan interest — let-out property",
                            "Home loan interest — self-occupied property"]),
        ("Other sources", ["Savings bank interest", "FD / RD / post office interest", "Dividends",
                           "Family pension", "Any other income"]),
        ("Capital gains", ["Short-term gains — listed shares / equity funds", "Long-term gains — listed shares / equity funds",
                           "Gains on property / gold / debt funds / others"]),
        ("Business / profession", ["Gross receipts / turnover", "Net profit as per books (or presumptive income)"]),
    ]
    r = 4
    sum_rows = []
    for sec, items in sections:
        c = io.cell(row=r, column=1, value=sec)
        c.font = H_FONT
        c.fill = H_FILL
        io.cell(row=r, column=2).fill = H_FILL
        io.cell(row=r, column=3).fill = H_FILL
        r += 1
        for it in items:
            io.cell(row=r, column=1, value=it)
            input_block(io, r, r, (2, 3))
            io.cell(row=r, column=2).number_format = MONEY
            sum_rows.append(r)
            r += 1
        r += 1
    io.cell(row=r, column=1, value="Total of all figures entered (for reference only — not taxable income)").font = Font(italic=True)
    io.cell(row=r, column=2, value="=" + "+".join(f"B{x}" for x in sum_rows)).number_format = MONEY

    # Deductions
    de = wb.create_sheet("Deductions", 4)
    title(de, "Deductions — only if you choose the OLD regime",
          "In the new regime only a few items (standard deduction, employer's NPS contribution) are allowed.", 4)
    header(de, 4, ["Deduction", "Limit (old regime)", "Amount you have", "Allowed"], [58, 20, 18, 16])
    deds = [
        ("Investments — PPF, ELSS, LIC, EPF, tuition fees, home loan principal (earlier Section 80C)", 150000),
        ("Additional NPS contribution (earlier Section 80CCD(1B))", 50000),
        ("Health insurance — self & family, below 60 (earlier 80D)", 25000),
        ("Health insurance — parents (₹50,000 if parents are senior citizens)", 50000),
        ("Interest on home loan — self-occupied house", 200000),
        ("Savings bank interest (earlier 80TTA; 80TTB up to ₹50,000 for senior citizens)", 10000),
        ("Education loan interest (no upper limit)", None),
        ("Donations (eligible part)", None),
    ]
    for i, (k, lim) in enumerate(deds):
        r = 5 + i
        de.cell(row=r, column=1, value=k).alignment = Alignment(wrap_text=True)
        de.cell(row=r, column=2, value=lim if lim else "No fixed limit").number_format = MONEY
        input_block(de, r, r, (3,))
        de.cell(row=r, column=3).number_format = MONEY
        de.cell(row=r, column=4, value=f'=IF(ISNUMBER(B{r}),MIN(B{r},N(C{r})),N(C{r}))').number_format = MONEY
        de.row_dimensions[r].height = 30
    tr = 5 + len(deds)
    de.cell(row=tr, column=1, value="Total deductions (old regime)").font = Font(bold=True)
    c = de.cell(row=tr, column=4, value=f"=SUM(D5:D{tr - 1})")
    c.number_format = MONEY
    c.font = Font(bold=True)
    c.fill = TOT_FILL
    de.cell(row=tr + 2, column=1, value=(
        "Limits shown are the long-standing old-regime figures; the 2025 Act renumbered these sections. Confirm the "
        "limits for your year before filing.")).alignment = Alignment(wrap_text=True)
    de.merge_cells(start_row=tr + 2, start_column=1, end_row=tr + 2, end_column=4)
    de.row_dimensions[tr + 2].height = 30

    # AIS match
    am = wb.create_sheet("AIS & 26AS Match", 5)
    title(am, "AIS & Form 26AS — match before you file", "Enter what AIS/26AS shows and what your records show", 6)
    header(am, 4, ["Item", "As per AIS / 26AS", "As per your records", "Difference", "Action", "Notes"],
           [40, 18, 18, 16, 34, 30])
    items = ["Salary", "TDS on salary", "Savings interest", "FD interest", "Dividend", "Rent received", "TDS on rent",
             "Sale of shares / mutual funds", "Sale of property", "TDS on property sale", "Business receipts",
             "TDS on professional / contract receipts", "Advance tax / self-assessment tax paid", "Other"]
    for i, it in enumerate(items):
        r = 5 + i
        am.cell(row=r, column=1, value=it)
        input_block(am, r, r, (2, 3, 6))
        am.cell(row=r, column=4, value=f'=IF(AND(B{r}="",C{r}=""),"",N(B{r})-N(C{r}))')
        am.cell(row=r, column=5, value=(f'=IF(D{r}="","",IF(ABS(D{r})<1,"OK",IF(D{r}>0,'
                                        f'"Report it, or give feedback on AIS if wrong","Check — your figure is higher")))'))
        money_cols(am, r, (2, 3, 4))

    # Filing details
    fd = wb.create_sheet("Filing Details", 6)
    title(fd, "Filing Details", "Keep this for your records", 2)
    label_rows(fd, 3, [
        ("Tax year", "2026-27 (FY April 2026 – March 2027)"),
        ("Name", ""), ("PAN", ""), ("Regime chosen (new / old)", "New"), ("ITR form", ""),
        ("Bank for refund (name, last 4 digits)", ""), ("Date filed", ""), ("Acknowledgement no.", ""),
        ("Date e-verified", ""), ("Refund / tax paid", ""), ("Filed by", ""),
    ], col_w=(38, 46))

    wb.active = 1
    wb.save(OUT / "ITR-Filing-Checklist-Organizer.xlsx")


# --------------------------------------------------------- Freelancer kit --

def freelancer_tax_kit():
    wb = Workbook()
    how_to(
        wb,
        "Freelancer Tax Kit — Invoices, Income Register, Presumptive Tax & Advance Tax",
        [
            "Fill 'Settings' once: your name, PAN, address, bank/SWIFT details and — if GST registered — GSTIN and LUT ARN.",
            "Make invoices in 'Invoice'. For foreign clients pick 'Export'; the LUT line is printed automatically when "
            "you have a LUT ARN. For Indian clients pick 'Domestic' — 18% GST is added only if you are GST registered.",
            "Log every payment you receive in 'Income Register' (foreign amount, exchange rate, INR credited, mode, TDS).",
            "Log business expenses in 'Expenses' — needed if you compare presumptive vs actual profit.",
            "'Presumptive Tax' shows whether you can use the 50% presumptive scheme, your taxable income and an "
            "estimate of tax under the new regime.",
            "'Advance Tax' tells you how much to pay and by when. 'GST Check' warns you as you get close to the "
            "₹20 lakh registration limit.",
        ],
        [
            "Presumptive scheme for professionals: Section 58 of the Income-tax Act, 2025 (old Section 44ADA). 50% of "
            "gross receipts is treated as income; limit ₹75 lakh if cash receipts are within 5% of total, else ₹50 lakh. "
            "It covers 'specified professions' (IT, design, consulting, technical, etc.) — check that yours qualifies.",
            "Presumptive taxpayers can pay the whole advance tax by 15 March.",
            "Export of services is zero-rated. If you are GST registered, file the LUT (RFD-11) every financial year to "
            "invoice at 0% without paying IGST.",
            "Keep FIRA / bank advice for every foreign receipt — it proves export of services.",
            "Tax estimate uses FY 2026-27 new-regime slabs (unchanged from FY 2025-26) and ignores surcharge (income "
            "above ₹50 lakh) and special-rate income. It's a planning aid, not a tax computation.",
        ],
    )

    s = wb.create_sheet("Settings", 1)
    title(s, "Your details", "Fill once", 2)
    label_rows(s, 3, [
        ("Name / trade name", "Your Name"),
        ("Address", "City, State - PIN, India"),
        ("PAN", "ABCDE1234F"),
        ("Email / phone", ""),
        ("GST registered? (Y/N)", "N"),
        ("GSTIN (if registered)", ""),
        ("LUT ARN for FY 2026-27 (if any)", ""),
        ("Bank name & account no.", ""),
        ("IFSC", ""),
        ("SWIFT code (for foreign payments)", ""),
        ("Profession", "Software development / design / consulting"),
    ])
    dv = DataValidation(type="list", formula1='"Y,N"', allow_blank=False)
    s.add_data_validation(dv)
    dv.add("B7")

    # Invoice
    iv = wb.create_sheet("Invoice", 2)
    for col, w in zip("ABCDEF", (6, 44, 10, 14, 16, 4)):
        iv.column_dimensions[col].width = w
    iv.merge_cells("A1:E1")
    iv["A1"] = "INVOICE"
    iv["A1"].font = Font(bold=True, size=18, color="FFFFFF")
    iv["A1"].fill = H_FILL
    iv["A1"].alignment = Alignment(horizontal="center")
    iv["A3"] = "=Settings!B3"
    iv["A3"].font = Font(bold=True, size=13, color=NAVY)
    iv["A4"] = "=Settings!B4"
    iv["A5"] = '="PAN: "&Settings!B5&IF(Settings!B7="Y","   GSTIN: "&Settings!B8,"")'
    iv["A6"] = "=Settings!B6"
    meta = [("Invoice No.", "FL/26-27/001"), ("Date", None), ("Type (Export / Domestic)", "Export"),
            ("Currency", "USD"), ("Exchange rate (₹ per unit)", 85)]
    for i, (k, v) in enumerate(meta):
        r = 3 + i
        iv.cell(row=r, column=4, value=k).font = Font(bold=True)
        c = iv.cell(row=r, column=5, value=v)
        c.fill = IN_FILL
        c.border = BOX
    iv["E4"].number_format = DATE
    tdv = DataValidation(type="list", formula1='"Export,Domestic"', allow_blank=False)
    iv.add_data_validation(tdv)
    tdv.add("E5")
    iv["A9"] = "Bill to:"
    iv["A9"].font = Font(bold=True)
    for r in range(10, 13):
        iv.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        iv.cell(row=r, column=1).fill = IN_FILL
    iv["A10"] = "Client name"
    iv["A11"] = "Client address, Country"
    iv["A12"] = "Client tax ID / GSTIN (if any)"
    header(iv, 14, ["#", "Description of service", "Qty / hrs", "Rate", "Amount"])
    for i in range(1, 9):
        r = 14 + i
        iv.cell(row=r, column=1, value=i)
        input_block(iv, r, r, (2, 3, 4))
        iv.cell(row=r, column=5, value=f'=IF(OR(C{r}="",D{r}=""),"",C{r}*D{r})').number_format = MONEY
        iv.cell(row=r, column=4).number_format = MONEY
    iv["D24"] = "Sub-total"
    iv["E24"] = "=SUM(E15:E22)"
    iv["D25"] = '=IF(AND(E5="Domestic",Settings!B7="Y"),"GST @ 18%","GST")'
    iv["E25"] = '=IF(AND(E5="Domestic",Settings!B7="Y"),ROUND(E24*18%,2),0)'
    iv["D26"] = '="TOTAL ("&IF(E5="Export",E6,"INR")&")"'
    iv["E26"] = "=E24+E25"
    iv["D27"] = "Approx. value in INR"
    iv["E27"] = '=IF(E5="Export",ROUND(E26*E7,0),E26)'
    for ref in ("E24", "E25", "E26", "E27"):
        iv[ref].number_format = MONEY
    for ref in ("D26", "E26"):
        iv[ref].font = Font(bold=True, color=NAVY, size=12)
    iv.merge_cells("A29:E29")
    iv["A29"] = ('=IF(E5="Export",IF(AND(Settings!B7="Y",Settings!B9<>""),"Supply meant for export of services under '
                 'LUT without payment of IGST — LUT ARN "&Settings!B9,"Export of services — zero-rated supply"),"")')
    iv["A29"].font = Font(italic=True)
    iv.merge_cells("A31:E33")
    iv["A31"] = '="Payment to: "&Settings!B10&"  |  IFSC: "&Settings!B11&"  |  SWIFT: "&Settings!B12'
    iv["A31"].alignment = Alignment(wrap_text=True, vertical="top")
    iv.page_setup.paperSize = iv.PAPERSIZE_A4
    iv.sheet_properties.pageSetUpPr.fitToPage = True
    iv.page_setup.fitToHeight = 1

    # Income register
    ir = wb.create_sheet("Income Register", 3)
    title(ir, "Income Register — every receipt", "One row per payment received", 12)
    header(ir, 4, ["Date received", "Client", "Country", "Invoice no.", "Currency", "Foreign amount", "Exchange rate",
                   "INR credited", "Mode (Bank/UPI/Cash)", "TDS deducted (Indian clients)", "FIRA received? (Y/N)",
                   "Month"], [12, 26, 14, 14, 10, 14, 12, 16, 14, 14, 12, 10])
    md = DataValidation(type="list", formula1='"Bank,UPI,Cash,Platform"', allow_blank=True)
    ir.add_data_validation(md)
    md.add("I5:I504")
    fy = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
    ir.add_data_validation(fy)
    fy.add("K5:K504")
    for r in range(5, 505):
        input_block(ir, r, r, range(1, 12))
        ir.cell(row=r, column=1).number_format = DATE
        money_cols(ir, r, (6, 8, 10))
        ir.cell(row=r, column=12, value=f'=IF(A{r}="","",TEXT(A{r},"mmm-yy"))')
    ir.freeze_panes = "A5"

    ex = wb.create_sheet("Expenses", 4)
    title(ex, "Business Expenses", "Laptop, software, internet, co-working, travel for work…", 5)
    header(ex, 4, ["Date", "Expense", "Paid to", "Amount (₹)", "Mode"], [12, 34, 26, 14, 12])
    for r in range(5, 305):
        input_block(ex, r, r, range(1, 6))
        ex.cell(row=r, column=1).number_format = DATE
        ex.cell(row=r, column=4).number_format = MONEY
    ex.freeze_panes = "A5"

    # Presumptive tax
    pt = wb.create_sheet("Presumptive Tax", 5)
    title(pt, "Presumptive Tax (Section 58 / old 44ADA) — FY 2026-27", "Fills from your registers; yellow = optional entries", 3)
    pt.column_dimensions["A"].width = 52
    pt.column_dimensions["B"].width = 20
    pt.column_dimensions["C"].width = 56
    IRH = "'Income Register'!$H$5:$H$504"
    rows = [
        ("Gross receipts (INR)", f"=SUM({IRH})", None),
        ("Receipts in cash", f"=SUMIF('Income Register'!$I$5:$I$504,\"Cash\",{IRH})", None),
        ("Cash share of receipts", "=IF(B4=0,0,B5/B4)", "0.0%"),
        ("Presumptive limit that applies to you", "=IF(B6<=5%,7500000,5000000)", None),
        ("Eligible for the presumptive scheme?", '=IF(B4<=B7,"Yes","No — receipts above the limit; books & audit rules apply")', "@"),
        ("Presumptive income (50% of receipts)", "=ROUND(B4*50%,0)", None),
        ("Actual profit (receipts − expenses)", "=B4-SUM(Expenses!D5:D304)", None),
        ("Income you may declare (higher of 50% or actual if you want)", "=MAX(B9,0)", None),
        ("Other income (interest, rent, salary etc.)", 0, "in"),
        ("Total income", "=B11+B12", None),
        ("Tax on total income — new regime slabs", (
            "=MAX(0,MIN(B13,800000)-400000)*5%+MAX(0,MIN(B13,1200000)-800000)*10%+MAX(0,MIN(B13,1600000)-1200000)*15%"
            "+MAX(0,MIN(B13,2000000)-1600000)*20%+MAX(0,MIN(B13,2400000)-2000000)*25%+MAX(0,B13-2400000)*30%"), None),
        ("After rebate / marginal relief (income up to ₹12 lakh)", "=IF(B13<=1200000,0,MIN(B14,B13-1200000))", None),
        ("Health & education cess @ 4%", "=ROUND(B15*4%,0)", None),
        ("ESTIMATED TAX FOR THE YEAR", "=B15+B16", None),
        ("Less: TDS already deducted", f"=SUM('Income Register'!$J$5:$J$504)", None),
        ("Tax still to pay (advance tax)", "=MAX(0,B17-B18)", None),
    ]
    for i, (k, f, fmt) in enumerate(rows):
        r = 4 + i
        pt.cell(row=r, column=1, value=k).font = Font(bold=k.isupper() or k.startswith("Tax still"))
        c = pt.cell(row=r, column=2, value=f)
        if fmt == "in":
            c.fill = IN_FILL
            c.border = BOX
            c.number_format = MONEY
        else:
            c.fill = TOT_FILL
            c.number_format = fmt or MONEY
    pt["B17"].font = Font(bold=True, color=NAVY, size=12)
    pt["C9"] = "You may declare more than 50% — never less, unless you keep books (and get them audited if required)."
    pt["C11"] = "Change this cell if you want to declare actual profit when it is higher than 50%."
    pt["B11"].fill = IN_FILL
    pt["C15"] = ("New regime: no tax if total income is up to ₹12 lakh (rebate). Slightly above ₹12 lakh, tax "
                 "can't exceed the income above ₹12 lakh (marginal relief).")
    pt["C17"] = "Surcharge (income above ₹50 lakh) and special-rate income (capital gains) are not included."
    for ref in ("C9", "C11", "C15", "C17"):
        pt[ref].alignment = Alignment(wrap_text=True, vertical="top")
        pt[ref].font = Font(italic=True, color="6B7280")
    for rr in (9, 11, 15, 17):
        pt.row_dimensions[rr].height = 30

    # Advance tax
    at = wb.create_sheet("Advance Tax", 6)
    title(at, "Advance Tax Planner — FY 2026-27", "Due only if the year's tax after TDS is ₹10,000 or more", 6)
    header(at, 4, ["Due date", "Cumulative % due", "Amount due by this date", "Paid on (date)", "Amount paid",
                   "Short / (excess)"], [14, 18, 20, 14, 16, 18])
    at["H4"] = "Using the presumptive scheme? (Y/N)"
    at["H4"].font = Font(bold=True)
    at["I4"] = "Y"
    at["I4"].fill = IN_FILL
    at.column_dimensions["H"].width = 36
    pdv = DataValidation(type="list", formula1='"Y,N"', allow_blank=False)
    at.add_data_validation(pdv)
    pdv.add("I4")
    sched = [("=DATE(2026,6,15)", 0.15), ("=DATE(2026,9,15)", 0.45), ("=DATE(2026,12,15)", 0.75), ("=DATE(2027,3,15)", 1.0)]
    for i, (d, p) in enumerate(sched):
        r = 5 + i
        at.cell(row=r, column=1, value=d).number_format = DATE
        at.cell(row=r, column=2, value=(f'=IF($I$4="Y",{1 if p == 1.0 else 0},{p})')).number_format = "0%"
        at.cell(row=r, column=3, value=f"=IF('Presumptive Tax'!$B$19<10000,0,ROUND('Presumptive Tax'!$B$19*B{r},0))").number_format = MONEY
        input_block(at, r, r, (4, 5))
        at.cell(row=r, column=4).number_format = DATE
        at.cell(row=r, column=5).number_format = MONEY
        at.cell(row=r, column=6, value=f"=C{r}-SUM($E$5:E{r})").number_format = MONEY
    at["A11"] = ("Pay online: e-filing portal → e-Pay Tax → Advance tax. Interest applies if you pay less than the "
                 "cumulative amount by each date (presumptive: by 15 March), and if total paid by 31 March is under 90%.")
    at["A11"].alignment = Alignment(wrap_text=True)
    at.merge_cells("A11:F12")
    at.row_dimensions[11].height = 30

    # GST check
    gc = wb.create_sheet("GST Check", 7)
    title(gc, "GST Registration Check", "Running total of receipts this financial year", 3)
    gc.column_dimensions["A"].width = 44
    gc.column_dimensions["B"].width = 20
    gc.column_dimensions["C"].width = 60
    gc["A4"] = "Turnover so far this FY (from Income Register)"
    gc["B4"] = "='Presumptive Tax'!B4"
    gc["A5"] = "Registration limit for service providers"
    gc["B5"] = 2000000
    gc["B5"].fill = IN_FILL
    gc["C5"] = "₹20 lakh in most states (₹10 lakh in some special-category states). Change if it applies to you."
    gc["A6"] = "Used so far"
    gc["B6"] = "=IF(B5=0,0,B4/B5)"
    gc["B6"].number_format = "0%"
    gc["A7"] = "Status"
    gc["B7"] = ('=IF(Settings!B7="Y","Registered — remember LUT every April",IF(B6>=1,"Register now (within 30 days)",'
                'IF(B6>=0.8,"Close to the limit — plan registration + LUT","Below the limit")))')
    for ref in ("B4", "B5"):
        gc[ref].number_format = MONEY
    gc["B7"].font = Font(bold=True, color=NAVY)
    gc["A9"] = ("Export of services counts towards the limit. Freelancers exporting services are not required to register "
                "until aggregate turnover crosses the limit (Notification 10/2017-Integrated Tax).")
    gc["A9"].alignment = Alignment(wrap_text=True)
    gc.merge_cells("A9:C10")
    gc.row_dimensions[9].height = 30

    wb.active = 1
    wb.save(OUT / "Freelancer-Tax-Kit.xlsx")


if __name__ == "__main__":
    # Old file names from the first edition — replaced by the upgraded kits below.
    for old in ("GST-Invoice-Billing-Kit.xlsx", "Small-Business-Bookkeeping-Kit.xlsx"):
        (OUT / old).unlink(missing_ok=True)
    gst_invoice_kit()
    bookkeeping_kit()
    rent_receipt_kit()
    notice_reply_kit()
    itr_organizer_kit()
    freelancer_tax_kit()
    for f in sorted(OUT.iterdir()):
        print(f"{f.name}: {f.stat().st_size:,} bytes")
